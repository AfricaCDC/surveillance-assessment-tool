from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "System_Usability_Scale_Questionnaire.docx"
NAVY = "16324F"
BLUE = "2E74B5"
LIGHT = "E8EEF5"
PALE = "F4F6F9"
GRAY = "5E6873"
WHITE = "FFFFFF"

items = [
    "I think that I would like to use this web application frequently.",
    "I found the web application unnecessarily complex.",
    "I thought the web application was easy to use.",
    "I think that I would need the support of a technical person to be able to use this web application.",
    "I found the various functions in this web application were well integrated.",
    "I thought there was too much inconsistency in this web application.",
    "I would imagine that most people would learn to use this web application very quickly.",
    "I found the web application very cumbersome to use.",
    "I felt very confident using the web application.",
    "I needed to learn a lot of things before I could get going with this web application.",
]

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=100, start=100, bottom=100, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v)); node.set(qn("w:type"), "dxa")

def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)

def set_table_widths(table, widths_dxa, indent=120):
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW"); tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths_dxa))); tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd"); tblPr.append(tblInd)
    tblInd.set(qn("w:w"), str(indent)); tblInd.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol"); col.set(qn("w:w"), str(width)); grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tcW = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW"); cell._tc.get_or_add_tcPr().append(tcW)
            tcW.set(qn("w:w"), str(widths_dxa[idx])); tcW.set(qn("w:type"), "dxa")

def style_run(run, size=11, bold=False, color="000000", italic=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size); run.bold = bold; run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)

def add_field_line(doc, label, width=52):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(label + "  " + "_" * width)
    style_run(r, 10.5)

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5); sec.page_height = Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"; normal.font.size = Pt(11); normal.font.color.rgb = RGBColor(0,0,0)
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri"); normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.paragraph_format.space_before = Pt(0); normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.1
for name, size, before, after, color in [("Heading 1",16,16,8,BLUE),("Heading 2",13,12,6,BLUE),("Heading 3",12,8,4,NAVY)]:
    st=styles[name]; st.font.name="Calibri"; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color)
    st._element.rPr.rFonts.set(qn("w:ascii"),"Calibri"); st._element.rPr.rFonts.set(qn("w:hAnsi"),"Calibri")
    st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after)

# Running header/footer
hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = hp.add_run("SYSTEM USABILITY SCALE | QUESTIONNAIRE")
style_run(hr, 8.5, True, GRAY)
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run("Confidential research instrument")
style_run(fr, 8.5, False, GRAY)

# Masthead
p = doc.add_paragraph(); p.paragraph_format.space_after=Pt(4)
r=p.add_run("SYSTEM USABILITY SCALE"); style_run(r,24,True,NAVY)
p = doc.add_paragraph(); p.paragraph_format.space_after=Pt(14)
r=p.add_run("Web Application Assessment Questionnaire"); style_run(r,14,False,GRAY)

meta = doc.add_table(rows=3, cols=2); meta.alignment=WD_TABLE_ALIGNMENT.LEFT
set_table_widths(meta,[4680,4680],120)
fields=[("Web application name","[Enter application name]"),("Study/project","[Enter study or project]"),("Assessment date","____ / ____ / ______"),("Participant ID","____________________"),("Device used","Desktop / Laptop / Tablet / Phone"),("Browser","____________________")]
for i,(label,value) in enumerate(fields):
    cell=meta.cell(i//2,i%2); set_cell_shading(cell,PALE); set_cell_margins(cell,110,140,110,140)
    p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(1)
    rr=p.add_run(label.upper()); style_run(rr,8.5,True,BLUE)
    p=cell.add_paragraph(); p.paragraph_format.space_after=Pt(0)
    rr=p.add_run(value); style_run(rr,10.5)

doc.add_heading("Participant instructions", level=1)
p=doc.add_paragraph("Please complete this questionnaire immediately after using the web application. For each statement, select one response based on your first impression. Answer every item. Do not spend too long considering individual statements.")
p.paragraph_format.keep_with_next=True

scale = doc.add_table(rows=1, cols=5); scale.alignment=WD_TABLE_ALIGNMENT.LEFT
set_table_widths(scale,[1872]*5,120)
labels=["1\nStrongly disagree","2\nDisagree","3\nNeither agree nor disagree","4\nAgree","5\nStrongly agree"]
for i,label in enumerate(labels):
    c=scale.cell(0,i); set_cell_shading(c,LIGHT); set_cell_margins(c,100,80,100,80); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0)
    rr=p.add_run(label); style_run(rr,9,True,NAVY)

doc.add_heading("Usability statements", level=1)
table=doc.add_table(rows=1, cols=7); table.style="Table Grid"; table.alignment=WD_TABLE_ALIGNMENT.LEFT
widths=[500,5860,600,600,600,600,600]; set_table_widths(table,widths,120)
headers=["No.","Statement","1","2","3","4","5"]
for i,h in enumerate(headers):
    c=table.cell(0,i); set_cell_shading(c,NAVY); set_cell_margins(c,110,90,110,90); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0)
    rr=p.add_run(h); style_run(rr,9.5,True,WHITE)
set_repeat_header(table.rows[0])
for idx,item in enumerate(items,1):
    cells=table.add_row().cells
    for c in cells: set_cell_margins(c,115,90,115,90); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p=cells[0].paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0); style_run(p.add_run(str(idx)),10,True,NAVY)
    p=cells[1].paragraphs[0]; p.paragraph_format.space_after=Pt(0); style_run(p.add_run(item),9.5)
    for j in range(2,7):
        p=cells[j].paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0); style_run(p.add_run("☐"),14,False,NAVY)

doc.add_heading("Optional comments", level=1)
for prompt in ["What did you like most about the web application?","What was the most difficult or frustrating part?","What is the single most important improvement you would recommend?"]:
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(3); style_run(p.add_run(prompt),10.5,True,NAVY)
    for _ in range(2):
        p=doc.add_paragraph("_"*96); p.paragraph_format.space_after=Pt(2); style_run(p.runs[0],9,False,GRAY)

doc.add_page_break()
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(4)
r=p.add_run("SCORING GUIDE"); style_run(r,22,True,NAVY)
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(14)
r=p.add_run("For researchers and administrators - do not show while the participant is responding"); style_run(r,10.5,False,GRAY,True)

doc.add_heading("Calculate the SUS score",level=1)
steps=[
    "For odd-numbered items (1, 3, 5, 7 and 9), subtract 1 from the participant's response.",
    "For even-numbered items (2, 4, 6, 8 and 10), subtract the participant's response from 5.",
    "Add the ten adjusted item scores. The adjusted total will range from 0 to 40.",
    "Multiply the adjusted total by 2.5. The final SUS score ranges from 0 to 100.",
]
for i,s in enumerate(steps,1):
    p=doc.add_paragraph(style="List Number"); p.paragraph_format.space_after=Pt(8); p.paragraph_format.line_spacing=1.167
    style_run(p.add_run(s),11)

calc=doc.add_table(rows=4,cols=2); calc.style="Table Grid"; calc.alignment=WD_TABLE_ALIGNMENT.LEFT
set_table_widths(calc,[3600,5760],120)
calc_rows=[("Odd-item adjusted total","__________  (maximum 20)"),("Even-item adjusted total","__________  (maximum 20)"),("Combined adjusted total","__________  (maximum 40)"),("Final SUS score","__________ × 2.5 = __________ / 100")]
for i,(a,b) in enumerate(calc_rows):
    for c in calc.rows[i].cells: set_cell_margins(c,130,140,130,140); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if i==3:
        set_cell_shading(calc.cell(i,0),LIGHT); set_cell_shading(calc.cell(i,1),LIGHT)
    p=calc.cell(i,0).paragraphs[0]; p.paragraph_format.space_after=Pt(0); style_run(p.add_run(a),10.5,True,NAVY)
    p=calc.cell(i,1).paragraphs[0]; p.paragraph_format.space_after=Pt(0); style_run(p.add_run(b),10.5)

doc.add_heading("Interpretation notes",level=1)
for text in [
    "A SUS score is a standardized usability score, not a percentage of tasks completed correctly.",
    "Higher scores indicate better perceived usability. Interpret scores alongside the study context, sample size, user characteristics, task-completion data and qualitative comments.",
    "Report the mean, median, standard deviation, range and number of valid responses. Define in advance how missing item responses will be handled.",
]:
    p=doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after=Pt(8); p.paragraph_format.line_spacing=1.167; style_run(p.add_run(text),11)

doc.add_heading("Administration record",level=1)
for label in ["Number of valid questionnaires:","Mean SUS score:","Median SUS score:","Standard deviation:","Notes:"]:
    add_field_line(doc,label,52 if label!="Notes:" else 78)

doc.add_heading("Instrument reference",level=1)
p=doc.add_paragraph("Brooke, J. (1996). SUS: A 'quick and dirty' usability scale. In P. W. Jordan, B. Thomas, B. A. Weerdmeester, & A. L. McClelland (Eds.), Usability Evaluation in Industry (pp. 189-194). Taylor & Francis.")
p.paragraph_format.space_after=Pt(0); style_run(p.runs[0],9.5,False,GRAY)

doc.core_properties.title="System Usability Scale Assessment Questionnaire"
doc.core_properties.subject="Web application usability assessment"
doc.core_properties.author=""
doc.save(OUT)
print(OUT)

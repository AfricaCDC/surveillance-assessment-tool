from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "Application_Refinement_and_Regression_Testing_Register.docx"
NAVY, BLUE, PALE, LIGHT, GRAY, WHITE = "16324F", "2E74B5", "F4F6F9", "E8EEF5", "5E6873", "FFFFFF"

def font(run, size=10, bold=False, color="000000", italic=False):
    run.font.name="Calibri"; run.font.size=Pt(size); run.bold=bold; run.italic=italic
    run.font.color.rgb=RGBColor.from_string(color)
    rpr=run._element.get_or_add_rPr(); rpr.rFonts.set(qn("w:ascii"),"Calibri"); rpr.rFonts.set(qn("w:hAnsi"),"Calibri")

def shade(cell, color):
    tcpr=cell._tc.get_or_add_tcPr(); shd=tcpr.find(qn("w:shd"))
    if shd is None: shd=OxmlElement("w:shd"); tcpr.append(shd)
    shd.set(qn("w:fill"),color)

def margins(cell, top=80, start=100, bottom=80, end=100):
    tcpr=cell._tc.get_or_add_tcPr(); mar=tcpr.first_child_found_in("w:tcMar")
    if mar is None: mar=OxmlElement("w:tcMar"); tcpr.append(mar)
    for key,val in (("top",top),("start",start),("bottom",bottom),("end",end)):
        el=mar.find(qn(f"w:{key}"))
        if el is None: el=OxmlElement(f"w:{key}"); mar.append(el)
        el.set(qn("w:w"),str(val)); el.set(qn("w:type"),"dxa")

def geometry(table, widths, indent=120):
    table.autofit=False; tblpr=table._tbl.tblPr
    tw=tblpr.find(qn("w:tblW"))
    if tw is None: tw=OxmlElement("w:tblW"); tblpr.append(tw)
    tw.set(qn("w:w"),str(sum(widths))); tw.set(qn("w:type"),"dxa")
    ti=tblpr.find(qn("w:tblInd"))
    if ti is None: ti=OxmlElement("w:tblInd"); tblpr.append(ti)
    ti.set(qn("w:w"),str(indent)); ti.set(qn("w:type"),"dxa")
    grid=table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col=OxmlElement("w:gridCol"); col.set(qn("w:w"),str(width)); grid.append(col)
    for row in table.rows:
        for i,cell in enumerate(row.cells):
            tcpr=cell._tc.get_or_add_tcPr(); tcw=tcpr.find(qn("w:tcW"))
            if tcw is None: tcw=OxmlElement("w:tcW"); tcpr.append(tcw)
            tcw.set(qn("w:w"),str(widths[i])); tcw.set(qn("w:type"),"dxa")

def repeat_header(row):
    trpr=row._tr.get_or_add_trPr(); el=OxmlElement("w:tblHeader"); el.set(qn("w:val"),"true"); trpr.append(el)

def add_table(doc, headers, widths, rows, font_size=8.3):
    table=doc.add_table(rows=1,cols=len(headers)); table.style="Table Grid"; table.alignment=WD_TABLE_ALIGNMENT.LEFT
    for i,h in enumerate(headers):
        c=table.cell(0,i); shade(c,NAVY); margins(c,90,90,90,90); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0); font(p.add_run(h),8.2,True,WHITE)
    repeat_header(table.rows[0])
    for values in rows:
        cells=table.add_row().cells
        for i,c in enumerate(cells):
            margins(c,90,90,90,90); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if len(table.rows)%2==1: shade(c,"FAFBFC")
            p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
            if i==0 or (headers[i] in {"Priority","Status","Result","Severity"}): p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            font(p.add_run(str(values[i]) if i<len(values) else ""),font_size,False,"000000")
    geometry(table,widths)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return table

def heading(doc,text,level=1):
    p=doc.add_heading(text,level=level); p.paragraph_format.keep_with_next=True; return p

def page_title(doc,title,subtitle=None):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(8); p.paragraph_format.space_after=Pt(4)
    font(p.add_run(title),23,True,NAVY)
    if subtitle:
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(14); font(p.add_run(subtitle),13,False,GRAY)

doc=Document(); sec=doc.sections[0]
sec.orientation=WD_ORIENT.LANDSCAPE; sec.page_width=Inches(11); sec.page_height=Inches(8.5)
sec.top_margin=sec.bottom_margin=Inches(.65); sec.left_margin=sec.right_margin=Inches(.65)
sec.header_distance=sec.footer_distance=Inches(.35)
normal=doc.styles["Normal"]; normal.font.name="Calibri"; normal.font.size=Pt(10); normal.paragraph_format.space_after=Pt(5); normal.paragraph_format.line_spacing=1.1
normal._element.rPr.rFonts.set(qn("w:ascii"),"Calibri"); normal._element.rPr.rFonts.set(qn("w:hAnsi"),"Calibri")
for name,size,before,after,color in (("Heading 1",16,14,7,BLUE),("Heading 2",13,10,5,BLUE),("Heading 3",11,7,4,NAVY)):
    st=doc.styles[name]; st.font.name="Calibri"; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color)
    st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.keep_with_next=True

hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; font(hp.add_run("APPLICATION QUALITY ASSURANCE | CONTROLLED REGISTER"),8,True,GRAY)
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(fp.add_run("Africa CDC Surveillance Digital Tools Assessment Web Application"),8,False,GRAY)

page_title(doc,"APPLICATION REFINEMENT AND REGRESSION TESTING REGISTER","Change control, test execution, defect management and release approval")
meta=[("Application","Africa CDC Surveillance Digital Tools Assessment Web Application","Document owner","____________________________"),("Version under test","1.0.40","Test cycle","____________________________"),("Environment","Development / Test / UAT / Production","Testing period","From __________ to __________"),("REDCap project","____________________________","Prepared by","____________________________")]
add_table(doc,["Field","Details","Field","Details"],[1500,3780,1500,3780],meta,9)

heading(doc,"Purpose",1)
doc.add_paragraph("Use this register to document proposed refinements, approved changes, regression test cases, execution evidence, defects, retesting and release decisions. Complete one copy for each application release or formal test cycle.")
heading(doc,"Status and severity conventions",2)
add_table(doc,["Category","Permitted values","Meaning"],[1800,3100,5660],[
    ("Change status","Proposed | Approved | In progress | Implemented | Deferred | Rejected","Current position of a requested application change."),
    ("Test result","Not run | Pass | Fail | Blocked | Not applicable","Outcome of a test execution."),
    ("Defect severity","Critical | High | Medium | Low","Impact on data integrity, security, workflow completion or presentation."),
    ("Release decision","Approved | Approved with conditions | Not approved","Final authorization following testing and risk review."),
],8.8)

doc.add_page_break(); page_title(doc,"1. CHANGE AND REFINEMENT LOG","Record each requested improvement and its approval history")
change_headers=["Change ID","Date raised","Source","Area / feature","Problem or requested change","Priority","Decision / rationale","Owner","Target version","Status"]
change_widths=[650,800,900,1050,2200,700,1800,850,850,810]
change_rows=[
    ("CHG-001","2026-08-26","User feedback","Profiling import","Map all standard questions, per-question comments, custom multi-select values, and Opportunities & Challenges responses.","Critical","Implemented and regression tested.","","1.0.32","Implemented"),
    ("CHG-002","2026-08-26","User feedback","Profiling report","Use the brief five-part summary and count only explicit national-level coverage responses.","High","Implemented and regression tested.","","1.0.34","Implemented"),
    ("CHG-003","2026-08-26","User feedback","Report assistant","Complete evidence-linked consolidated recommendations without unsupported attribution or truncation.","High","Implemented with compact local-AI evidence.","","1.0.36","Implemented"),
    ("CHG-004","2026-08-26","Data protection","Test Country","Keep Test Country temporary in memory and prevent every REDCap write path.","Critical","Implemented; existing test record deleted.","","1.0.37","Implemented"),
    ("CHG-005","2026-08-26","Authentication","Password stability","Use the adjacent .env as the authoritative password-security configuration for every launch mode.","Critical","Implemented; existing password rehashed and login verified.","","1.0.38","Implemented"),
    ("CHG-006","2026-08-26","User feedback","Profiling report","Separate affirmative answer counts from system counts and clarify Inventory versus Profiling denominators.","High","Implemented with an explicit interpretation note.","","1.0.39","Implemented"),
]+[(f"CHG-{i:03d}","","","","","","","","","") for i in range(7,11)]
add_table(doc,change_headers,change_widths,change_rows,7.6)

doc.add_page_break(); page_title(doc,"2. REGRESSION TEST CASE CATALOGUE","Reusable baseline tests for every release")
test_headers=["Test ID","Area","Test objective","Preconditions / test data","Test steps","Expected result","Priority"]
test_widths=[650,1050,1800,1800,2600,2400,750]
tests=[
    ("TC-001","Authentication","Confirm a valid active user can sign in.","Active test account exists.","Enter valid credentials and select Sign in.","Dashboard opens with the correct role and permitted functions.","Critical"),
    ("TC-002","Authorization","Confirm group users cannot modify another group's assigned work.","Two groups and assignments exist.","Sign in as Group A; attempt to access or save Group B work.","Access is denied and Group B data remains unchanged.","Critical"),
    ("TC-003","REDCap import","Load the correct country assessment and reporting period.","Unique REDCap test record exists.","Choose country and period; select Load REDCap data.","Correct record, tools, profiles, assignments and responses load.","Critical"),
    ("TC-004","Inventory","Add and save a new surveillance tool.","Approved persistent UAT assessment is open.","Add a tool; complete fields; save.","REDCap confirms the save and the tool appears in the linked record.","High"),
    ("TC-005","Inventory","Edit an imported tool without affecting unrelated tools.","Assessment contains at least two tools.","Change one field in Tool A; save; inspect Tool A and Tool B in REDCap.","Tool A changes; unrelated Tool B values remain correct.","Critical"),
    ("TC-006","Assignments","Assign and unassign tools from working groups.","Tools and groups exist.","Assign selected tools; save; unassign one tool.","Assignments display correctly and synchronize to REDCap.","High"),
    ("TC-007","Profiling","Edit and save an imported system profile.","Profile linked to a REDCap record exists.","Change one profile response; save; verify in REDCap.","The profile and audit trail reflect the approved change.","Critical"),
    ("TC-008","Gap analysis","Save response, explanation and comment.","Domain assigned to signed-in group.","Enter all three values and save.","Values persist against the correct question and group.","High"),
    ("TC-009","Validation","Prevent an invalid or incomplete critical submission.","Required-field rules are configured.","Leave required data blank or enter invalid values; save.","Clear validation appears and invalid data is not synchronized.","High"),
    ("TC-010","Sync failure","Handle a REDCap connection or rejection failure safely.","Use test endpoint/token that will fail.","Edit a value and save.","Error is shown; prior server state is restored; retry guidance is clear.","Critical"),
    ("TC-011","Re-import","Confirm the effect of reloading REDCap data.","Local data and REDCap test record differ.","Reload the same assessment from REDCap.","The documented reload behaviour occurs without silent data loss.","Critical"),
    ("TC-012","Excel import","Validate and import a standard workbook.","Valid and invalid sample workbooks exist.","Import both files separately.","Valid data loads; invalid structure produces a clear error.","Medium"),
    ("TC-013","Reporting","Generate Word and Excel outputs from current data.","Completed test assessment exists.","Generate each export and compare key values.","Files open successfully and contain current, correct data.","High"),
    ("TC-014","Session/logout","Confirm logout handling for new and REDCap-loaded assessments.","One new and one imported workspace exist.","Log out from each workflow after saving.","Messages are accurate; no unintended duplicate synchronization occurs.","High"),
    ("TC-015","Auditability","Verify traceability of saved REDCap changes.","REDCap logging access is available.","Edit and save a test field; inspect REDCap log.","User, date/time and modification are traceable.","Critical"),
    ("TC-016","Test data","Prevent Test Country from entering REDCap.","Blank Test Country workspace exists.","Enter/import data; save; inspect status and REDCap; restart app.","Temporary-only status appears; REDCap has zero Test Country records; workspace is discarded after restart.","Critical"),
    ("TC-017","Excel import","Populate the complete Profiling workbook immediately.","Approved standard workbook contains Inventory and Profiling sheets.","Import once; inspect every system without refreshing the browser.","Names, platforms, years, diseases, coverage, all responses and comments display under the correct questions.","Critical"),
    ("TC-018","Profiling","Preserve unanswered dropdowns.","Workbook contains blank categorical responses.","Import; inspect blank coverage and other dropdowns; generate report.","Controls show No response, retain an empty value, and do not default to National or another first option.","Critical"),
    ("TC-019","Profiling report","Validate the brief summary and coverage denominator.","Profiles contain national, subnational and blank coverage responses.","Generate Profiling report; compare counts with source responses.","Five-part summary appears; only explicit national-level responses are counted.","High"),
    ("TC-020","Data quality","Prevent stale profiling tools and duplicates.","Workbook contains Inventory plus unmatched/stale profile sheet.","Import workbook and inspect tool/profile lists and report.","Inventory remains authoritative; unmatched profiles do not create tools; repeated findings are deduplicated.","High"),
    ("TC-021","Report assistant","Complete evidence-linked recommendations.","Profiling report with findings and recommendations is generated.","Ask to consolidate and prioritise recommendations with supporting findings.","Complete deduplicated response links each recommendation to recorded evidence and finishes without truncation.","High"),
    ("TC-022","Local AI","Avoid CPU first-token timeout for recommendation requests.","qwen3:4b-instruct runs on deployment computer.","Submit the evidence-linked recommendation request and time first response.","Compact evidence is used and response begins within the approved local performance threshold.","Medium"),
    ("TC-023","Authentication","Keep the administrator password valid across restarts.","Adjacent .env and active admin account exist.","Sign in; restart through terminal, shortcut and packaged launch; sign in again.","The same password succeeds for every supported launch and .env remains authoritative.","Critical"),
    ("TC-024","Profiling report","Keep response counts distinct from system counts.","Inventory has 14 systems; 11 have profiles with interoperability answers.","Generate the report and ask the assistant to discuss Domain capability scores.","The table labels answer counts; data exchange uses an explicit system denominator; no answer count is called a system count.","Critical"),
]
add_table(doc,test_headers,test_widths,tests,7.3)

doc.add_page_break(); page_title(doc,"3. TEST EXECUTION RECORD","Complete one row for every executed test case and build")
execution_headers=["Run ID","Test ID","Build / version","Environment","Tester","Date","Result","Evidence / actual result","Defect ID","Retest result"]
execution_widths=[700,650,950,850,900,750,700,2600,850,860]
execution_rows=[(f"RUN-{i:03d}","","","","","","Not run","","","") for i in range(1,13)]
add_table(doc,execution_headers,execution_widths,execution_rows,7.8)
heading(doc,"Evidence guidance",2)
for text in ["Reference screenshots, exported files, REDCap record IDs, log entries or other evidence without including passwords or API tokens.","For a failed test, record what occurred, where it occurred and how the actual result differed from the expected result.","Retest every corrected Critical or High defect and link the retest to both the original Run ID and Defect ID."]:
    p=doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after=Pt(5); font(p.add_run(text),9.5)

doc.add_page_break(); page_title(doc,"4. DEFECT AND RETEST REGISTER","Track failures from discovery through verified closure")
defect_headers=["Defect ID","Related test","Date","Summary","Severity","Steps / evidence","Owner","Target fix","Status","Retest evidence"]
defect_widths=[700,750,750,1900,750,2350,850,850,750,1160]
defect_rows=[(f"DEF-{i:03d}","","","","","","","","Open","") for i in range(1,9)]
add_table(doc,defect_headers,defect_widths,defect_rows,7.7)

doc.add_page_break(); page_title(doc,"5. TEST CYCLE SUMMARY AND RELEASE DECISION","Summarise coverage, residual risk and formal approval")
summary_rows=[
    ("Application version / build",""),("Test environment",""),("Test period",""),("Total tests planned",""),("Tests passed",""),("Tests failed",""),("Tests blocked / not run",""),("Open Critical defects",""),("Open High defects",""),("Overall test conclusion",""),
]
add_table(doc,["Summary field","Entry"],[2800,7760],summary_rows,9)
heading(doc,"Residual risks and release conditions",1)
for _ in range(5):
    p=doc.add_paragraph("____________________________________________________________________________________________________________________________"); p.paragraph_format.space_after=Pt(4); font(p.runs[0],9,False,GRAY)
heading(doc,"Approval",1)
approval_rows=[("Test lead","","",""),("Application owner","","",""),("REDCap/data manager","","",""),("Business or study owner","","","")]
add_table(doc,["Role","Name","Decision / signature","Date"],[2100,2400,4200,1860],approval_rows,9)
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(8); p.paragraph_format.space_after=Pt(0)
font(p.add_run("Release decision:  ☐ Approved     ☐ Approved with conditions     ☐ Not approved"),11,True,NAVY)

doc.core_properties.title="Application Refinement and Regression Testing Register"
doc.core_properties.subject="Africa CDC Surveillance Digital Tools Assessment Web Application"
doc.core_properties.author=""
doc.save(OUT)
print(OUT)

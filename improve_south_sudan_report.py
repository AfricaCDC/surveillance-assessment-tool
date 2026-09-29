from copy import deepcopy
from docx import Document
from docx.oxml.ns import qn

SOURCE = "Africa_CDC_South_Sudan_Assistant_Analysis_7_source.docx"
OUTPUT = "Africa_CDC_South_Sudan_Assistant_Analysis_Improved.docx"

doc = Document(SOURCE)
table = doc.tables[0]

def replace_paragraph_text(paragraph, new_text):
    if paragraph.runs:
        first = paragraph.runs[0]
        first.text = new_text
        for run in paragraph.runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(new_text)

def replace_cell_text(cell, new_text):
    first = cell.paragraphs[0]
    replace_paragraph_text(first, new_text)
    for paragraph in list(cell.paragraphs[1:]):
        paragraph._element.getparent().remove(paragraph._element)

# Correct the matrix content before removing the validation-gap column.
replace_cell_text(table.cell(0, 5), "Recommendation")

for row_index in range(1, len(table.rows)):
    tool = table.cell(row_index, 0).text.strip()
    covered = table.cell(row_index, 2).text.strip()
    missing = table.cell(row_index, 3).text.strip()

    if covered == "All 12 functions":
        replace_cell_text(
            table.cell(row_index, 2),
            "All 12 functions, including case investigation, contact follow-up, and case classification",
        )

    if covered.startswith("10 of 12 functions"):
        replace_cell_text(
            table.cell(row_index, 2),
            "10 of 12 functions, including case investigation, contact tracing, case classification, contact follow-up, outcome monitoring, case notification, case outcome, contact registration, laboratory data, and vaccination registry",
        )

    if missing == "Case investigation, contact follow-up, specimen tracking, case classification, contact assignment":
        replace_cell_text(table.cell(row_index, 3), "Specimen tracking; contact assignment")

    if tool == "HMIS/DHIS2":
        replace_cell_text(
            table.cell(row_index, 5),
            "Retain and enhance existing DHIS2 functionality through Unified DHIS2, subject to configuration and validation",
        )

# Remove the full Country validation gap column and its grid definition.
for row in table.rows:
    tc = row.cells[4]._tc
    tc.getparent().remove(tc)
grid = table._tbl.tblGrid
grid_cols = grid.findall(qn("w:gridCol"))
if len(grid_cols) > 4:
    grid.remove(grid_cols[4])

# Align narrative recommendations with the corrected Unified DHIS2 position.
paragraph_replacements = {
    "Finalize data migration plans for tools to be retired (HMIS/DHIS2, ODK).":
        "Finalize the ODK migration plan after validation, while retaining and enhancing existing HMIS/DHIS2 functionality through Unified DHIS2.",
    "Complete integration of EWARS, Eyer, and selected external tools (GoDATA, CoDA, NIS, OCV/ONA, eLIMS, PoE screening form, VLSM, Call center) into Unified DHIS2.":
        "Complete integration of EWARS, Eyer, and selected external tools (GoDATA, CoDA, NIS, OCV/ONA, eLIMS, PoE screening form, VLSM, Call center) with Unified DHIS2, using it to enhance and extend existing DHIS2 functionality.",
    "At least 80% of field data entry must be completed via DHIS2 with offline capability.":
        "At least 80% of field data entry must be completed through Unified DHIS2 with offline capability.",
}
for paragraph in doc.paragraphs:
    old = paragraph.text.strip()
    if old in paragraph_replacements:
        replace_paragraph_text(paragraph, paragraph_replacements[old])

# Remove the unsupported validation-gates section while preserving the
# following evidence-limitations section.
removing_validation_gates = False
for paragraph in list(doc.paragraphs):
    text = paragraph.text.strip()
    if text == "Validation gates and indicators":
        removing_validation_gates = True
    elif removing_validation_gates and text == "Evidence limitations":
        break

    if removing_validation_gates:
        paragraph._element.getparent().remove(paragraph._element)

doc.core_properties.title = "Improved South Sudan Unified DHIS2 Analysis"
doc.save(OUTPUT)
print(OUTPUT)

import io, json
from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

GREEN, DARK, GOLD, RED, WHITE, INK, ALT = "176B45", "0D4F32", "D5A83F", "9B2335", "FFFFFF", "20362B", "F3F8F5"

def display(value):
    if isinstance(value, list): return "; ".join(map(str, value))
    if value is None: return ""
    text = str(value)
    try:
        parsed = json.loads(text)
        return "; ".join(map(str, parsed)) if isinstance(parsed, list) else text
    except (TypeError, ValueError): return text

def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr(); node = props.find(qn("w:shd"))
    if node is None: node = OxmlElement("w:shd"); props.append(node)
    node.set(qn("w:fill"), fill)

def width(cell, inches):
    props = cell._tc.get_or_add_tcPr(); node = props.find(qn("w:tcW"))
    if node is None: node = OxmlElement("w:tcW"); props.append(node)
    node.set(qn("w:w"), str(round(inches * 1440))); node.set(qn("w:type"), "dxa")

class Builder:
    def __init__(self, profile, state, domains, summary, logo):
        self.profile, self.state, self.domains, self.summary = profile, state, domains, summary
        self.logo, self.doc = Path(logo), Document(); self.setup()

    def setup(self):
        sec = self.doc.sections[0]
        sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(.75)
        sec.header_distance = sec.footer_distance = Inches(.35)
        normal = self.doc.styles["Normal"]; normal.font.name = "Arial"; normal.font.size = Pt(10); normal.font.color.rgb = RGBColor.from_string(INK)
        normal.paragraph_format.space_after = Pt(5); normal.paragraph_format.line_spacing = 1.1
        for name, size, color, before, after in [("Title",26,DARK,0,10),("Heading 1",18,GREEN,14,8),("Heading 2",14,DARK,11,6),("Heading 3",11,RED,8,4)]:
            style = self.doc.styles[name]; style.font.name = "Arial"; style.font.size = Pt(size); style.font.bold = True; style.font.color.rgb = RGBColor.from_string(color)
            style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after); style.paragraph_format.keep_with_next = True
        header = sec.header.paragraphs[0]; header.text = "AFRICA CDC  |  SURVEILLANCE SYSTEMS ASSESSMENT"
        header.runs[0].font.name = "Arial"; header.runs[0].font.size = Pt(8); header.runs[0].font.bold = True; header.runs[0].font.color.rgb = RGBColor.from_string(GREEN)
        footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = footer.add_run("Africa CDC Assessment  |  "); run.font.name = "Arial"; run.font.size = Pt(8); run.font.color.rgb = RGBColor.from_string(GREEN)
        field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)

    def table(self, headers, rows, widths, size=8.5):
        rows = [row for row in rows if any(display(value).strip() and display(value).strip() != "-" for value in row)]
        if not rows:
            return
        table = self.doc.add_table(rows=1, cols=len(headers)); table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
        for i, (cell, label) in enumerate(zip(table.rows[0].cells, headers)):
            cell.text = str(label); shade(cell, GREEN); width(cell, widths[i]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cell.paragraphs[0].runs: run.font.name = "Arial"; run.font.size = Pt(size); run.font.bold = True; run.font.color.rgb = RGBColor.from_string(WHITE)
        props = table.rows[0]._tr.get_or_add_trPr(); repeat = OxmlElement("w:tblHeader"); repeat.set(qn("w:val"), "true"); props.append(repeat)
        for row_no, values in enumerate(rows):
            cells = table.add_row().cells
            for i, cell in enumerate(cells):
                cell.text = str(values[i] if i < len(values) and values[i] not in (None, "") else "-"); width(cell, widths[i]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if row_no % 2: shade(cell, ALT)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2)
                    for run in p.runs: run.font.name = "Arial"; run.font.size = Pt(size)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(1)

    def cover(self):
        if self.logo.exists():
            p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p.add_run().add_picture(str(self.logo), width=Inches(1.75))
        p = self.doc.add_paragraph("AFRICA CENTRES FOR DISEASE CONTROL AND PREVENTION"); r = p.runs[0]; r.font.name = "Arial"; r.font.size = Pt(10); r.font.bold = True; r.font.color.rgb = RGBColor.from_string(RED)
        title = self.doc.add_paragraph("Surveillance Systems Profiling\nand Gap Analysis", style="Title"); title.paragraph_format.space_before = Pt(30)
        p = self.doc.add_paragraph("Complete Assessment Report"); r = p.runs[0]; r.font.name = "Arial"; r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = RGBColor.from_string(GOLD)
        phase1 = self.profile.get("phase1", {})
        self.table(["Report information","Value"], [("Country",phase1.get("country_name") or "Not provided"),("Reporting period",phase1.get("reporting_period") or "Not provided"),("Assessment",self.summary["assessment"].get("title") or "Africa CDC Gap Analysis")], [1.6,4.9], 10)
        self.doc.add_page_break()

    def add_phase1(self):
        phase1 = self.profile.get("phase1", {}); tools = [x for x in phase1.get("tools",[]) if str(x.get("inventory_name","")).strip()]; phase2 = [x for x in self.profile.get("phase2",[]) if str(x.get("official_name","")).strip()]
        self.doc.add_heading("Executive Summary",1); self.doc.add_paragraph("This report synthesises the available Inventory, Profiling and Gap Analysis evidence for the selected country assessment.")
        self.table(["Assessment component","Coverage"], [("Inventory",f"{len(tools)} systems"),("Profiling",f"{len(phase2)} system profiles"),("Gap Analysis",f"{len(self.state['responses'])} responses; {len(self.state['assignments'])} domains assigned")], [3.25,3.25], 10)
        analysis = self.summary.get("inventory_analysis") or {}
        findings = analysis.get("findings") or []
        recommendations = analysis.get("recommendations") or []
        if findings:
            p = self.doc.add_paragraph("Key findings from the Inventory:"); p.runs[0].font.bold = True
            for item in findings: self.doc.add_paragraph(str(item), style="List Bullet")
        self.doc.add_heading("Inventory",1)
        rows = [[display(x.get(k)) for k in ("inventory_name","unit_responsible","system_type","surveillance_functions","geographical_coverage","has_api")] for x in tools]
        self.table(["System","Responsible unit","Type","Functions","Coverage","API"], rows, [1,1.05,1,1.8,1,.65], 7.5)
        if recommendations:
            self.doc.add_heading("Recommendations",2)
            for item in recommendations: self.doc.add_paragraph(str(item), style="List Bullet")

    def add_phase2(self):
        items = [x for x in self.profile.get("phase2",[]) if str(x.get("official_name","")).strip()]
        self.doc.add_heading("Profiling",1)
        rows = [[x.get("official_name"),x.get("implementation_status"),display(x.get("platforms")),display(x.get("ownership_institution")),x.get("api_available"),x.get("funding_source")] for x in items]
        self.table(["System","Status","Platform","Owner","API","Funding"], rows, [1.25,.8,1.25,1.25,.65,1.3], 8)
        for item in items:
            self.doc.add_heading(str(item.get("official_name") or "System assessment"),2)
            rows = [[k.replace("_"," ").title(),display(v)] for k,v in item.items() if k != "official_name" and display(v).strip()]
            self.table(["Assessment item","Response"], rows, [2.25,4.25], 8.5)

    def add_phase3(self):
        self.doc.add_heading("Gap Analysis",1)
        assignments = {x["domain_id"]:x for x in self.state["assignments"]}; responses = {x["question_id"]:x for x in self.state["responses"]}
        for domain in self.domains:
            rows=[]
            for q in domain["questions"]:
                answer = responses.get(q["id"],{}); value = display(answer.get("response",""))
                explanation = display(answer.get("explanation", "")).strip(); comment = display(answer.get("comment", "")).strip()
                if not any((value.strip(), explanation, comment)):
                    continue
                if explanation: value += ("\n\nExplanation: " if value else "Explanation: ") + explanation
                rows.append([q.get("source_id",q["id"]).upper(),q["text"],value,comment])
            if not rows:
                continue
            owner = assignments.get(domain["id"],{}).get("group_name","Unassigned"); self.doc.add_heading(domain["name"],2)
            p = self.doc.add_paragraph(f"Participant group: {owner}"); p.runs[0].font.bold = True; p.runs[0].font.color.rgb = RGBColor.from_string(GREEN)
            self.table(["ID","Question","Response / explanation","Comment / gap / recommendation"], rows, [.55,2.65,2,1.3], 7.5)
        self.doc.add_heading("Gap Analysis: Gaps and Recommendations Summary",1)
        for item in self.summary["sections"]:
            gaps=item["gaps"]+item["other"]
            if not gaps and not item["recommendations"]:
                continue
            self.doc.add_heading(item["domain"],2)
            self.table(["Gaps and other findings","Recommendations"], [["\n".join(gaps),"\n".join(item["recommendations"])]], [3.25,3.25], 9)

    def build(self, phases=None):
        phases = tuple(phases or (1, 2, 3)); self.cover()
        for index, phase in enumerate(phases):
            if index: self.doc.add_page_break()
            {1:self.add_phase1, 2:self.add_phase2, 3:self.add_phase3}[phase]()
        props=self.doc.core_properties; props.title="Africa CDC Surveillance Systems Assessment Report"; props.subject="Inventory, Profiling and Gap Analysis results"; props.author="Africa CDC"
        output=io.BytesIO(); self.doc.save(output); return output.getvalue()

def build_report(profile,state,domains,summary,logo,phases=None): return Builder(profile,state,domains,summary,logo).build(phases)

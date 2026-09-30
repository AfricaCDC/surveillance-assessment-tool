import argparse
import base64
import hashlib
import hmac
import io
import json
import mimetypes
import os
import re
import secrets
import shutil
import socket
import sqlite3
import subprocess
import sys
import threading
import time
import traceback
import webbrowser
import zipfile
from xml.etree import ElementTree
from recommendation_engine import draft_report, format_audit, is_recommendation_request
from datetime import datetime, timedelta, timezone
from http.cookies import SimpleCookie
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlencode, urlparse
from urllib.request import Request, urlopen
from urllib.error import HTTPError, URLError

try:
    from argon2 import PasswordHasher
    from argon2.exceptions import InvalidHashError, VerificationError, VerifyMismatchError
except ImportError:
    PasswordHasher = None
    InvalidHashError = VerificationError = VerifyMismatchError = ValueError
try:
    from dotenv import load_dotenv
except ImportError:
    load_dotenv = None

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None
try:
    from docx import Document
except ImportError:
    Document = None
try:
    from surveillance_app import generate_report as generate_desktop_report
except ImportError:
    generate_desktop_report = None


FROZEN = bool(getattr(sys, "frozen", False) and hasattr(sys, "_MEIPASS"))
APP_DIR = Path(sys.executable).resolve().parent if FROZEN else Path(__file__).resolve().parent
BUNDLE_DIR = Path(sys._MEIPASS).resolve() if FROZEN else APP_DIR


def load_app_environment():
    env_path = APP_DIR / ".env"
    if load_dotenv:
        # The application and its adjacent .env form one configuration unit.
        # Always prefer that file so launching from a terminal, shortcut, or
        # packaged executable cannot silently change the password pepper.
        load_dotenv(env_path, override=True)
        return
    if not env_path.is_file():
        return
    for raw_line in env_path.read_text(encoding="utf-8-sig").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip()
        if key.startswith("export "):
            key = key[7:].strip()
        if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", key):
            continue
        value = value.strip()
        if len(value) >= 2 and value[0] == value[-1] and value[0] in {'"', "'"}:
            value = value[1:-1]
        os.environ[key] = value


load_app_environment()
STATIC_DIR = BUNDLE_DIR / "static"
DB_PATH = Path(os.environ.get("AFRICA_CDC_DB_PATH", APP_DIR / "africa_cdc_web.db"))
STANDARD_TEMPLATE_PATH = BUNDLE_DIR / "standard_assessment_template.xlsx"
DEFAULT_REDCAP_API_URL = os.environ.get("AFRICA_CDC_REDCAP_API_URL", "https://tools.africacdc.org/africacdcrc/api/").strip()
APP_VERSION = "1.0.41"
PROFILE_FORMAT_VERSION = 3
OLLAMA_URL = os.environ.get("AFRICA_CDC_OLLAMA_URL", "http://127.0.0.1:11434").strip().rstrip("/")
OLLAMA_MODEL = os.environ.get("AFRICA_CDC_OLLAMA_MODEL", "qwen3:4b-instruct").strip()
OLLAMA_START_LOCK = threading.Lock()
OPENAI_API_KEY = os.environ.get("AFRICA_CDC_OPENAI_API_KEY", "").strip()
OPENAI_MODEL = os.environ.get("AFRICA_CDC_OPENAI_MODEL", "gpt-5-mini").strip()
OPENAI_RESPONSES_URL = "https://api.openai.com/v1/responses"

STANDARD_GAP_SHEETS = {
    "domain 1,8,9": ({"d1", "d8", "d9"}, 0),
    "domain 2,4,7": ({"d2", "d4", "d7"}, 1),
    "domain 3,5,6": ({"d3", "d5", "d6"}, 2),
}

DOMAINS = [
    {"id": "d1", "name": "Governance and Coordination", "questions": [
        {"id": "g1", "text": "Who are the main institutions and actors involved in surveillance and emergency response in the country?", "type": "multi", "options": ["Ministry of Health", "IES&PHE", "PHEOC", "National and reference laboratories", "District/county local governments", "Health facilities", "Points of entry", "One Health sectors", "Research/academic institutions", "Implementing partners", "Other (specify)"]},
        {"id": "g2", "text": "Describe the mandates and responsibilities of the different institutions for collecting, managing, reporting and using surveillance and outbreak data.", "type": "text"},
        {"id": "g3", "text": "What coordination mechanisms or structures are in place for surveillance and emergency response, and how regularly do they meet and follow up agreed actions?", "type": "text"},
        {"id": "g4", "text": "How are decisions made and approved for introducing or replacing surveillance systems, changing reporting tools, and upgrading software platforms and systems?", "type": "text"},
        {"id": "g5", "text": "What are the main governance and coordination gaps, and which actions should be prioritised to address them?", "type": "text"},
    ]},
    {"id": "d2", "name": "Architecture and Design", "questions": [
        {"id": "a1", "text": "Describe the overall architecture of the country's routine surveillance and emergency response information systems.", "type": "text"},
        {"id": "a2", "text": "How do the different surveillance systems relate to each other and to other national systems such as laboratory, EMR, immunisation, CRVS and points-of-entry systems?", "type": "text"},
        {"id": "a3", "text": "Is there a documented long-term national health information or digital health architecture, and does it adequately address surveillance and emergency response?", "type": "select_explain", "options": ["Yes, approved and implemented", "Available but only partly addresses surveillance", "Draft/under development", "No", "Unknown"]},
        {"id": "a4", "text": "How are case notification, investigation and follow-up managed electronically, including the relationship between case-based systems such as DHIS2 Tracker and aggregate IDSR reporting?", "type": "text"},
        {"id": "a5", "text": "How are laboratory results linked to cases, and what identifiers are used to link persons, cases, samples, contacts and events?", "type": "text"},
    ]},
    {"id": "d3", "name": "Data Collection Processes and Tools", "questions": [
        {"id": "dc1", "text": "Describe the main surveillance data collection and reporting processes from community and Village Health Team level through health facilities, laboratories, districts and national level.", "type": "text"},
        {"id": "dc2", "text": "Which groups of users are involved in surveillance data collection, and what reporting tools, devices and communication channels do they use?", "type": "text"},
        {"id": "dc3", "text": "How complete and timely is case notification and routine surveillance reporting at the different levels of the health system?", "type": "text"},
        {"id": "dc4", "text": "How are mobile and offline data collection supported, including device provision, use of personal devices, connectivity, application updates and hardware support?", "type": "text"},
        {"id": "dc5", "text": "What are the main gaps in surveillance data collection processes and tools, and which improvements should be prioritised?", "type": "text"},
    ]},
    {"id": "d4", "name": "Data Standards and Quality", "questions": [
        {"id": "ds1", "text": "Which diseases, conditions and public health events are under surveillance, and which are included in electronic surveillance systems?", "type": "text"},
        {"id": "ds2", "text": "How well do the variables and indicators in national surveillance tools align with the Gap Analysis Toolkit data dictionary and current national, WHO and Africa CDC guidance?", "type": "text"},
        {"id": "ds3", "text": "Are clear and current case definitions, classification rules and reporting guidance available and accessible for all priority diseases and events?", "type": "select_explain", "options": ["Available and current for all", "Available for most", "Available for some", "Not available", "Unknown"]},
        {"id": "ds4", "text": "Are outbreak alert and action thresholds defined for priority diseases and events, and are they configured and used in the surveillance systems?", "type": "text"},
        {"id": "ds5", "text": "What are the main data standards and data quality gaps, and which actions should be prioritised to address them?", "type": "text"},
    ]},
    {"id": "d5", "name": "Surveillance and Response Functionality", "questions": [
        {"id": "f1", "text": "How are rumours, signals and alerts collected, verified, triaged and linked to case or event investigation?", "type": "text"},
        {"id": "f2", "text": "How is contact tracing currently supported, including linkage of contacts to index cases and monitoring of follow-up and outcomes?", "type": "text"},
        {"id": "f3", "text": "How are vaccination activities for routine services and outbreak response recorded and linked to cases, contacts or target populations?", "type": "text"},
        {"id": "f4", "text": "How do the systems support automated notifications, escalation of key events, and Rapid Response Team rostering and deployment?", "type": "text"},
        {"id": "f5", "text": "Which surveillance and response functions are unavailable or not working adequately, and which should be prioritised for strengthening or rapid activation?", "type": "text"},
    ]},
    {"id": "d6", "name": "Analytics, Dissemination, and Data Use", "questions": [
        {"id": "an1", "text": "What surveillance indicators, dashboards and analytical outputs are currently available, and how are they defined and maintained?", "type": "text"},
        {"id": "an2", "text": "Are dashboards and analytical products tailored to the needs of different users such as the Ministry of Health, PHEOC, laboratories, districts, facilities and partners?", "type": "select_explain", "options": ["Yes, for all key user groups", "For most groups", "For a few groups", "No", "Unknown"]},
        {"id": "an3", "text": "How are surveillance data analysed and used for early warning, trend monitoring, outbreak threshold detection and management of response activities?", "type": "text"},
        {"id": "an4", "text": "Describe the current practices for disseminating surveillance information and producing routine information products.", "type": "text"},
        {"id": "an5", "text": "What are the main gaps in analytics, dissemination and data use, and which improvements should be prioritised?", "type": "text"},
    ]},
    {"id": "d7", "name": "Security, Privacy, and Resilience", "questions": [
        {"id": "sp1", "text": "Is there a clearly designated senior person or institution responsible for information security and privacy for surveillance systems?", "type": "select_explain", "options": ["Yes, clearly designated", "Partly defined", "No", "Unknown"]},
        {"id": "sp2", "text": "What policies and procedures are in place for data protection, data sharing, user access, confidentiality, retention and acceptable use of surveillance data?", "type": "text"},
        {"id": "sp3", "text": "Is there a current security and privacy risk register, and are risk or privacy assessments conducted and followed up?", "type": "select_explain", "options": ["Yes, current and routinely reviewed", "Available but not current", "Assessments are ad hoc", "No", "Unknown"]},
        {"id": "sp4", "text": "Are system backups automated and routinely tested, and are recovery and business-continuity procedures documented and tested?", "type": "text"},
        {"id": "sp5", "text": "What are the main security, privacy and system resilience gaps, and which actions should be prioritised?", "type": "text"},
    ]},
    {"id": "d8", "name": "Capacity for Sustained System Management", "questions": [
        {"id": "cm1", "text": "Which teams are responsible for managing the country's surveillance information systems, and what roles do they perform?", "type": "text"},
        {"id": "cm2", "text": "Are all the required roles and competencies available within the system management teams?", "type": "text"},
        {"id": "cm3", "text": "Do system administrators and support teams receive appropriate induction, training, refresher training and mentoring to perform their roles?", "type": "select_explain", "options": ["Consistently", "Sometimes", "Rarely", "Never", "Unknown"]},
        {"id": "cm4", "text": "Do the responsible teams have sufficient time, staffing, tools, infrastructure and operational resources to manage and support the systems?", "type": "select_explain", "options": ["Sufficient", "Mostly sufficient", "Partly sufficient", "Insufficient", "Unknown"]},
        {"id": "cm5", "text": "Which standard operating procedures and support arrangements are available for user management, configuration changes, testing, upgrades, documentation and helpdesk support, and what gaps remain?", "type": "text"},
    ]},
    {"id": "d9", "name": "Sustainable Financing", "questions": [
        {"id": "fin1", "source_id": "sf1", "text": "Describe the current funding model for surveillance information systems and emergency response activities in the country.", "type": "text"},
        {"id": "fin2", "source_id": "sf2", "text": "Which routine and emergency cost categories are currently funded, and which institutions or partners cover them?", "type": "text"},
        {"id": "fin3", "source_id": "sf3", "text": "For how long is current funding secured for routine surveillance system operations and priority strengthening activities?", "type": "text"},
        {"id": "fin4", "source_id": "sf4", "text": "Are there ongoing funded surveillance or outbreak-response projects, and how are they aligned with national priorities, architecture and transition to government ownership?", "type": "text"},
        {"id": "fin5", "source_id": "sf5", "text": "What are the main financing gaps and sustainability risks, and which actions should be prioritised?", "type": "text"},
    ]},
]

QUESTION_INDEX = {}
for domain in DOMAINS:
    for question in domain["questions"]:
        QUESTION_INDEX[question["id"]] = (domain["id"], question)
        QUESTION_INDEX[question.get("source_id", question["id"])] = (domain["id"], question)


def utc_now():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


_RUNTIME_CONNECTION = None
_RUNTIME_LOCK = threading.RLock()
# REDCap is the sole persistent store for Member State assessment data. Only
# authentication and application configuration may be written to local disk.
_PERSISTENT_TABLES = ("users", "sessions", "app_settings")


def _disk_connect():
    connection = sqlite3.connect(DB_PATH, timeout=20)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA journal_mode=WAL")
    connection.execute("PRAGMA foreign_keys=ON")
    return connection


def _persist_runtime_tables():
    if _RUNTIME_CONNECTION is None:
        return
    disk = _disk_connect()
    try:
        disk.execute("PRAGMA foreign_keys=OFF")
        for table in _PERSISTENT_TABLES:
            columns = [row[1] for row in _RUNTIME_CONNECTION.execute(f"PRAGMA table_info({table})")]
            if not columns:
                continue
            rows = [list(row) for row in _RUNTIME_CONNECTION.execute(f"SELECT {','.join(columns)} FROM {table}")]
            if table == "country_projects" and "payload" in columns:
                payload_index = columns.index("payload")
                for row in rows:
                    row[payload_index] = None
            disk.execute(f"DELETE FROM {table}")
            if rows:
                placeholders = ",".join("?" for _ in columns)
                disk.executemany(f"INSERT INTO {table}({','.join(columns)}) VALUES({placeholders})", rows)
        disk.commit()
    finally:
        disk.close()


class _RuntimeContext:
    def __enter__(self):
        _RUNTIME_LOCK.acquire()
        self.changes = _RUNTIME_CONNECTION.total_changes
        return _RUNTIME_CONNECTION

    def __exit__(self, exc_type, exc, traceback):
        try:
            if exc_type is None:
                _RUNTIME_CONNECTION.commit()
                if _RUNTIME_CONNECTION.total_changes != self.changes:
                    _persist_runtime_tables()
            else:
                _RUNTIME_CONNECTION.rollback()
        finally:
            _RUNTIME_LOCK.release()
        return False


def connect():
    if _RUNTIME_CONNECTION is None:
        return _disk_connect()
    return _RuntimeContext()


def initialize_runtime_workspace():
    global _RUNTIME_CONNECTION
    disk = _disk_connect()
    runtime = sqlite3.connect(":memory:", check_same_thread=False)
    runtime.row_factory = sqlite3.Row
    disk.backup(runtime)
    runtime.execute("PRAGMA foreign_keys=ON")
    assessment_id = runtime.execute("SELECT id FROM assessments ORDER BY id LIMIT 1").fetchone()[0]
    for table in ("responses", "assignments", "inventory_tool_assignments", "profiling_tool_assignments", "groups", "profile_data", "country_projects"):
        runtime.execute(f"DELETE FROM {table}")
    runtime.execute("UPDATE assessments SET title='Africa CDC Gap Analysis',scope='' WHERE id=?", (assessment_id,))
    for name in ("Group 1", "Group 2", "Group 3"):
        runtime.execute("INSERT INTO groups(assessment_id,name) VALUES(?,?)", (assessment_id, name))
    now = utc_now()
    project_id = runtime.execute(
        "INSERT INTO country_projects(country_name,reporting_period,title,payload,redcap_record_id,created_at,updated_at) VALUES('','','Country Assessment',NULL,NULL,?,?)",
        (now, now),
    ).lastrowid
    runtime.execute("INSERT INTO app_settings(key,value) VALUES('active_project_id',?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (str(project_id),))
    runtime.commit()
    disk_assessment_id = disk.execute("SELECT id FROM assessments ORDER BY id LIMIT 1").fetchone()[0]
    disk.execute("PRAGMA foreign_keys=OFF")
    for table in ("responses", "assignments", "inventory_tool_assignments", "profiling_tool_assignments", "groups", "profile_data", "country_projects"):
        disk.execute(f"DELETE FROM {table}")
    disk.execute("UPDATE assessments SET title='Africa CDC Gap Analysis',scope='' WHERE id=?", (disk_assessment_id,))
    for name in ("Group 1", "Group 2", "Group 3"):
        disk.execute("INSERT INTO groups(assessment_id,name) VALUES(?,?)", (disk_assessment_id, name))
    disk_project_id = disk.execute(
        "INSERT INTO country_projects(country_name,reporting_period,title,payload,redcap_record_id,created_at,updated_at) VALUES('','','Country Assessment',NULL,NULL,?,?)",
        (now, now),
    ).lastrowid
    disk.execute("INSERT INTO app_settings(key,value) VALUES('active_project_id',?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (str(disk_project_id),))
    disk.commit()
    disk.close()
    _RUNTIME_CONNECTION = runtime


REDCAP_TOKEN = os.environ.get("AFRICA_CDC_REDCAP_TOKEN", "")
REDCAP_TOKEN_LOCK = threading.Lock()


def configured_redcap_api_url():
    return os.environ.get("AFRICA_CDC_REDCAP_API_URL", "").strip() or app_setting("redcap_api_url", DEFAULT_REDCAP_API_URL)


def app_setting(key, default=""):
    with connect() as db:
        row = db.execute("SELECT value FROM app_settings WHERE key=?", (key,)).fetchone()
    return row[0] if row else default


def save_app_setting(key, value):
    with connect() as db:
        db.execute("INSERT INTO app_settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (key, value))


def redcap_status():
    url = configured_redcap_api_url()
    return {
        "configured": bool(url and REDCAP_TOKEN),
        "api_url": url,
        "token_saved": bool(REDCAP_TOKEN),
        "last_tested_at": app_setting("redcap_last_tested_at"),
        "project_title": app_setting("redcap_project_title"),
        "project_id": app_setting("redcap_project_id"),
        "last_sync_at": app_setting("redcap_last_sync_at"),
        "last_record_id": app_setting("redcap_last_record_id"),
        "last_pull_at": app_setting("redcap_last_pull_at"),
    }


def redcap_request(parameters):
    api_url = configured_redcap_api_url().strip()
    with REDCAP_TOKEN_LOCK:
        token = REDCAP_TOKEN
    if not api_url or not token:
        raise ValueError("Save the REDCap API URL and token first")
    normalized_url = api_url.rstrip("/").lower()
    if not api_url.lower().startswith("https://") or not (normalized_url.endswith("/api") or normalized_url.endswith("/api/index.php")):
        raise ValueError("REDCap API URL must use HTTPS and end with /api/ or /api/index.php")
    body = urlencode({"token": token, "format": "json", "returnFormat": "json", **parameters}).encode("utf-8")
    request = Request(api_url, data=body, headers={"Content-Type": "application/x-www-form-urlencoded"}, method="POST")
    for attempt in range(3):
        try:
            with urlopen(request, timeout=25) as response:
                raw_response = response.read().decode("utf-8", "replace").strip()
                if not raw_response:
                    raise RuntimeError("REDCap returned an empty response. Ask the REDCap administrator to check API access and IP allow-listing.")
                try:
                    result = json.loads(raw_response)
                except json.JSONDecodeError as exc:
                    clean = re.sub(r"<[^>]+>", " ", raw_response)
                    clean = " ".join(clean.split())[:500]
                    raise RuntimeError(f"REDCap returned a non-JSON response: {clean}") from exc
            break
        except HTTPError as exc:
            detail = exc.read().decode("utf-8", "replace")
            try: detail = json.loads(detail).get("error", detail)
            except Exception: pass
            raise RuntimeError(f"REDCap rejected the connection: {detail}") from exc
        except URLError as exc:
            if attempt < 2:
                time.sleep(attempt + 1)
                continue
            raise RuntimeError(
                "Could not connect to REDCap after three attempts. Check the internet connection, "
                "VPN or proxy, and Windows DNS, then try again."
            ) from exc
    if isinstance(result, dict) and result.get("error"):
        raise RuntimeError(result["error"])
    return result


def test_redcap_connection():
    project = redcap_request({"content": "project"})
    title = str(project.get("project_title") or project.get("projectTitle") or "REDCap project")
    project_id = str(project.get("project_id") or project.get("projectId") or "")
    tested_at = utc_now()
    save_app_setting("redcap_project_title", title)
    save_app_setting("redcap_project_id", project_id)
    save_app_setting("redcap_last_tested_at", tested_at)
    return {**redcap_status(), "ok": True, "message": f"Connected to {title}"}


def _redcap_value(value):
    if value is None:
        return ""
    if isinstance(value, list):
        return "; ".join(str(item) for item in value if item not in (None, ""))
    if isinstance(value, (dict, tuple)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def redcap_records_for_current_project():
    """Flatten the active assessment into REDCap's repeating-instrument format."""
    capture_active_project()
    project_id = active_project_id()
    profile = profile_state()
    state = current_state()
    phase1 = profile.get("phase1") or {}
    phase2 = profile.get("phase2") or []
    country = str(phase1.get("country_name") or "").strip()
    period = str(phase1.get("reporting_period") or "").strip()
    with connect() as db:
        project = db.execute("SELECT redcap_record_id FROM country_projects WHERE id=?", (project_id,)).fetchone()
    record_id = str(project[0] or "").strip() if project else ""
    if not record_id:
        record_id = f"acdc-{project_id}"
    now = utc_now()
    records = [{
        "record_id": record_id,
        "country_name": country,
        "reporting_period": period,
        "assessment_title": state["assessment"].get("title") or f"{country or 'Country'} Assessment",
        "assessment_scope": state["assessment"].get("scope") or "",
        "assessment_status": "active",
        "updated_at": now.replace("T", " ").split("+")[0],
        "country_assessment_complete": "2",
    }]
    groups = {item["id"]: item for item in state.get("groups", [])}
    with connect() as db:
        accounts = [dict(row) for row in db.execute("SELECT username,active,created_at,phase,group_id FROM users WHERE role='clerk' ORDER BY phase,group_id")]
    account_by_phase_group = {(item.get("phase"), item.get("group_id")): item for item in accounts}
    phase_groups = set(account_by_phase_group)
    phase_groups.update(("inventory", item.get("group_id")) for item in state.get("inventory_assignments", []) if item.get("group_id"))
    phase_groups.update(("profiling", item.get("group_id")) for item in state.get("profiling_assignments", []) if item.get("group_id"))
    phase_groups.update(("gap", item.get("group_id")) for item in state.get("assignments", []) if item.get("group_id"))
    for instance, (phase, group_id) in enumerate(sorted(phase_groups, key=lambda item: (item[0], int(item[1]))), 1):
        account = account_by_phase_group.get((phase, group_id), {})
        group = groups.get(group_id, {})
        records.append({
            "record_id": record_id, "redcap_repeat_instrument": "working_groups", "redcap_repeat_instance": instance,
            "wg_group_id": group_id, "wg_phase": phase,
            "wg_group_name": group.get("name", ""), "wg_username": account.get("username", ""),
            "wg_active": "1" if account.get("active", True) else "0",
            "wg_created_at": str(account.get("created_at") or "").replace("T", " ").split("+")[0],
            "working_groups_complete": "2",
        })
    inventory_owner = {item["tool_id"]: item for item in state.get("inventory_assignments", [])}
    for instance, tool in enumerate(phase1.get("tools", []), 1):
        owner = inventory_owner.get(tool.get("_id"), {})
        records.append({
            "record_id": record_id, "redcap_repeat_instrument": "inventory_tools", "redcap_repeat_instance": instance,
            "inv_tool_id": tool.get("_id", ""), "inv_tool_name": tool.get("inventory_name", ""),
            "inv_unit": tool.get("unit_responsible", ""), "inv_system_type": _redcap_value(tool.get("system_type")),
            "inv_surv_functions": _redcap_value(tool.get("surveillance_functions")), "inv_geo_coverage": tool.get("geographical_coverage", ""),
            "inv_data_entry": tool.get("point_of_data_entry", ""), "inv_data_captured": tool.get("data_captured", ""),
            "inv_info_users": _redcap_value(tool.get("information_users")), "inv_technology": tool.get("technology", ""),
            "inv_has_api": "1" if str(tool.get("has_api", "")).strip().lower() == "yes" else "0" if str(tool.get("has_api", "")).strip().lower() == "no" else "",
            "inv_linked_systems": tool.get("linked_to_other_systems", ""), "inv_comments": tool.get("comments", ""),
            "inv_entered_by": owner.get("group_name", ""), "inv_updated_at": now.replace("T", " ").split("+")[0],
            "inventory_tools_complete": "2",
        })
    for instance, item in enumerate(phase2, 1):
        row = {"record_id": record_id, "redcap_repeat_instrument": "system_profiling", "redcap_repeat_instance": instance,
               "prof_tool_id": item.get("_tool_id", ""), "prof_updated_at": now.replace("T", " ").split("+")[0],
               "system_profiling_complete": "2"}
        for key in PROFILE_ID_MAP.values():
            row[key] = _redcap_value(item.get(key))
        records.append(row)
    work = []
    for phase, key, item_key in (("inventory", "inventory_assignments", "tool_id"), ("profiling", "profiling_assignments", "tool_id"), ("gap", "assignments", "domain_id")):
        for item in state.get(key, []):
            item_id = item.get(item_key, "")
            if phase == "gap":
                item_name = next((domain["name"] for domain in DOMAINS if domain["id"] == item_id), item_id)
            else:
                item_name = next((tool.get("inventory_name", item_id) for tool in phase1.get("tools", []) if tool.get("_id") == item_id), item_id)
            work.append((phase, item, item_id, item_name))
    for instance, (phase, item, item_id, item_name) in enumerate(work, 1):
        records.append({
            "record_id": record_id, "redcap_repeat_instrument": "work_assignments", "redcap_repeat_instance": instance,
            "assign_phase": phase, "assign_item_id": item_id, "assign_item_name": item_name,
            "assign_group_id": item.get("group_id", ""), "assign_group_name": item.get("group_name", ""),
            "assigned_at": str(item.get("assigned_at") or "").replace("T", " ").split("+")[0],
            "work_assignments_complete": "2",
        })
    domain_names = {domain["id"]: domain["name"] for domain in DOMAINS}
    question_text = {question["id"]: question["text"] for domain in DOMAINS for question in domain["questions"]}
    for instance, item in enumerate(state.get("responses", []), 1):
        records.append({
            "record_id": record_id, "redcap_repeat_instrument": "gap_analysis_responses", "redcap_repeat_instance": instance,
            "gap_domain_id": item.get("domain_id", ""), "gap_domain_name": domain_names.get(item.get("domain_id"), ""),
            "gap_question_id": item.get("question_id", ""), "gap_question_text": question_text.get(item.get("question_id"), ""),
            "gap_response": item.get("response", ""), "gap_explanation": item.get("explanation", ""),
            "gap_comment": item.get("comment", ""), "gap_group_id": item.get("group_id", ""),
            "gap_group_name": groups.get(item.get("group_id"), {}).get("name", ""),
            "gap_updated_at": str(item.get("updated_at") or now).replace("T", " ").split("+")[0],
            "gap_analysis_responses_complete": "2",
        })
    return record_id, records


def is_test_country_name(value):
    return re.sub(r"\s+", " ", str(value or "")).strip().casefold() == "test country"


def sync_current_project_to_redcap():
    country = str((profile_state().get("phase1") or {}).get("country_name") or "").strip()
    if is_test_country_name(country):
        return {
            "ok": True,
            "temporary": True,
            "record_id": "",
            "rows_sent": 0,
            "synced_at": utc_now(),
            "message": "Test Country data is temporary and was not saved to REDCap.",
        }
    record_id, records = redcap_records_for_current_project()
    result = redcap_request({
        "content": "record", "type": "flat", "overwriteBehavior": "overwrite",
        "data": json.dumps(records, ensure_ascii=False), "returnContent": "count",
    })
    synced_at = utc_now()
    with connect() as db:
        db.execute("UPDATE country_projects SET redcap_record_id=? WHERE id=?", (record_id, active_project_id()))
    save_app_setting("redcap_last_sync_at", synced_at)
    save_app_setting("redcap_last_record_id", record_id)
    return {"ok": True, "record_id": record_id, "rows_sent": len(records), "result": result, "synced_at": synced_at}


def redcap_auto_sync_minutes():
    raw = os.environ.get("AFRICA_CDC_REDCAP_AUTO_SYNC_MINUTES", "0").strip()
    try:
        minutes = int(raw)
    except ValueError as exc:
        raise ValueError("AFRICA_CDC_REDCAP_AUTO_SYNC_MINUTES must be a whole number") from exc
    if minutes < 0:
        raise ValueError("AFRICA_CDC_REDCAP_AUTO_SYNC_MINUTES cannot be negative")
    return minutes


def start_redcap_auto_sync():
    # REDCap-primary mode writes synchronously on every assessment mutation.
    # A background writer could overwrite a record with an unloaded workspace.
    return None


def _split_redcap_multi(value):
    return [part.strip() for part in str(value or "").split(";") if part.strip()]


def pull_current_project_from_redcap(country_name=None, reporting_period=None, reset_group_accounts=None):
    local_phase1 = profile_state().get("phase1") or {}
    country = str(country_name or local_phase1.get("country_name") or "").strip()
    reporting_period = str(reporting_period or local_phase1.get("reporting_period") or "").strip()
    if not country:
        raise ValueError("Select a named country assessment before loading REDCap data")
    all_rows = redcap_request({
        "content": "record", "type": "flat",
        "rawOrLabel": "raw", "rawOrLabelHeaders": "raw", "exportDataAccessGroups": "false",
    })
    if not isinstance(all_rows, list) or not all_rows:
        raise ValueError("No assessment records were returned by REDCap")
    main_rows = [row for row in all_rows if not str(row.get("redcap_repeat_instrument") or "").strip()]
    country_rows = [row for row in main_rows if str(row.get("country_name") or "").strip().casefold() == country.casefold()]
    matches = [row for row in country_rows if str(row.get("reporting_period") or "").strip().casefold() == reporting_period.casefold()]
    if not matches and re.fullmatch(r"\d{4}", reporting_period):
        year_matches = [
            row for row in country_rows
            if re.search(rf"\b{re.escape(reporting_period)}\b", str(row.get("reporting_period") or ""))
        ]
        if len(year_matches) == 1:
            matches = year_matches
        elif len(year_matches) > 1:
            periods = sorted({str(row.get("reporting_period") or "Not specified").strip() for row in year_matches})
            raise ValueError(f"REDCap has several {reporting_period} records for {country}. Select a specific period: {', '.join(periods)}")
    if not matches:
        if country_rows:
            periods = sorted({str(row.get("reporting_period") or "Not specified").strip() for row in country_rows})
            raise ValueError(f"REDCap has data for {country}, but not for {reporting_period}. Available periods: {', '.join(periods)}")
        raise ValueError(f"No REDCap data was found for {country} ({reporting_period})")
    if len(matches) > 1:
        raise ValueError(f"REDCap contains more than one record for {country} ({reporting_period}); remove or merge the duplicate records first")
    reporting_period = str(matches[0].get("reporting_period") or reporting_period).strip()
    record_id = str(matches[0].get("record_id") or "").strip()
    if not record_id:
        raise ValueError(f"The REDCap record for {country} ({reporting_period}) has no record ID")
    exported = [row for row in all_rows if str(row.get("record_id") or "").strip() == record_id]
    if country_name or reporting_period:
        capture_active_project()
        with connect() as db:
            project = db.execute(
                "SELECT id FROM country_projects WHERE country_name=? COLLATE NOCASE AND reporting_period=? COLLATE NOCASE ORDER BY id LIMIT 1",
                (country, reporting_period),
            ).fetchone()
            if project:
                project_id = project[0]
            else:
                title = f"{country} Surveillance Systems Assessment"
                now = utc_now()
                blank = blank_project_payload(country, reporting_period, title)
                cursor = db.execute(
                    "INSERT INTO country_projects(country_name,reporting_period,title,payload,created_at,updated_at) VALUES(?,?,?,?,?,?)",
                    (country, reporting_period, title, json.dumps(blank, ensure_ascii=False), now, now),
                )
                project_id = cursor.lastrowid
            db.execute("UPDATE app_settings SET value=? WHERE key='active_project_id'", (str(project_id),))
    save_app_setting("redcap_last_record_id", record_id)
    # Loading is read-only in REDCap. Retain the source ID so later user edits
    # update this record instead of creating one from the local project ID.
    with connect() as db:
        db.execute("UPDATE country_projects SET redcap_record_id=? WHERE id=?", (record_id, active_project_id()))
    main = next((row for row in exported if not row.get("redcap_repeat_instrument")), exported[0])
    repeat_rows = {}
    for row in exported:
        instrument = str(row.get("redcap_repeat_instrument") or "").strip()
        if instrument:
            repeat_rows.setdefault(instrument, []).append(row)
    for rows in repeat_rows.values():
        rows.sort(key=lambda row: int(row.get("redcap_repeat_instance") or 0))

    phase1 = {
        "country_name": main.get("country_name", ""), "reporting_period": main.get("reporting_period", ""),
        "comments": "", "tools": [],
    }
    for row in repeat_rows.get("inventory_tools", []):
        if not str(row.get("inv_tool_id") or row.get("inv_tool_name") or "").strip():
            continue
        has_api = str(row.get("inv_has_api") or "").strip()
        phase1["tools"].append({
            "_id": row.get("inv_tool_id") or secrets.token_urlsafe(12), "inventory_name": row.get("inv_tool_name", ""),
            "unit_responsible": row.get("inv_unit", ""), "system_type": _split_redcap_multi(row.get("inv_system_type")),
            "surveillance_functions": _split_redcap_multi(row.get("inv_surv_functions")),
            "geographical_coverage": row.get("inv_geo_coverage", ""), "point_of_data_entry": row.get("inv_data_entry", ""),
            "data_captured": row.get("inv_data_captured", ""), "information_users": row.get("inv_info_users", ""),
            "technology": row.get("inv_technology", ""), "has_api": "Yes" if has_api == "1" else "No" if has_api == "0" else "",
            "linked_to_other_systems": row.get("inv_linked_systems", ""), "comments": row.get("inv_comments", ""),
        })

    phase2 = []
    for row in repeat_rows.get("system_profiling", []):
        if not str(row.get("prof_tool_id") or row.get("official_name") or "").strip():
            continue
        profile = {"_tool_id": row.get("prof_tool_id", "")}
        for key in PROFILE_ID_MAP.values():
            value = row.get(key, "")
            profile[key] = _split_redcap_multi(value) if key in PROFILE_MULTI_FIELDS else value
        phase2.append(profile)

    group_rows = [row for row in repeat_rows.get("working_groups", []) if str(row.get("wg_group_id") or row.get("wg_group_name") or "").strip()]
    groups_by_id = {}
    for row in group_rows:
        try: group_id = int(row.get("wg_group_id"))
        except (TypeError, ValueError): continue
        groups_by_id[group_id] = {"id": group_id, "name": row.get("wg_group_name") or f"Group {group_id}"}
    assignment_rows = repeat_rows.get("work_assignments", [])
    for row in assignment_rows:
        try: group_id = int(row.get("assign_group_id"))
        except (TypeError, ValueError): continue
        groups_by_id.setdefault(group_id, {"id": group_id, "name": row.get("assign_group_name") or f"Group {group_id}"})
    if not groups_by_id:
        groups_by_id = {index: {"id": index, "name": f"Group {index}"} for index in range(1, 4)}

    inventory_assignments, profiling_assignments, assignments = [], [], []
    for row in assignment_rows:
        phase = row.get("assign_phase")
        item_id = row.get("assign_item_id")
        try: group_id = int(row.get("assign_group_id"))
        except (TypeError, ValueError): continue
        if not phase or not item_id: continue
        item = {"group_id": group_id, "group_name": groups_by_id.get(group_id, {}).get("name", ""), "assigned_at": row.get("assigned_at") or utc_now()}
        if phase == "inventory": inventory_assignments.append({**item, "tool_id": item_id})
        elif phase == "profiling": profiling_assignments.append({**item, "tool_id": item_id})
        elif phase == "gap": assignments.append({**item, "domain_id": item_id})

    responses = []
    for row in repeat_rows.get("gap_analysis_responses", []):
        if not row.get("gap_domain_id") or not row.get("gap_question_id"): continue
        try: group_id = int(row.get("gap_group_id"))
        except (TypeError, ValueError):
            owner = next((item for item in assignments if item["domain_id"] == row.get("gap_domain_id")), None)
            if not owner: continue
            group_id = owner["group_id"]
        responses.append({
            "domain_id": row.get("gap_domain_id"), "question_id": row.get("gap_question_id"), "group_id": group_id,
            "response": row.get("gap_response", ""), "explanation": row.get("gap_explanation", ""),
            "comment": row.get("gap_comment", ""), "updated_at": row.get("gap_updated_at") or utc_now(),
        })

    payload = {
        "profile": {"phase1": phase1, "phase2": phase2},
        "state": {
            "assessment": {"title": main.get("assessment_title") or "Africa CDC Gap Analysis", "scope": main.get("assessment_scope", "")},
            "groups": list(groups_by_id.values()), "inventory_assignments": inventory_assignments,
            "profiling_assignments": profiling_assignments, "assignments": assignments, "responses": responses,
        },
        "accounts": [],
    }
    # A refresh of this country keeps its group logins; a different country or
    # period must not inherit them. False is reserved for locked report previews.
    if reset_group_accounts is None:
        reset_group_accounts = workspace_scope(local_phase1) != workspace_scope(phase1)
    restore_project(payload, reset_group_accounts=reset_group_accounts)
    capture_active_project()
    pulled_at = utc_now()
    save_app_setting("redcap_last_pull_at", pulled_at)
    return {
        "ok": True, "record_id": record_id, "country_name": country, "reporting_period": reporting_period,
        "pulled_at": pulled_at, "inventory_tools": len(phase1["tools"]),
        "profiles": len(phase2), "working_groups": len(groups_by_id), "assignments": len(assignment_rows),
        "gap_responses": len(responses),
    }


def initialize_database():
    with connect() as db:
        db.executescript("""
            CREATE TABLE IF NOT EXISTS assessments (
                id INTEGER PRIMARY KEY, title TEXT NOT NULL, scope TEXT NOT NULL DEFAULT '', created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS groups (
                id INTEGER PRIMARY KEY, assessment_id INTEGER NOT NULL, name TEXT NOT NULL,
                UNIQUE(assessment_id, name COLLATE NOCASE), FOREIGN KEY(assessment_id) REFERENCES assessments(id)
            );
            CREATE TABLE IF NOT EXISTS assignments (
                assessment_id INTEGER NOT NULL, domain_id TEXT NOT NULL, group_id INTEGER NOT NULL, assigned_at TEXT NOT NULL,
                PRIMARY KEY(assessment_id, domain_id), FOREIGN KEY(group_id) REFERENCES groups(id)
            );
            CREATE TABLE IF NOT EXISTS inventory_tool_assignments (
                assessment_id INTEGER NOT NULL, tool_id TEXT NOT NULL, group_id INTEGER NOT NULL, assigned_at TEXT NOT NULL,
                PRIMARY KEY(assessment_id, tool_id), FOREIGN KEY(group_id) REFERENCES groups(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS profiling_tool_assignments (
                assessment_id INTEGER NOT NULL, tool_id TEXT NOT NULL, group_id INTEGER NOT NULL, assigned_at TEXT NOT NULL,
                PRIMARY KEY(assessment_id, tool_id), FOREIGN KEY(group_id) REFERENCES groups(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS responses (
                assessment_id INTEGER NOT NULL, domain_id TEXT NOT NULL, question_id TEXT NOT NULL, group_id INTEGER NOT NULL,
                response TEXT NOT NULL DEFAULT '', explanation TEXT NOT NULL DEFAULT '', comment TEXT NOT NULL DEFAULT '',
                updated_at TEXT NOT NULL, PRIMARY KEY(assessment_id, question_id), FOREIGN KEY(group_id) REFERENCES groups(id)
            );
            CREATE TABLE IF NOT EXISTS profile_data (
                phase TEXT PRIMARY KEY, payload TEXT NOT NULL DEFAULT '{}', updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY, username TEXT NOT NULL UNIQUE COLLATE NOCASE,
                display_name TEXT NOT NULL, password_hash TEXT NOT NULL,
                role TEXT NOT NULL DEFAULT 'user', active INTEGER NOT NULL DEFAULT 1,
                created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS sessions (
                token_hash TEXT PRIMARY KEY, user_id INTEGER NOT NULL, expires_at TEXT NOT NULL,
                created_at TEXT NOT NULL, FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS country_projects (
                id INTEGER PRIMARY KEY, country_name TEXT NOT NULL DEFAULT '', reporting_period TEXT NOT NULL DEFAULT '',
                title TEXT NOT NULL, payload TEXT, redcap_record_id TEXT, created_at TEXT NOT NULL, updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS app_settings (
                key TEXT PRIMARY KEY, value TEXT NOT NULL
            );
        """)
        user_columns = {row[1] for row in db.execute("PRAGMA table_info(users)")}
        if "phase" not in user_columns:
            db.execute("ALTER TABLE users ADD COLUMN phase TEXT")
        if "group_id" not in user_columns:
            db.execute("ALTER TABLE users ADD COLUMN group_id INTEGER")
        if "last_login_at" not in user_columns:
            db.execute("ALTER TABLE users ADD COLUMN last_login_at TEXT")
        if "project_scope" not in user_columns:
            # Older group accounts have no trustworthy country binding. They
            # remain unusable until the coordinator issues current credentials.
            db.execute("ALTER TABLE users ADD COLUMN project_scope TEXT")
        project_columns = {row[1] for row in db.execute("PRAGMA table_info(country_projects)")}
        if "redcap_record_id" not in project_columns:
            db.execute("ALTER TABLE country_projects ADD COLUMN redcap_record_id TEXT")
        if not db.execute("SELECT 1 FROM assessments LIMIT 1").fetchone():
            db.execute("INSERT INTO assessments(title, scope, created_at) VALUES (?, ?, ?)", ("Africa CDC Gap Analysis", "", utc_now()))
        assessment_id = db.execute("SELECT id FROM assessments ORDER BY id LIMIT 1").fetchone()[0]
        for name in ("Group 1", "Group 2", "Group 3"):
            db.execute("INSERT OR IGNORE INTO groups(assessment_id, name) VALUES (?, ?)", (assessment_id, name))
        active_project = db.execute("SELECT value FROM app_settings WHERE key='active_project_id'").fetchone()
        if not active_project:
            now = utc_now()
            cursor = db.execute("INSERT INTO country_projects(country_name,reporting_period,title,payload,created_at,updated_at) VALUES('','','Country Assessment',NULL,?,?)", (now, now))
            db.execute("INSERT INTO app_settings(key,value) VALUES('active_project_id',?)", (str(cursor.lastrowid),))
        saved_redcap_url = db.execute("SELECT value FROM app_settings WHERE key='redcap_api_url'").fetchone()
        if not saved_redcap_url:
            db.execute("INSERT INTO app_settings(key,value) VALUES('redcap_api_url',?)", (DEFAULT_REDCAP_API_URL,))
        elif "/redcap_v" in str(saved_redcap_url[0]).lower():
            db.execute("UPDATE app_settings SET value=? WHERE key='redcap_api_url'", (DEFAULT_REDCAP_API_URL,))
        # Project snapshots are assessment data, not an authentication store.
        for project in db.execute("SELECT id,payload FROM country_projects WHERE payload IS NOT NULL").fetchall():
            try:
                payload = json.loads(project["payload"])
            except (TypeError, json.JSONDecodeError):
                continue
            if payload.pop("accounts", None) is not None:
                db.execute("UPDATE country_projects SET payload=? WHERE id=?", (json.dumps(payload, ensure_ascii=False), project["id"]))
        db.execute("DELETE FROM sessions WHERE expires_at<=?", (utc_now(),))
    initialize_runtime_workspace()


PASSWORD_ITERATIONS = 310_000
SESSION_HOURS = 12
SECURE_COOKIES = os.environ.get("AFRICA_CDC_SECURE_COOKIES", "").strip().lower() in {"1", "true", "yes", "on"}
PASSWORD_PEPPER = os.environ.get("AFRICA_CDC_PASSWORD_PEPPER", "")
ARGON2_HASHER = PasswordHasher(time_cost=3, memory_cost=65536, parallelism=4, hash_len=32, salt_len=16) if PasswordHasher else None


def password_material(password):
    password = str(password or "")
    if not PASSWORD_PEPPER:
        return password
    return hmac.new(PASSWORD_PEPPER.encode("utf-8"), password.encode("utf-8"), hashlib.sha256).hexdigest()


def password_digest(password, salt=None):
    if not ARGON2_HASHER:
        raise RuntimeError("argon2-cffi is required; run: pip install -r requirements.txt")
    return ARGON2_HASHER.hash(password_material(password))


def password_matches(password, stored):
    if str(stored or "").startswith("$argon2id$"):
        if not ARGON2_HASHER:
            return False
        try:
            return ARGON2_HASHER.verify(stored, password_material(password))
        except (InvalidHashError, VerificationError, VerifyMismatchError):
            return False
    try:
        algorithm, iterations, salt, expected = stored.split("$", 3)
        if algorithm != "pbkdf2_sha256":
            return False
        actual = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), bytes.fromhex(salt), int(iterations)).hex()
        return hmac.compare_digest(actual, expected)
    except (TypeError, ValueError):
        return False


def password_needs_rehash(stored):
    if not str(stored or "").startswith("$argon2id$"):
        return True
    try:
        return ARGON2_HASHER is not None and ARGON2_HASHER.check_needs_rehash(stored)
    except (InvalidHashError, VerificationError):
        return True


def validate_credentials(username, password, display_name=""):
    username = str(username or "").strip().lower()
    display_name = str(display_name or "").strip()
    is_email = bool(re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", username))
    is_username = bool(re.fullmatch(r"[a-z0-9._-]{3,40}", username))
    if not (is_email or is_username) or len(username) > 100:
        raise ValueError("Enter a valid email address or a username using letters, numbers, dots, dashes or underscores")
    if len(str(password or "")) < 15:
        raise ValueError("New passwords must contain at least 15 characters")
    if len(str(password)) > 1024:
        raise ValueError("New passwords must contain at most 1,024 characters")
    if not re.search(r"[A-Z]", str(password)):
        raise ValueError("Password must contain at least one capital letter")
    if not re.search(r"[a-z]", str(password)):
        raise ValueError("Password must contain at least one lowercase letter")
    if not re.search(r"[0-9]", str(password)):
        raise ValueError("Password must contain at least one number")
    if not re.search(r"[^A-Za-z0-9]", str(password)):
        raise ValueError("Password must contain at least one symbol")
    if display_name and len(display_name) > 80:
        raise ValueError("Display name must be 80 characters or fewer")
    return username, str(password), display_name or username


def workspace_scope(phase1=None):
    phase1 = profile_state()["phase1"] if phase1 is None else phase1
    identity = [str(phase1.get(key) or "").strip().casefold()
                for key in ("country_name", "reporting_period")]
    if not all(identity):
        return None
    return hashlib.sha256(json.dumps(identity, ensure_ascii=False).encode("utf-8")).hexdigest()


def clerk_workspace_matches(user):
    if user.get("role") in {"admin", "coordinator"}:
        return True
    scope = workspace_scope()
    return (user.get("role") == "clerk"
            and user.get("phase") in {"inventory", "profiling", "gap"}
            and bool(user.get("group_id")) and bool(scope)
            and hmac.compare_digest(str(user.get("project_scope") or ""), scope))


def session_user(cookie_header):
    cookie = SimpleCookie(); cookie.load(cookie_header or "")
    morsel = cookie.get("africa_cdc_session")
    if not morsel:
        return None
    token_hash = hashlib.sha256(morsel.value.encode("utf-8")).hexdigest()
    now = utc_now()
    with connect() as db:
        db.execute("DELETE FROM sessions WHERE expires_at<=?", (now,))
        row = db.execute("""
            SELECT u.id,u.username,u.display_name,u.role,u.phase,u.group_id,u.last_login_at,u.project_scope FROM sessions s
            JOIN users u ON u.id=s.user_id
            WHERE s.token_hash=? AND s.expires_at>? AND u.active=1
        """, (token_hash, now)).fetchone()
        return dict(row) if row and clerk_workspace_matches(dict(row)) else None


def new_session(user_id):
    token = secrets.token_urlsafe(32)
    token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest()
    expires = (datetime.now(timezone.utc) + timedelta(hours=SESSION_HOURS)).isoformat(timespec="seconds")
    with connect() as db:
        db.execute("INSERT INTO sessions(token_hash,user_id,expires_at,created_at) VALUES(?,?,?,?)", (token_hash, user_id, expires, utc_now()))
    return token


def username_slug(value):
    value = re.sub(r"[^a-z0-9]+", "_", str(value or "").lower()).strip("_")
    return value[:32] or "assessment"


def generated_password():
    # New credentials are longer; existing password hashes are left untouched.
    characters = [
        secrets.choice("ABCDEFGHJKLMNPQRSTUVWXYZ"),
        secrets.choice("abcdefghijkmnopqrstuvwxyz"),
        secrets.choice("23456789"),
        secrets.choice("!@#$%"),
        secrets.choice("abcdefghijkmnopqrstuvwxyz"),
        secrets.choice("23456789"),
        secrets.choice("ABCDEFGHJKLMNPQRSTUVWXYZ"),
    ]
    characters.extend(secrets.choice("ABCDEFGHJKLMNPQRSTUVWXYZabcdefghijkmnopqrstuvwxyz23456789!@#$%") for _ in range(9))
    secrets.SystemRandom().shuffle(characters)
    return "".join(characters)


def application_users():
    with connect() as db:
        rows = db.execute(
            "SELECT id,username,display_name,role,active,created_at,last_login_at "
            "FROM users WHERE role IN ('admin','coordinator') ORDER BY role,display_name,username"
        ).fetchall()
    return [dict(row) for row in rows]


def create_coordinator_account(username, display_name):
    password = generated_password()
    username, password, display_name = validate_credentials(username, password, display_name)
    with connect() as db:
        if db.execute("SELECT 1 FROM users WHERE username=? COLLATE NOCASE", (username,)).fetchone():
            raise ValueError("That username is already in use")
        cursor = db.execute(
            "INSERT INTO users(username,display_name,password_hash,role,active,created_at,phase,group_id) "
            "VALUES(?,?,?,?,1,?,?,?)",
            (username, display_name, password_digest(password), "coordinator", utc_now(), None, None),
        )
    return {
        "users": application_users(),
        "credentials": {"user_id": cursor.lastrowid, "username": username, "password": password, "display_name": display_name},
    }


def coordinator_account(user_id):
    with connect() as db:
        row = db.execute(
            "SELECT id,username,display_name,role,active FROM users WHERE id=?",
            (int(user_id),),
        ).fetchone()
    if not row or row["role"] != "coordinator":
        raise ValueError("Coordinator account not found")
    return dict(row)


def reset_coordinator_password(user_id):
    account = coordinator_account(user_id)
    password = generated_password()
    with connect() as db:
        db.execute(
            "UPDATE users SET password_hash=?,active=1,last_login_at=NULL WHERE id=?",
            (password_digest(password), account["id"]),
        )
        db.execute("DELETE FROM sessions WHERE user_id=?", (account["id"],))
    return {
        "users": application_users(),
        "credentials": {"user_id": account["id"], "username": account["username"], "password": password, "display_name": account["display_name"]},
    }


def admin_change_user_password(user_id, new_password, confirmation):
    if str(new_password or "") != str(confirmation or ""):
        raise ValueError("The new password and confirmation do not match")
    with connect() as db:
        account = db.execute(
            "SELECT id,username,display_name,password_hash,role FROM users WHERE id=? AND active=1",
            (int(user_id),),
        ).fetchone()
    if not account:
        raise ValueError("Active user account not found")
    _, validated_password, _ = validate_credentials(account["username"], new_password, account["display_name"])
    if password_matches(validated_password, account["password_hash"]):
        raise ValueError("Choose a new password that is different from the current password")
    with connect() as db:
        db.execute("UPDATE users SET password_hash=?,last_login_at=NULL WHERE id=?", (password_digest(validated_password), account["id"]))
        db.execute("DELETE FROM sessions WHERE user_id=?", (account["id"],))
    return {"ok": True, "users": application_users()}


def change_own_password(user_id, current_password, new_password, confirmation):
    if str(new_password or "") != str(confirmation or ""):
        raise ValueError("The new password and confirmation do not match")
    with connect() as db:
        account = db.execute(
            "SELECT id,username,display_name,password_hash FROM users WHERE id=? AND active=1",
            (int(user_id),),
        ).fetchone()
    if not account or not password_matches(current_password, account["password_hash"]):
        raise ValueError("The current password is incorrect")
    _, validated_password, _ = validate_credentials(account["username"], new_password, account["display_name"])
    if password_matches(validated_password, account["password_hash"]):
        raise ValueError("Choose a new password that is different from the current password")
    with connect() as db:
        db.execute("UPDATE users SET password_hash=? WHERE id=?", (password_digest(validated_password), account["id"]))
        db.execute("DELETE FROM sessions WHERE user_id=?", (account["id"],))
    return {"ok": True}


def set_coordinator_active(user_id, active):
    account = coordinator_account(user_id)
    active = 1 if bool(active) else 0
    with connect() as db:
        db.execute("UPDATE users SET active=? WHERE id=?", (active, account["id"]))
        if not active:
            db.execute("DELETE FROM sessions WHERE user_id=?", (account["id"],))
    return {"users": application_users()}


def short_group_username(country, phase, group_name, group_id):
    country_code = re.sub(r"[^a-z0-9]", "", str(country or "").lower())[:2] or "ct"
    phase_code = {"inventory": "i", "profiling": "p", "gap": "g"}[phase]
    match = re.search(r"\d+", str(group_name or ""))
    group_code = (match.group(0) if match else str(group_id))[-3:]
    return f"{country_code}{phase_code}{group_code}"[:7]


def ensure_group_account(assessment_id, phase, group_id):
    if phase not in {"inventory", "profiling", "gap"}:
        raise ValueError("Unknown assessment phase")
    with connect() as db:
        scope = workspace_scope()
        if not scope:
            raise ValueError("Select a country and reporting period before creating group credentials")
        existing = db.execute("SELECT id,username FROM users WHERE role='clerk' AND phase=? AND group_id=? AND project_scope=? AND active=1", (phase, group_id, scope)).fetchone()
        if existing:
            return None
        group = db.execute("SELECT name FROM groups WHERE id=? AND assessment_id=?", (group_id, assessment_id)).fetchone()
        if not group:
            raise ValueError("Participant group not found")
        country = profile_state()["phase1"].get("country_name") or "country"
        base = short_group_username(country, phase, group["name"], group_id)
        username, suffix = base, 1
        while db.execute("SELECT 1 FROM users WHERE username=? COLLATE NOCASE", (username,)).fetchone():
            suffix += 1
            suffix_text = str(suffix)
            username = f"{base[:7-len(suffix_text)]}{suffix_text}"
        password = generated_password()
        db.execute(
            "INSERT INTO users(username,display_name,password_hash,role,active,created_at,phase,group_id,project_scope) VALUES(?,?,?,?,1,?,?,?,?)",
            (username, f"{group['name']} — {phase.title()}", password_digest(password), "clerk", utc_now(), phase, group_id, scope),
        )
    return {"username": username, "password": password, "phase": phase, "group_name": group["name"], "country": country}


def reset_group_account(assessment_id, phase, group_id):
    if phase not in {"inventory", "profiling", "gap"}:
        raise ValueError("Unknown assessment phase")
    with connect() as db:
        scope = workspace_scope()
        if not scope:
            raise ValueError("Select a country and reporting period before creating group credentials")
        group = db.execute("SELECT name FROM groups WHERE id=? AND assessment_id=?", (group_id, assessment_id)).fetchone()
        if not group:
            raise ValueError("Participant group not found")
        existing = db.execute("SELECT id FROM users WHERE role='clerk' AND phase=? AND group_id=? ORDER BY (project_scope=?) DESC,id LIMIT 1", (phase, group_id, scope)).fetchone()
        country = profile_state()["phase1"].get("country_name") or "country"
        base = short_group_username(country, phase, group["name"], group_id)
        username, suffix = base, 1
        while db.execute("SELECT 1 FROM users WHERE username=? COLLATE NOCASE AND id<>?", (username, existing["id"] if existing else -1)).fetchone():
            suffix += 1
            suffix_text = str(suffix)
            username = f"{base[:7-len(suffix_text)]}{suffix_text}"
        password = generated_password()
        if existing:
            db.execute("DELETE FROM sessions WHERE user_id=?", (existing["id"],))
            db.execute("UPDATE users SET username=?,password_hash=?,active=1,last_login_at=NULL,project_scope=? WHERE id=?", (username, password_digest(password), scope, existing["id"]))
        else:
            db.execute(
                "INSERT INTO users(username,display_name,password_hash,role,active,created_at,phase,group_id,project_scope) VALUES(?,?,?,?,1,?,?,?,?)",
                (username, f"{group['name']} — {phase.title()}", password_digest(password), "clerk", utc_now(), phase, group_id, scope),
            )
    capture_active_project()
    return {"username": username, "password": password, "phase": phase, "group_name": group["name"], "country": country}


def current_state():
    with connect() as db:
        assessment = dict(db.execute("SELECT * FROM assessments ORDER BY id LIMIT 1").fetchone())
        aid = assessment["id"]
        groups = [dict(row) for row in db.execute("SELECT * FROM groups WHERE assessment_id=? ORDER BY id", (aid,))]
        assignments = [dict(row) for row in db.execute("""
            SELECT a.domain_id, a.group_id, g.name AS group_name, a.assigned_at
            FROM assignments a JOIN groups g ON g.id=a.group_id WHERE a.assessment_id=?
        """, (aid,))]
        responses = [dict(row) for row in db.execute("SELECT * FROM responses WHERE assessment_id=?", (aid,))]
        inventory_assignments = [dict(row) for row in db.execute("""
            SELECT a.tool_id, a.group_id, g.name AS group_name, a.assigned_at
            FROM inventory_tool_assignments a JOIN groups g ON g.id=a.group_id WHERE a.assessment_id=?
        """, (aid,))]
        profiling_assignments = [dict(row) for row in db.execute("""
            SELECT a.tool_id, a.group_id, g.name AS group_name, a.assigned_at
            FROM profiling_tool_assignments a JOIN groups g ON g.id=a.group_id WHERE a.assessment_id=?
        """, (aid,))]
    return {"app_version": APP_VERSION, "assessment": assessment, "groups": groups, "assignments": assignments, "inventory_assignments": inventory_assignments, "profiling_assignments": profiling_assignments, "responses": responses, "domains": DOMAINS}


def scoped_state(user):
    state = current_state()
    if not user or user.get("role") in {"admin", "coordinator"}:
        return state
    group_id, phase = user.get("group_id"), user.get("phase")
    state["groups"] = [group for group in state["groups"] if group["id"] == group_id]
    if phase == "inventory":
        state["inventory_assignments"] = [item for item in state["inventory_assignments"] if item["group_id"] == group_id]
        state["profiling_assignments"], state["assignments"], state["responses"], state["domains"] = [], [], [], []
    elif phase == "profiling":
        state["profiling_assignments"] = [item for item in state["profiling_assignments"] if item["group_id"] == group_id]
        state["inventory_assignments"], state["assignments"], state["responses"], state["domains"] = [], [], [], []
    else:
        state["assignments"] = [item for item in state["assignments"] if item["group_id"] == group_id]
        domain_ids = {item["domain_id"] for item in state["assignments"]}
        state["domains"] = [domain for domain in state["domains"] if domain["id"] in domain_ids]
        state["responses"] = [item for item in state["responses"] if item["domain_id"] in domain_ids]
        state["inventory_assignments"], state["profiling_assignments"] = [], []
    return state


def ensure_inventory_tool_ids(payload):
    payload = dict(payload or {})
    payload["tools"] = [dict(tool or {}) for tool in payload.get("tools", [])]
    seen = set()
    for tool in payload["tools"]:
        tool_id = str(tool.get("_id") or "").strip()
        if not tool_id or tool_id in seen:
            tool_id = secrets.token_urlsafe(12)
        tool["_id"] = tool_id
        seen.add(tool_id)
    return payload


def profile_state():
    with connect() as db:
        rows = db.execute("SELECT phase, payload FROM profile_data").fetchall()
    stored = {row["phase"]: json.loads(row["payload"]) for row in rows}
    original_phase1 = stored.get("phase1", {"country_name": "", "reporting_period": "", "comments": "", "tools": []})
    phase1 = ensure_inventory_tool_ids(original_phase1)
    if phase1 != original_phase1:
        with connect() as db:
            db.execute(
                "INSERT INTO profile_data(phase,payload,updated_at) VALUES('phase1',?,?) "
                "ON CONFLICT(phase) DO UPDATE SET payload=excluded.payload,updated_at=excluded.updated_at",
                (json.dumps(phase1, ensure_ascii=False), utc_now()),
            )
    return {
        "phase1": phase1,
        "phase2": stored.get("phase2", []),
    }


def scoped_profile(user):
    profile = profile_state()
    if not user or user.get("role") in {"admin", "coordinator"}:
        return profile
    group_id, phase = user.get("group_id"), user.get("phase")
    state = current_state()
    if phase == "inventory":
        tool_ids = {item["tool_id"] for item in state["inventory_assignments"] if item["group_id"] == group_id}
        profile["phase1"]["tools"] = [tool for tool in profile["phase1"].get("tools", []) if tool.get("_id") in tool_ids]
        profile["phase2"] = []
    elif phase == "profiling":
        tool_ids = {item["tool_id"] for item in state["profiling_assignments"] if item["group_id"] == group_id}
        profile["phase1"]["tools"] = [tool for tool in profile["phase1"].get("tools", []) if tool.get("_id") in tool_ids]
        profile["phase2"] = [item for item in profile["phase2"] if item.get("_tool_id") in tool_ids]
    else:
        profile = {"phase1": {"country_name": profile["phase1"].get("country_name", ""), "reporting_period": profile["phase1"].get("reporting_period", ""), "comments": "", "tools": []}, "phase2": []}
    return profile


def my_entries_data(user):
    if not user or user.get("role") != "clerk":
        raise PermissionError("A group account is required")
    state, profile = scoped_state(user), scoped_profile(user)
    group = next((item for item in state.get("groups", []) if item.get("id") == user.get("group_id")), None)
    phase1 = profile.get("phase1") or {}
    phase = user.get("phase")
    if phase == "inventory":
        items = (profile.get("phase1") or {}).get("tools", [])
        summary = {"assigned": len(items), "started": sum(any(value not in (None, "", []) for key, value in item.items() if not key.startswith("_") and key != "inventory_name") for item in items), "responses": sum(sum(value not in (None, "", []) for key, value in item.items() if not key.startswith("_")) for item in items), "findings": 0}
    elif phase == "profiling":
        assigned = len((profile.get("phase1") or {}).get("tools", [])); items = profile.get("phase2") or []
        summary = {"assigned": assigned, "started": len(items), "responses": sum(sum(value not in (None, "", []) for key, value in item.items() if not key.startswith("_")) for item in items), "findings": 0}
    else:
        assigned = len(state.get("assignments", [])); responses = state.get("responses", [])
        summary = {"assigned": assigned, "started": len({item["domain_id"] for item in responses if item.get("response") or item.get("explanation") or item.get("comment")}), "responses": sum(bool(item.get("response")) for item in responses), "findings": sum(bool(item.get("comment")) for item in responses)}
    return {
        "phase": phase,
        "group": (group or {}).get("name") or user.get("display_name") or "Group",
        "country": phase1.get("country_name") or "",
        "reporting_period": phase1.get("reporting_period") or "",
        "profile": profile,
        "state": state,
        "summary": summary,
    }


def my_entries_text(user):
    data = my_entries_data(user)
    summary = data["summary"]
    lines = ["AFRICA CDC — GROUP OVERVIEW REPORT", f"Country: {data['country']}", f"Reporting period: {data['reporting_period']}", f"Group: {data['group']}", f"Section: {data['phase'].title()}", "OVERVIEW", f"Assigned items: {summary['assigned']}", f"Items started: {summary['started']}", f"Saved responses: {summary['responses']}", f"Recorded findings/comments: {summary['findings']}", "", "ENTERED DETAILS", ""]
    if data["phase"] == "inventory":
        for index, tool in enumerate((data["profile"].get("phase1") or {}).get("tools", []), 1):
            lines.append(f"TOOL {index}: {tool.get('inventory_name') or 'Unnamed tool'}")
            lines.extend(f"{key.replace('_', ' ').title()}: {', '.join(value) if isinstance(value, list) else value}" for key, value in tool.items() if not key.startswith("_") and value not in (None, "", []))
            lines.append("")
    elif data["phase"] == "profiling":
        for index, profile in enumerate(data["profile"].get("phase2") or [], 1):
            lines.append(f"SYSTEM {index}: {profile.get('official_name') or 'Unnamed system'}")
            lines.extend(f"{key.replace('_', ' ').title()}: {', '.join(value) if isinstance(value, list) else value}" for key, value in profile.items() if not key.startswith("_") and value not in (None, "", []))
            lines.append("")
    else:
        responses = {item["question_id"]: item for item in data["state"].get("responses", [])}
        assigned = {item["domain_id"] for item in data["state"].get("assignments", [])}
        for domain in data["state"].get("domains", []):
            if domain["id"] not in assigned: continue
            lines.append(domain["name"].upper())
            for question in domain["questions"]:
                saved = responses.get(question["id"], {})
                answer = saved.get("response") or ""
                if answer: lines.extend([question["text"], f"Response: {answer}"])
                if saved.get("explanation"): lines.append(f"Explanation: {saved['explanation']}")
                if saved.get("comment"): lines.append(f"Comment: {saved['comment']}")
                if answer or saved.get("explanation") or saved.get("comment"): lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def save_profile_phase(phase, payload):
    if phase not in {"phase1", "phase2"}:
        raise ValueError("Unknown assessment phase")
    if phase == "phase1" and not isinstance(payload, dict):
        raise ValueError("Inventory data must be an object")
    if phase == "phase1":
        payload = ensure_inventory_tool_ids(payload)
    if phase == "phase2" and isinstance(payload, dict):
        payload = [payload]
    if phase == "phase2" and not isinstance(payload, list):
        raise ValueError("Profiling data must be a list")
    with connect() as db:
        db.execute(
            "INSERT INTO profile_data(phase,payload,updated_at) VALUES(?,?,?) "
            "ON CONFLICT(phase) DO UPDATE SET payload=excluded.payload, updated_at=excluded.updated_at",
            (phase, json.dumps(payload, ensure_ascii=False), utc_now()),
        )
    return profile_state()


def save_inventory_tool(payload, user=None):
    tool = dict(payload.get("tool") or {})
    tool_id = str(tool.get("_id") or payload.get("tool_id") or "").strip()
    if not tool_id:
        raise ValueError("The inventory tool identifier is required")
    tool["_id"] = tool_id
    try:
        group_id = int(payload.get("group_id")) if payload.get("group_id") not in (None, "") else None
    except (TypeError, ValueError):
        raise ValueError("A valid working group is required")
    with connect() as db:
        clerk = user is not None and user.get("role") not in {"admin", "coordinator"}
        if clerk:
            if user.get("phase") != "inventory" or not clerk_workspace_matches(user):
                raise PermissionError("This group account cannot edit this country assessment")
            group_id = user["group_id"]
        assessment_id = db.execute("SELECT id FROM assessments ORDER BY id LIMIT 1").fetchone()[0]
        owner = db.execute("SELECT group_id FROM inventory_tool_assignments WHERE assessment_id=? AND tool_id=?", (assessment_id, tool_id)).fetchone()
        if clerk and (not owner or owner[0] != group_id):
            raise PermissionError("This inventory tool is not assigned to your group")
        if owner and group_id is not None and group_id != owner[0]:
            raise PermissionError("This inventory tool is assigned to another group")
        row = db.execute("SELECT payload FROM profile_data WHERE phase='phase1'").fetchone()
        phase1 = ensure_inventory_tool_ids(json.loads(row[0]) if row else {})
        tools = phase1.get("tools", [])
        index = next((i for i, item in enumerate(tools) if item.get("_id") == tool_id), None)
        if clerk:
            if index is None:
                raise PermissionError("This inventory tool is not available to your group")
            for key in ("country_name", "reporting_period", "comments"):
                if key in payload and str(payload.get(key) or "").strip() != str(phase1.get(key) or "").strip():
                    raise PermissionError("Only a coordinator can change shared assessment details")
        if index is None:
            tools.append(tool)
        else:
            tools[index] = tool
        for key in ("country_name", "reporting_period", "comments"):
            if key in payload and not clerk:
                phase1[key] = str(payload.get(key) or "")
        db.execute(
            "INSERT INTO profile_data(phase,payload,updated_at) VALUES('phase1',?,?) "
            "ON CONFLICT(phase) DO UPDATE SET payload=excluded.payload,updated_at=excluded.updated_at",
            (json.dumps(phase1, ensure_ascii=False), utc_now()),
        )
    return profile_state()


def save_profiling_tool(payload, user=None):
    profile = dict(payload.get("profile") or {})
    tool_id = str(profile.get("_tool_id") or payload.get("tool_id") or "").strip()
    if not tool_id:
        raise ValueError("The profiling system identifier is required")
    profile["_tool_id"] = tool_id
    try:
        group_id = int(payload.get("group_id")) if payload.get("group_id") not in (None, "") else None
    except (TypeError, ValueError):
        raise ValueError("A valid working group is required")
    with connect() as db:
        clerk = user is not None and user.get("role") not in {"admin", "coordinator"}
        if clerk:
            if user.get("phase") != "profiling" or not clerk_workspace_matches(user):
                raise PermissionError("This group account cannot edit this country assessment")
            group_id = user["group_id"]
        assessment_id = db.execute("SELECT id FROM assessments ORDER BY id LIMIT 1").fetchone()[0]
        owner = db.execute("SELECT group_id FROM profiling_tool_assignments WHERE assessment_id=? AND tool_id=?", (assessment_id, tool_id)).fetchone()
        if clerk and (not owner or owner[0] != group_id):
            raise PermissionError("This profiling system is not assigned to your group")
        if clerk and not any(item.get("_id") == tool_id for item in profile_state()["phase1"].get("tools", [])):
            raise PermissionError("This profiling system is not available to your group")
        if owner and group_id is not None and group_id != owner[0]:
            raise PermissionError("This profiling system is assigned to another group")
        row = db.execute("SELECT payload FROM profile_data WHERE phase='phase2'").fetchone()
        profiles = json.loads(row[0]) if row else []
        index = next((i for i, item in enumerate(profiles) if item.get("_tool_id") == tool_id), None)
        if index is None:
            profiles.append(profile)
        else:
            profiles[index] = profile
        db.execute(
            "INSERT INTO profile_data(phase,payload,updated_at) VALUES('phase2',?,?) "
            "ON CONFLICT(phase) DO UPDATE SET payload=excluded.payload,updated_at=excluded.updated_at",
            (json.dumps(profiles, ensure_ascii=False), utc_now()),
        )
    return profile_state()


def reset_assessment():
    """Start like the desktop application: a fresh unsaved form on every launch."""
    with connect() as db:
        aid = db.execute("SELECT id FROM assessments ORDER BY id LIMIT 1").fetchone()[0]
        db.execute("DELETE FROM responses WHERE assessment_id=?", (aid,))
        db.execute("DELETE FROM assignments WHERE assessment_id=?", (aid,))
        db.execute("DELETE FROM inventory_tool_assignments WHERE assessment_id=?", (aid,))
        db.execute("DELETE FROM profiling_tool_assignments WHERE assessment_id=?", (aid,))
        db.execute("DELETE FROM groups WHERE assessment_id=?", (aid,))
        for name in ("Group 1", "Group 2", "Group 3"):
            db.execute("INSERT INTO groups(assessment_id,name) VALUES(?,?)", (aid, name))
        db.execute("DELETE FROM profile_data")
        db.execute("UPDATE assessments SET title=?, scope=? WHERE id=?", ("Africa CDC Gap Analysis", "", aid))


def active_project_id():
    with connect() as db:
        row = db.execute("SELECT value FROM app_settings WHERE key='active_project_id'").fetchone()
    return int(row[0])


def capture_active_project():
    profile, state = profile_state(), current_state()
    payload = {"profile": profile, "state": state}
    project_id = active_project_id()
    country = str(profile["phase1"].get("country_name") or "").strip()
    period = str(profile["phase1"].get("reporting_period") or "").strip()
    title = state["assessment"].get("title") or f"{country or 'Country'} Assessment"
    with connect() as db:
        db.execute("UPDATE country_projects SET country_name=?,reporting_period=?,title=?,payload=?,updated_at=? WHERE id=?", (country, period, title, json.dumps(payload, ensure_ascii=False), utc_now(), project_id))
    return project_id


def blank_project_payload(country_name, reporting_period, title):
    groups = [{"id": index, "name": f"Group {index}"} for index in range(1, 4)]
    return {
        "profile": {"phase1": {"country_name": country_name, "reporting_period": reporting_period, "comments": "", "tools": []}, "phase2": []},
        "state": {"assessment": {"title": title, "scope": ""}, "groups": groups, "assignments": [], "inventory_assignments": [], "profiling_assignments": [], "responses": []},
        "accounts": [],
    }


def restore_project(payload, reset_group_accounts=True):
    payload = json.loads(payload) if isinstance(payload, str) else dict(payload or {})
    profile, state = payload.get("profile") or {}, payload.get("state") or {}
    with connect() as db:
        assessment = db.execute("SELECT id FROM assessments ORDER BY id LIMIT 1").fetchone()
        aid = assessment[0]
        if reset_group_accounts:
            clerk_ids = [row[0] for row in db.execute("SELECT id FROM users WHERE role='clerk'")]
            if clerk_ids:
                placeholders = ",".join("?" for _ in clerk_ids)
                db.execute(f"DELETE FROM sessions WHERE user_id IN ({placeholders})", clerk_ids)
            db.execute("DELETE FROM users WHERE role='clerk'")
        db.execute("DELETE FROM responses WHERE assessment_id=?", (aid,))
        db.execute("DELETE FROM assignments WHERE assessment_id=?", (aid,))
        db.execute("DELETE FROM inventory_tool_assignments WHERE assessment_id=?", (aid,))
        db.execute("DELETE FROM profiling_tool_assignments WHERE assessment_id=?", (aid,))
        db.execute("DELETE FROM groups WHERE assessment_id=?", (aid,))
        db.execute("DELETE FROM profile_data")
        assessment_data = state.get("assessment") or {}
        db.execute("UPDATE assessments SET title=?,scope=? WHERE id=?", (assessment_data.get("title") or "Africa CDC Gap Analysis", assessment_data.get("scope") or "", aid))
        old_to_new = {}
        for group in state.get("groups") or [{"id": i, "name": f"Group {i}"} for i in range(1, 4)]:
            if not reset_group_accounts and group.get("id"):
                cursor = db.execute("INSERT INTO groups(id,assessment_id,name) VALUES(?,?,?)", (group.get("id"), aid, group.get("name") or "Group"))
            else:
                cursor = db.execute("INSERT INTO groups(assessment_id,name) VALUES(?,?)", (aid, group.get("name") or "Group"))
            old_to_new[group.get("id")] = cursor.lastrowid
        for phase in ("phase1", "phase2"):
            value = profile.get(phase, {"country_name": "", "reporting_period": "", "comments": "", "tools": []} if phase == "phase1" else [])
            db.execute("INSERT INTO profile_data(phase,payload,updated_at) VALUES(?,?,?)", (phase, json.dumps(value, ensure_ascii=False), utc_now()))
        for item in state.get("inventory_assignments", []):
            group_id = old_to_new.get(item.get("group_id"))
            if group_id: db.execute("INSERT INTO inventory_tool_assignments VALUES(?,?,?,?)", (aid, item.get("tool_id"), group_id, item.get("assigned_at") or utc_now()))
        for item in state.get("profiling_assignments", []):
            group_id = old_to_new.get(item.get("group_id"))
            if group_id: db.execute("INSERT INTO profiling_tool_assignments VALUES(?,?,?,?)", (aid, item.get("tool_id"), group_id, item.get("assigned_at") or utc_now()))
        for item in state.get("assignments", []):
            group_id = old_to_new.get(item.get("group_id"))
            if group_id: db.execute("INSERT INTO assignments VALUES(?,?,?,?)", (aid, item.get("domain_id"), group_id, item.get("assigned_at") or utc_now()))
        for item in state.get("responses", []):
            group_id = old_to_new.get(item.get("group_id"))
            if group_id: db.execute("INSERT INTO responses VALUES(?,?,?,?,?,?,?,?)", (aid, item.get("domain_id"), item.get("question_id"), group_id, item.get("response", ""), item.get("explanation", ""), item.get("comment", ""), item.get("updated_at") or utc_now()))


def country_projects_state():
    active = active_project_id()
    profile = profile_state()["phase1"]
    with connect() as db:
        rows = [dict(row) for row in db.execute("SELECT id,country_name,reporting_period,title,redcap_record_id,created_at,updated_at FROM country_projects ORDER BY created_at,id")]
    for row in rows:
        row["source"] = "redcap" if str(row.pop("redcap_record_id", "") or "").strip() else "new"
        row["active"] = row["id"] == active
        if row["active"]:
            row["country_name"] = profile.get("country_name") or row.get("country_name", "")
            row["reporting_period"] = profile.get("reporting_period") or row.get("reporting_period", "")
    rows = [
        row for row in rows
        if str(row.get("country_name") or "").strip()
        and str(row.get("reporting_period") or "").strip()
    ]
    return {"active_project_id": active, "projects": rows}


def create_country_project(country_name, reporting_period, open_existing=False):
    country_name, reporting_period = str(country_name or "").strip(), str(reporting_period or "").strip()
    if not country_name or not reporting_period:
        raise ValueError("Country and reporting period are required")
    temporary_test = is_test_country_name(country_name)
    with connect() as db:
        existing = db.execute(
            "SELECT id,redcap_record_id,payload FROM country_projects WHERE country_name=? COLLATE NOCASE AND reporting_period=? COLLATE NOCASE LIMIT 1",
            (country_name, reporting_period),
        ).fetchone()
    if existing:
        if not temporary_test:
            if not open_existing:
                raise ValueError(f"An assessment already exists for {country_name} ({reporting_period}). Open it to update the information?")
            capture_active_project()
            with connect() as db:
                existing = db.execute(
                    "SELECT id,redcap_record_id,payload FROM country_projects WHERE id=?", (existing["id"],)
                ).fetchone()
            if str(existing["redcap_record_id"] or "").strip():
                pull_current_project_from_redcap(country_name, reporting_period)
            else:
                with connect() as db:
                    db.execute("UPDATE app_settings SET value=? WHERE key='active_project_id'", (str(existing["id"]),))
                payload = json.loads(existing["payload"] or "{}")
                restore_project(payload or blank_project_payload(country_name, reporting_period))
            result = country_projects_state()
            result["opened_existing"] = True
            result["existing_source"] = "redcap" if str(existing["redcap_record_id"] or "").strip() else "new"
            return result
        # Test Country is a disposable in-memory workspace. Starting it again
        # means reset the prior temporary workspace instead of treating it as a
        # saved country assessment or reporting a duplicate.
        title = f"{country_name} Surveillance Systems Assessment"
        payload = blank_project_payload(country_name, reporting_period, title)
        now = utc_now()
        with connect() as db:
            db.execute(
                "UPDATE country_projects SET title=?,payload=?,redcap_record_id=NULL,updated_at=? WHERE id=?",
                (title, json.dumps(payload, ensure_ascii=False), now, existing[0]),
            )
            db.execute("UPDATE app_settings SET value=? WHERE key='active_project_id'", (str(existing[0]),))
        restore_project(payload)
        return country_projects_state()
    capture_active_project()
    title = f"{country_name} Surveillance Systems Assessment"
    payload = blank_project_payload(country_name, reporting_period, title)
    now = utc_now()
    with connect() as db:
        cursor = db.execute("INSERT INTO country_projects(country_name,reporting_period,title,payload,created_at,updated_at) VALUES(?,?,?,?,?,?)", (country_name, reporting_period, title, json.dumps(payload, ensure_ascii=False), now, now))
        project_id = cursor.lastrowid
        db.execute("UPDATE app_settings SET value=? WHERE key='active_project_id'", (str(project_id),))
    restore_project(payload)
    return country_projects_state()


def switch_country_project(project_id):
    project_id = int(project_id)
    if project_id == active_project_id(): return country_projects_state()
    with connect() as db:
        row = db.execute("SELECT country_name,reporting_period FROM country_projects WHERE id=?", (project_id,)).fetchone()
        if not row: raise ValueError("Country assessment not found")
    # The loader captures the current workspace before selecting the target.
    # Updating the active ID here would attach the old data to the new country.
    pull_current_project_from_redcap(row["country_name"], row["reporting_period"])
    return country_projects_state()


def inventory_findings_and_recommendations(tools):
    findings, recommendations = [], []
    if not tools:
        return findings, recommendations

    total = len(tools)
    national = []
    limited_coverage = []
    no_api = []
    api_unknown = []
    not_linked = []
    linkage_unknown = []
    access_unknown = []
    users_unknown = []
    manual_systems = []
    incomplete = []

    for tool in tools:
        name = str(tool.get("inventory_name") or "Unnamed system").strip()
        coverage = str(tool.get("geographical_coverage") or "").strip()
        api = str(tool.get("has_api") or "").strip().casefold()
        linked = str(tool.get("linked_to_other_systems") or "").strip().casefold()
        technology = str(tool.get("technology") or "").strip().casefold()
        coverage_key = normalized_key(coverage)
        if coverage_key in {"national", "national_scale", "national_level", "national_coverage", "countrywide", "country_wide"} or coverage_key.startswith("national_"):
            national.append(name)
        elif coverage:
            limited_coverage.append(f"{name} ({coverage})")
        else:
            limited_coverage.append(f"{name} (coverage not recorded)")
        if api == "no":
            no_api.append(name)
        elif api != "yes":
            api_unknown.append(name)
        if linked == "no":
            not_linked.append(name)
        elif linked != "yes":
            linkage_unknown.append(name)
        if not str(tool.get("point_of_data_entry") or "").strip():
            access_unknown.append(name)
        if not str(tool.get("information_users") or "").strip():
            users_unknown.append(name)
        if any(term in technology for term in ("paper", "manual", "spreadsheet", "excel")):
            manual_systems.append(name)
        missing = [
            label for key, label in (
                ("unit_responsible", "responsible unit"),
                ("geographical_coverage", "geographic coverage"),
                ("point_of_data_entry", "point of data entry"),
                ("information_users", "information users"),
                ("technology", "technology"),
                ("has_api", "API availability"),
                ("linked_to_other_systems", "system linkages"),
            ) if not str(tool.get(key) or "").strip()
        ]
        if missing:
            incomplete.append(f"{name}: {', '.join(missing)}")

    findings.append(f"The Inventory records {total} surveillance system{'s' if total != 1 else ''}.")
    findings.append(f"{len(national)} of {total} systems report national geographic coverage.")
    if limited_coverage:
        findings.append(f"Access limitation: systems without confirmed national coverage are {', '.join(limited_coverage)}.")
        recommendations.append("Review geographic access gaps and develop a phased plan to extend priority systems to underserved administrative levels and facilities.")
    if no_api:
        findings.append(f"Interoperability limitation: {', '.join(no_api)} {'do' if len(no_api) != 1 else 'does'} not provide an API.")
        recommendations.append(f"Assess standards-based API integration for {', '.join(no_api)} to reduce manual exchange and duplicate data entry.")
    if api_unknown:
        findings.append(f"API availability was not confirmed for {', '.join(api_unknown)}.")
        recommendations.append(f"Confirm and document API capability, ownership, and access procedures for {', '.join(api_unknown)}.")
    if not_linked:
        findings.append(f"Integration limitation: {', '.join(not_linked)} {'are' if len(not_linked) != 1 else 'is'} not linked to other information systems.")
        recommendations.append(f"Prioritise an interoperability plan for {', '.join(not_linked)}, including target systems, data standards, governance, and implementation milestones.")
    if linkage_unknown:
        findings.append(f"System linkages were not confirmed for {', '.join(linkage_unknown)}.")
        recommendations.append(f"Document current and planned system linkages for {', '.join(linkage_unknown)}.")
    if access_unknown:
        findings.append(f"Access limitation: points of data entry were not recorded for {', '.join(access_unknown)}.")
        recommendations.append(f"Document where and how users enter data for {', '.join(access_unknown)}, including facility, district, national, mobile, and offline access.")
    if users_unknown:
        findings.append(f"User access groups were not recorded for {', '.join(users_unknown)}.")
        recommendations.append(f"Define authorised user groups, roles, and access levels for {', '.join(users_unknown)}.")
    if manual_systems:
        findings.append(f"Potential access and timeliness limitation: {', '.join(manual_systems)} use manual, paper, or spreadsheet-based technology.")
        recommendations.append(f"Assess digitisation and offline-capable data-entry options for {', '.join(manual_systems)} while maintaining appropriate paper contingency procedures.")
    if incomplete:
        findings.append(f"Key Inventory details remain incomplete for {'; '.join(incomplete)}.")
        recommendations.append("Complete and validate the missing Inventory fields with system owners before finalising investment or integration decisions.")
    if not recommendations:
        recommendations.append("Maintain the documented coverage and interoperability arrangements, and validate them periodically with system owners and users.")
    return list(dict.fromkeys(findings)), list(dict.fromkeys(recommendations))


def complete_report_text():
    profile = profile_state()
    phase1 = dict(profile.get("phase1") or {})
    phase1["phase2"] = profile.get("phase2") or []
    comments = str(phase1.get("comments", "") or "")
    tools = [tool for tool in phase1.get("tools", []) if str(tool.get("inventory_name", "")).strip()]
    findings, recommendations = [], []

    def classify_lines(value, default_target):
        for item in re.split(r"[\n;]+", str(value or "")):
            item = item.strip()
            if not item:
                continue
            prefix, separator, detail = item.partition(":")
            category = prefix.strip().casefold()
            text = detail.strip() if separator and category in {"gap", "gaps", "finding", "findings", "recommendation", "recommendations"} else item
            target = recommendations if category in {"recommendation", "recommendations"} else findings if category in {"gap", "gaps", "finding", "findings"} else default_target
            target.append(text)

    classify_lines(comments, recommendations)
    for tool in tools:
        classify_lines(tool.get("comments", ""), findings)
    derived_findings, derived_recommendations = inventory_findings_and_recommendations(tools)
    findings.extend(derived_findings)
    recommendations.extend(derived_recommendations)
    findings = list(dict.fromkeys(findings))
    recommendations = list(dict.fromkeys(recommendations))
    if generate_desktop_report is not None:
        text = generate_desktop_report(phase1, findings, recommendations)
    else:
        text = "Africa CDC Surveillance Systems Profiling and Gap Analysis\n\n"
        text += f"Country: {phase1.get('country_name') or 'Not provided'}\nReporting Period: {phase1.get('reporting_period') or 'Not provided'}\n"
    return text.rstrip() + "\n\n" + report_text()


def report_data():
    state = current_state()
    inventory_tools = [
        tool for tool in (profile_state().get("phase1") or {}).get("tools", [])
        if str(tool.get("inventory_name") or "").strip()
    ]
    inventory_findings, inventory_recommendations = inventory_findings_and_recommendations(inventory_tools)
    assignment_map = {item["domain_id"]: item for item in state["assignments"]}
    response_map = {item["question_id"]: item for item in state["responses"]}
    sections = []
    for domain in DOMAINS:
        gaps, recommendations, other = [], [], []
        for question in domain["questions"]:
            comment = response_map.get(question["id"], {}).get("comment", "")
            for line in comment.splitlines():
                line = line.strip()
                if not line:
                    continue
                prefix, sep, detail = line.partition(":")
                detail = detail.strip() if sep else line
                if prefix.strip().lower().rstrip("s") == "gap":
                    gaps.append(detail)
                elif prefix.strip().lower().rstrip("s") == "recommendation":
                    recommendations.append(detail)
                else:
                    other.append(line)
        if not gaps and not recommendations and not other:
            continue
        sections.append({
            "domain_id": domain["id"], "domain": domain["name"],
            "group": assignment_map.get(domain["id"], {}).get("group_name", "Unassigned"),
            "gaps": list(dict.fromkeys(gaps)), "recommendations": list(dict.fromkeys(recommendations)),
            "other": list(dict.fromkeys(other)),
        })
    return {
        "assessment": state["assessment"],
        "sections": sections,
        "inventory_analysis": {
            "findings": inventory_findings,
            "recommendations": inventory_recommendations,
        },
    }


def report_text():
    report = report_data()
    lines = [report["assessment"]["title"], "Gap Analysis: Gaps and Recommendations Summary", ""]
    for section in report["sections"]:
        lines.extend([section["domain"], f"Assigned group: {section['group']}", "Gaps:"])
        lines.extend([f"- {item}" for item in section["gaps"] + section["other"]] or ["- None recorded."])
        lines.append("Recommendations:")
        lines.extend([f"- {item}" for item in section["recommendations"]] or ["- None recorded."])
        lines.append("")
    return "\n".join(lines)


def report_docx():
    if Document is None:
        raise RuntimeError("python-docx is required for Word export")
    from report_builder import build_report
    return build_report(profile_state(), current_state(), DOMAINS, report_data(), STATIC_DIR / "africa_cdc_logo_full.png")
    report = report_data()
    document = Document()
    document.add_heading(report["assessment"]["title"], 0)
    document.add_heading("Complete Assessment Report", level=1)
    for line in complete_report_text().splitlines():
        document.add_paragraph(line)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
    document.add_heading("Gap Analysis: Gaps and Recommendations Summary", level=1)
    for section in report["sections"]:
        document.add_heading(section["domain"], level=2)
        document.add_paragraph(f"Assigned group: {section['group']}")
        table = document.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "Gaps"
        table.rows[0].cells[1].text = "Recommendations"
        row = table.add_row().cells
        row[0].text = "\n".join(f"• {x}" for x in section["gaps"] + section["other"]) or "None recorded."
        row[1].text = "\n".join(f"• {x}" for x in section["recommendations"]) or "None recorded."
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def full_profile():
    profile = profile_state()
    state = current_state()
    phase3 = {
        "gap_analysis_name": state["assessment"].get("title", ""),
        "gap_analysis_scope": state["assessment"].get("scope", ""),
        "_groups": [group["name"] for group in state["groups"]],
        "_domain_assignments": {item["domain_id"]: item["group_name"] for item in state["assignments"]},
    }
    for item in state["responses"]:
        value = item.get("response", "")
        try:
            parsed = json.loads(value)
            value = parsed if isinstance(parsed, list) else value
        except (TypeError, ValueError):
            pass
        phase3[item["question_id"]] = value
        if item.get("explanation"):
            phase3[f'{item["question_id"]}_explanation'] = item["explanation"]
        if item.get("comment"):
            phase3[f'{item["question_id"]}_comment'] = item["comment"]
    phase1 = profile["phase1"]
    form = {"country_name": phase1.get("country_name", ""), "reporting_period": phase1.get("reporting_period", ""),
            "tools": phase1.get("tools", []), "findings": "", "recommendations": phase1.get("comments", ""),
            "phase2": profile["phase2"], "phase3": phase3}
    return {"version": PROFILE_FORMAT_VERSION, "form": form, "profile": profile, "phase3": state}


def import_full_profile(payload):
    profile = payload.get("profile", payload)
    if "form" in payload:  # Desktop profile compatibility.
        form = payload.get("form") or {}
        profile = {"phase1": {"country_name": form.get("country_name", ""), "reporting_period": form.get("reporting_period", ""), "comments": form.get("recommendations", ""), "tools": form.get("tools", [])}, "phase2": form.get("phase2", [])}
    save_profile_phase("phase1", profile.get("phase1", {}))
    save_profile_phase("phase2", profile.get("phase2", []))
    phase3 = payload.get("phase3")
    desktop_phase3 = (payload.get("form") or {}).get("phase3")
    if isinstance(desktop_phase3, dict):
        groups = desktop_phase3.get("_groups") or ["Group 1", "Group 2", "Group 3"]
        assignments = desktop_phase3.get("_domain_assignments") or {}
        with connect() as db:
            aid = db.execute("SELECT id FROM assessments ORDER BY id LIMIT 1").fetchone()[0]
            db.execute("UPDATE assessments SET title=?, scope=? WHERE id=?", (desktop_phase3.get("gap_analysis_name") or "Africa CDC Gap Analysis", desktop_phase3.get("gap_analysis_scope") or "", aid))
            name_to_id = {}
            for name in groups:
                name = str(name).strip()
                if not name: continue
                db.execute("INSERT OR IGNORE INTO groups(assessment_id,name) VALUES(?,?)", (aid, name))
                name_to_id[name] = db.execute("SELECT id FROM groups WHERE assessment_id=? AND name=? COLLATE NOCASE", (aid, name)).fetchone()[0]
            db.execute("DELETE FROM assignments WHERE assessment_id=?", (aid,))
            db.execute("DELETE FROM responses WHERE assessment_id=?", (aid,))
            default_group = next(iter(name_to_id.values()))
            for domain in DOMAINS:
                answered = [q for q in domain["questions"] if desktop_phase3.get(q["id"]) not in (None, "", []) or desktop_phase3.get(f'{q["id"]}_comment')]
                owner_name = assignments.get(domain["name"], assignments.get(domain["id"], ""))
                group_id = name_to_id.get(owner_name, default_group)
                if answered or owner_name:
                    db.execute("INSERT INTO assignments VALUES(?,?,?,?)", (aid, domain["id"], group_id, utc_now()))
                for question in answered:
                    qid = question["id"]
                    value = desktop_phase3.get(qid, "")
                    if isinstance(value, list): value = json.dumps(value)
                    db.execute("INSERT INTO responses VALUES(?,?,?,?,?,?,?,?)", (aid, domain["id"], qid, group_id, str(value), str(desktop_phase3.get(f'{qid}_explanation', '')), str(desktop_phase3.get(f'{qid}_comment', '')), utc_now()))
    if isinstance(phase3, dict) and isinstance(phase3.get("assessment"), dict):
        with connect() as db:
            assessment = db.execute("SELECT id FROM assessments ORDER BY id LIMIT 1").fetchone()
            aid = assessment[0]
            source_assessment = phase3["assessment"]
            db.execute("UPDATE assessments SET title=?, scope=? WHERE id=?", (source_assessment.get("title") or "Africa CDC Gap Analysis", source_assessment.get("scope") or "", aid))
            old_to_new = {}
            for group in phase3.get("groups", []):
                name = str(group.get("name", "")).strip()
                if not name: continue
                db.execute("INSERT OR IGNORE INTO groups(assessment_id,name) VALUES(?,?)", (aid, name))
                new_id = db.execute("SELECT id FROM groups WHERE assessment_id=? AND name=? COLLATE NOCASE", (aid, name)).fetchone()[0]
                old_to_new[group.get("id")] = new_id
            db.execute("DELETE FROM assignments WHERE assessment_id=?", (aid,))
            db.execute("DELETE FROM responses WHERE assessment_id=?", (aid,))
            for item in phase3.get("assignments", []):
                group_id = old_to_new.get(item.get("group_id"))
                if group_id and item.get("domain_id") in {d["id"] for d in DOMAINS}:
                    db.execute("INSERT INTO assignments(assessment_id,domain_id,group_id,assigned_at) VALUES(?,?,?,?)", (aid, item["domain_id"], group_id, item.get("assigned_at") or utc_now()))
            for item in phase3.get("responses", []):
                group_id = old_to_new.get(item.get("group_id"))
                if group_id and item.get("question_id") in QUESTION_INDEX:
                    domain_id = QUESTION_INDEX[item["question_id"]][0]
                    db.execute("INSERT INTO responses(assessment_id,domain_id,question_id,group_id,response,explanation,comment,updated_at) VALUES(?,?,?,?,?,?,?,?)", (aid, domain_id, item["question_id"], group_id, item.get("response", ""), item.get("explanation", ""), item.get("comment", ""), item.get("updated_at") or utc_now()))
    return profile_state()


def parse_excel_files(files):
    if openpyxl is None:
        raise RuntimeError("openpyxl is required for Excel import")
    imported = {}
    for file_data in files:
        workbook = workbook_from_upload(file_data)
        for sheet in workbook.worksheets:
            standard_gap = STANDARD_GAP_SHEETS.get(sheet.title.strip().casefold())
            indexes = None
            for row in sheet.iter_rows(values_only=True):
                normalized = [str(cell or "").strip().lower().rstrip(".") for cell in row]
                response_names = {"response", "responses", "answer", "answers"}
                if "question" in normalized and any(item in response_names for item in normalized):
                    indexes = {
                        "id": next((i for i, value in enumerate(normalized) if value in {"no", "question id", "id"}), None),
                        "question": normalized.index("question"),
                        "response": next(i for i, value in enumerate(normalized) if value in response_names),
                        "comments": [i for i, value in enumerate(normalized) if value in {"comment", "comments", "gap", "gaps", "recommendation", "recommendations", "comment / gap / recommendation"}],
                        "explanation": next((i for i, value in enumerate(normalized) if value in {"explanation", "explanation / details"}), None),
                        "group": next((i for i, value in enumerate(normalized) if value in {"group", "participant group", "assigned group"}), None),
                        "headers": normalized,
                    }
                    continue
                if not indexes:
                    continue
                qid = str(row[indexes["id"]] or "").strip().lower() if indexes["id"] is not None and indexes["id"] < len(row) else ""
                if qid not in QUESTION_INDEX:
                    continue
                domain_id, question = QUESTION_INDEX[qid]
                if standard_gap and domain_id not in standard_gap[0]:
                    continue
                response = row[indexes["response"]] if indexes["response"] < len(row) else ""
                standard_explanation = ""
                if standard_gap and isinstance(response, str) and re.search(r"\nExplanation:\s*", response, re.I):
                    response, standard_explanation = re.split(r"\nExplanation:\s*", response, maxsplit=1, flags=re.I)
                comments = []
                for index in indexes["comments"]:
                    if index < len(row) and row[index] not in (None, ""):
                        raw_heading = indexes["headers"][index]
                        text = str(row[index]).strip()
                        if raw_heading in {"comment", "comments", "comment / gap / recommendation"}:
                            comments.append(text)
                        else:
                            comments.append(f"{raw_heading.rstrip('s').title()}: {text}")
                current = imported.setdefault(question["id"], {"domain_id": domain_id, "response": "", "explanation": "", "comment": ""})
                if standard_gap:
                    current["group_slot"] = standard_gap[1]
                if response not in (None, ""):
                    current["response"] = str(response).strip()
                if standard_explanation:
                    current["explanation"] = standard_explanation.strip()
                if indexes["explanation"] is not None and indexes["explanation"] < len(row) and row[indexes["explanation"]] not in (None, ""):
                    current["explanation"] = str(row[indexes["explanation"]]).strip()
                if indexes["group"] is not None and indexes["group"] < len(row) and row[indexes["group"]] not in (None, ""):
                    current["group_name"] = str(row[indexes["group"]]).strip()
                if comments:
                    current["comment"] = "\n".join(comments)
    return imported


def normalized_key(value):
    return "_".join("".join(ch.lower() if ch.isalnum() else " " for ch in str(value or "")).split())


def profiling_name_keys(value):
    words = re.findall(r"[A-Za-z0-9]+", str(value or ""))
    lowered = [word.casefold() for word in words]
    full = "".join(lowered)
    acronym = "".join(
        word[0] for word in lowered
        if word not in {"and", "of", "the", "through", "for", "system"}
    )
    explicit = [word.casefold() for word in words if len(word) > 1 and (word.isupper() or any(ch.isdigit() for ch in word))]
    return {key for key in [full, acronym, *explicit] if len(key) >= 2}


PROFILE_ID_MAP = {
    "A1":"official_name","A2":"platforms","A4":"implementation_status","A5":"year_implemented","B1":"surveillance_approaches",
    "C1":"diseases_supported","C2":"captures_signals","C3":"captures_verification","C4":"alert_risk_assessment",
    "D1":"case_notifications","D2":"unique_identifier","D3":"case_investigation","D4":"specimen_requests","D5":"specimen_tracking","D6":"lab_results","D7":"lab_results_received","D8":"contact_identification","D9":"contact_followup","D10":"dashboards_surveillance","D11":"gis_mapping","D12":"epidemic_curves","D13":"trend_analysis","D14":"hotspot_analysis","D15":"contact_dashboards","D16":"lab_dashboards","D17":"mortality_dashboards","D18":"custom_visualisations","D19":"automated_alerts","D20":"threshold_monitoring","D21":"line_lists","D22":"situation_reports","D23":"weekly_reports","D24":"monthly_reports","D25":"scheduled_reports","D26":"bulletin_outputs","D27":"excel_export","D28":"pdf_export","D29":"api_access",
    "E1":"primary_users","E2":"active_users","F1":"coverage_level","F2":"district_count","F3":"facility_count",
    "G1":"web_access","G2":"mobile_interface","G3":"offline_access","G4":"api_available","G5":"open_source","G6":"hosting_model","G7":"data_exchange",
    "H1":"ownership_institution","H2":"hosting_institution","H3":"sop_available","H4":"user_manuals","H5":"data_policy","H6":"security_controls",
    "I1":"emergency_support","I2":"supported_events","I3":"lessons_learned","J1":"technical_support","J2":"expertise_sufficient","J3":"funding_source",
    "K1":"retained_capabilities","K2":"operational_challenges","K3":"planned_enhancements","K4":"implementation_lessons","K5":"known_limitations","K6":"priority_actions","L1":"other_information",
}
PROFILE_MULTI_FIELDS = {"platforms", "surveillance_approaches", "diseases_supported", "primary_users"}


def normalize_profile_import_value(key, value):
    text = str(value or "").strip()
    folded = re.sub(r"[^a-z0-9]+", " ", text.casefold()).strip()
    categorical = {
        "yes": "Yes", "no": "No", "planned": "Planned", "unknown": "Unknown",
        "na": "N/A", "n a": "N/A", "not applicable": "N/A",
    }
    if folded in categorical:
        return categorical[folded]
    if key == "implementation_status":
        statuses = {
            "active": "Active", "pilot": "Pilot", "planned": "Planned",
            "dormant": "Dormant", "retired": "Retired",
            "outbreak response only": "Outbreak Response Only", "unknown": "Unknown",
        }
        return statuses.get(folded, text)
    aliases = {
        "platforms": {
            "kobo": "Kobo system", "kobo system": "Kobo system",
            "kobotool": "Kobo system", "kobotoolbox": "Kobo system",
            "open srp": "OpenSRP",
        },
        "diseases_supported": {
            "ebola": "Ebola virus disease",
            "ebola virus disease": "Ebola virus disease",
            "ebola bundibugyo virus disease": "Ebola virus disease",
            "bundibugyo": "Ebola virus disease",
            "bundibugyo virus disease": "Ebola virus disease",
            "marburg": "Marburg virus disease",
        },
        "coverage_level": {
            "country wide": "National",
            "countrywide": "National",
            "national coverage": "National",
        },
        "primary_users": {
            "countyteams": "CountyTeams",
            "county teams": "CountyTeams",
        },
    }
    return aliases.get(key, {}).get(folded, text)


def split_profile_multi_value(key, value):
    # Slash is part of valid choices such as AFP/Polio and HL7/FHIR. It is a
    # separator only for institution and platform entries such as MoH/NPHI.
    separator = r"[;,\n/\\]+" if key in {"platforms", "ownership_institution", "hosting_institution"} else r"[;,\n]+"
    return [item.strip() for item in re.split(separator, str(value)) if item.strip()]


def parse_phase2_workbook(workbook, phase1):
    phase1 = ensure_inventory_tool_ids(phase1)
    tools = phase1.setdefault("tools", [])
    inventory_by_name = {}
    inventory_tools = []
    for tool in tools:
        inventory_name = str(tool.get("inventory_name") or "").strip()
        if not inventory_name:
            continue
        inventory_tools.append(tool)
        aliases = [inventory_name, *re.split(r"[/\\|]+", inventory_name)]
        for alias in aliases:
            key = normalized_key(alias)
            if key:
                inventory_by_name.setdefault(key, tool)
                inventory_by_name.setdefault(key.replace("_", ""), tool)
    profile_sheets = [sheet for sheet in workbook.worksheets if sheet.title.casefold().startswith("sp_")]
    if not profile_sheets:
        for sheet in workbook.worksheets:
            ids = {str(sheet.cell(row, 2).value or "").strip().upper() for row in range(2, min(sheet.max_row, 100) + 1)}
            headers = {normalized_key(cell.value) for cell in sheet[1]}
            if len(ids & set(PROFILE_ID_MAP)) >= 60 and {"question", "response"}.issubset(headers):
                if any(sheet.cell(row, 6).value not in (None, "") for row in range(2, min(sheet.max_row, 100) + 1)):
                    profile_sheets.append(sheet)

    assessments = []
    for sheet in profile_sheets:
        rows = list(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 200), values_only=True))
        assessment = {"official_name": sheet.title}
        field_comments = {}
        imported_comment_lines = []
        header_index = next((i for i, row in enumerate(rows) if {"question", "response"}.issubset({normalized_key(x) for x in row})), None)
        if header_index is not None:
            headers = [normalized_key(x) for x in rows[header_index]]
            qi, ri = headers.index("question"), headers.index("response")
            ci = headers.index("comments") if "comments" in headers else None
            for row in rows[header_index + 1:]:
                if qi >= len(row) or not row[qi]:
                    continue
                source_id = str(row[qi - 1] or "").strip().upper() if qi else ""
                key = PROFILE_ID_MAP.get(source_id, normalized_key(row[qi]))
                value = row[ri] if ri < len(row) else ""
                # Some completed workbooks put the connected system name in G7
                # instead of the expected Yes/No response. Preserve that useful
                # detail and record the exchange capability as affirmative.
                if source_id == "G7" and value not in (None, ""):
                    exchange_text = str(value).strip()
                    exchange_key = re.sub(r"[^a-z0-9]+", " ", exchange_text.casefold()).strip()
                    if exchange_key not in {"yes", "no", "unknown", "y", "n", "true", "false", "1", "0"}:
                        assessment["data_exchange"] = "Yes"
                        assessment["exchange_comments"] = exchange_text
                        value = "Yes"
                if value not in (None, ""):
                    if key in PROFILE_MULTI_FIELDS:
                        assessment[key] = [
                            normalize_profile_import_value(key, item)
                            for item in split_profile_multi_value(key, value)
                        ]
                    else:
                        assessment[key] = normalize_profile_import_value(key, value)
                if ci is not None and ci < len(row) and row[ci] not in (None, ""):
                    comment = str(row[ci]).strip()
                    field_comments[key] = comment
                    imported_comment_lines.append(f"{source_id or key}: {comment}")
                    if source_id in {"K1", "K2", "K3", "K4", "K5", "K6"} and not str(assessment.get(key) or "").strip():
                        assessment[key] = comment
                    if source_id == "G7":
                        assessment["exchange_comments"] = comment
        elif rows:
            headers = [normalized_key(x) for x in rows[0]]
            if len(rows) > 1:
                assessment.update({key: str(value).strip() for key, value in zip(headers, rows[1]) if key and value not in (None, "")})

        official_name = str(assessment.get("official_name") or sheet.title).strip()
        if field_comments:
            assessment["field_comments"] = field_comments
            assessment["imported_comments"] = "\n".join(imported_comment_lines)
        official_key, sheet_key = normalized_key(official_name), normalized_key(sheet.title)
        tool = (
            inventory_by_name.get(official_key)
            or inventory_by_name.get(official_key.replace("_", ""))
            or inventory_by_name.get(sheet_key)
            or inventory_by_name.get(sheet_key.replace("_", ""))
        )
        if tool is None:
            source_keys = profiling_name_keys(official_name) | profiling_name_keys(sheet.title)
            candidates = [
                candidate for candidate in inventory_tools
                if source_keys & profiling_name_keys(candidate.get("inventory_name"))
            ]
            if len(candidates) == 1:
                tool = candidates[0]
        if tool is None:
            # The Inventory is authoritative. A completed or stale Profiling
            # tab must never create a system that was not inventoried in the
            # workbook currently being imported.
            continue
        assessment["official_name"] = str(tool.get("inventory_name") or official_name).strip()
        assessment["_tool_id"] = tool["_id"]
        inventory_by_name[official_key] = tool
        inventory_by_name[official_key.replace("_", "")] = tool
        inventory_by_name[sheet_key] = tool
        inventory_by_name[sheet_key.replace("_", "")] = tool
        if len(assessment) > 2:
            assessments.append(assessment)
    return phase1, assessments


EXCEL_MAX_UPLOAD_BYTES = 10 * 1024 * 1024
EXCEL_MAX_EXPANDED_BYTES = 50 * 1024 * 1024
EXCEL_MAX_ZIP_ENTRIES = 2000
EXCEL_MAX_SHEETS = 100
EXCEL_MAX_ROWS = 10000
EXCEL_MAX_COLUMNS = 100
EXCEL_MAX_CELLS = 1000000
EXCEL_MAX_XML_ELEMENTS = 1000000


def _check_excel_archive(data):
    """Bound ZIP/XML expansion and sheet geometry before openpyxl allocates cells."""
    def check_range(reference):
        try:
            _, _, max_column, max_row = openpyxl.utils.range_boundaries(reference)
        except (TypeError, ValueError):
            raise ValueError("The Excel workbook contains an invalid cell range") from None
        if not max_column or not max_row or max_column > EXCEL_MAX_COLUMNS or max_row > EXCEL_MAX_ROWS:
            raise ValueError("Excel sheets may contain at most 10,000 rows and 100 columns")
        return max_row, max_column

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > EXCEL_MAX_ZIP_ENTRIES or sum(item.file_size for item in entries) > EXCEL_MAX_EXPANDED_BYTES:
                raise ValueError("The Excel workbook expands beyond the supported import size (50 MiB)")
            if len({item.filename for item in entries}) != len(entries):
                raise ValueError("The Excel workbook contains duplicate archive entries")
            total_elements = total_cells = worksheet_count = 0
            for item in entries:
                if item.flag_bits & 1:
                    raise ValueError("Password-protected Excel files cannot be imported")
                # Inspect every XML part, including worksheet parts named by relationships.
                with archive.open(item) as source:
                    part = source.read(EXCEL_MAX_EXPANDED_BYTES + 1)
                if len(part) > EXCEL_MAX_EXPANDED_BYTES:
                    raise ValueError("The Excel workbook expands beyond the supported import size (50 MiB)")
                xml_prefix = part[:200].replace(b"\x00", b"").lstrip()
                if not item.filename.lower().endswith((".xml", ".rels")) and not xml_prefix.startswith(b"<"):
                    continue
                if re.search(br"<!\s*(?:DOCTYPE|ENTITY)\b", part.replace(b"\x00", b""), re.I):
                    raise ValueError("XML document types and entities are not supported in Excel imports")
                sheet_rows = sheet_columns = row_index = column_index = 0
                worksheet = False
                for event, element in ElementTree.iterparse(io.BytesIO(part), events=("start", "end")):
                    if event == "end":
                        element.clear()
                        continue
                    total_elements += 1
                    if total_elements > EXCEL_MAX_XML_ELEMENTS:
                        raise ValueError("The Excel workbook contains too many XML elements")
                    tag = element.tag.rsplit("}", 1)[-1]
                    if tag == "worksheet":
                        worksheet = True
                        worksheet_count += 1
                        if worksheet_count > EXCEL_MAX_SHEETS:
                            raise ValueError("Excel imports support at most 100 worksheets")
                    if not worksheet:
                        continue
                    if tag in {"dimension", "mergeCell"} and element.get("ref"):
                        rows, columns = check_range(element.get("ref"))
                        sheet_rows, sheet_columns = max(sheet_rows, rows), max(sheet_columns, columns)
                    elif tag == "row":
                        row_index = int(element.get("r", row_index + 1))
                        if not 1 <= row_index <= EXCEL_MAX_ROWS:
                            raise ValueError("Excel sheets may contain at most 10,000 rows")
                        sheet_rows, column_index = max(sheet_rows, row_index), 0
                    elif tag == "c":
                        if element.get("r"):
                            rows, column_index = check_range(element.get("r"))
                        else:
                            rows, column_index = row_index or 1, column_index + 1
                        if column_index > EXCEL_MAX_COLUMNS:
                            raise ValueError("Excel sheets may contain at most 100 columns")
                        sheet_rows, sheet_columns = max(sheet_rows, rows), max(sheet_columns, column_index)
                total_cells += sheet_rows * sheet_columns
                if total_cells > EXCEL_MAX_CELLS:
                    raise ValueError("The Excel workbook contains too many cells (maximum 1,000,000 including blank ranges)")
    except (zipfile.BadZipFile, ElementTree.ParseError, RuntimeError, NotImplementedError) as exc:
        raise ValueError("The uploaded file is not a supported, readable Excel workbook") from exc


def workbook_from_upload(file_data):
    if openpyxl is None:
        raise RuntimeError("openpyxl is required for Excel import")
    payload = file_data.get("data", "")
    if not isinstance(payload, str):
        raise ValueError("The Excel upload must contain base64 file data")
    # Bound allocation before splitting a data URL or decoding its base64 payload.
    if len(payload) > 4 * ((EXCEL_MAX_UPLOAD_BYTES + 2) // 3) + 200:
        raise ValueError("Excel files must be 10 MiB or smaller")
    payload = payload.split(",", 1)[-1]
    try:
        data = base64.b64decode(payload, validate=True)
    except (ValueError, base64.binascii.Error) as exc:
        raise ValueError("The Excel upload contains invalid base64 file data") from exc
    if len(data) > EXCEL_MAX_UPLOAD_BYTES:
        raise ValueError("Excel files must be 10 MiB or smaller")
    _check_excel_archive(data)
    return openpyxl.load_workbook(io.BytesIO(data), data_only=True, keep_links=False)


def validate_standardized_excel(file_data):
    workbook = workbook_from_upload(file_data)
    errors, warnings = [], []
    standard_inventory = next((sheet for sheet in workbook.worksheets if sheet.title.strip().casefold() == "his inventory"), None)
    if standard_inventory:
        header_row = next((
            row_index for row_index, row in enumerate(
                standard_inventory.iter_rows(min_row=1, max_row=min(20, standard_inventory.max_row), values_only=True), 1
            ) if "system" in [normalized_key(value) for value in row]
        ), None)
        headers = [normalized_key(cell.value) for cell in standard_inventory[header_row][:12]] if header_row else []
        required_headers = ["system", "unit_responsible", "type_of_system", "surveillance_functions_select_all", "geographical_coverage", "point_of_data_entry", "data_captured", "information_user", "technology", "has_api", "linked_to_other_systems", "comments"]
        missing = [name for name in required_headers if name not in headers]
        if missing: errors.append("HIS Inventory is missing columns: " + ", ".join(missing))
        inventory_names = [
            str(standard_inventory.cell(row, 1).value or "").strip()
            for row in range((header_row or 1) + 1, standard_inventory.max_row + 1)
            if str(standard_inventory.cell(row, 1).value or "").strip()
        ]
        profile_names = []
        for sheet in workbook.worksheets:
            ids = {str(sheet.cell(row, 2).value or "").strip().upper() for row in range(2, min(sheet.max_row, 80) + 1)}
            if len(ids & set(PROFILE_ID_MAP)) >= 60:
                name = next((str(sheet.cell(row, 6).value or "").strip() for row in range(2, min(sheet.max_row, 80) + 1) if str(sheet.cell(row, 2).value or "").strip().upper() == "A1"), "")
                if name: profile_names.append(name)
        tool_names = inventory_names or profile_names
        if not tool_names: errors.append("No named tools were found in HIS Inventory or the tool profiling tabs")
        required_gap = set(STANDARD_GAP_SHEETS)
        found_gap = {sheet.title.strip().casefold() for sheet in workbook.worksheets} & required_gap
        if found_gap:
            missing_gap = sorted(required_gap - found_gap)
            if missing_gap: errors.append("Missing grouped Gap Analysis sheets: " + ", ".join(missing_gap))
        expected_ids = {str(question.get("source_id", question["id"])).upper() for domain in DOMAINS for question in domain["questions"]}
        found_ids = set()
        for sheet in workbook.worksheets:
            standard = STANDARD_GAP_SHEETS.get(sheet.title.strip().casefold())
            if not standard: continue
            allowed = standard[0]
            for row in range(1, min(sheet.max_row, 150) + 1):
                source_id = str(sheet.cell(row, 1).value or "").strip().upper()
                if source_id.lower() in QUESTION_INDEX and QUESTION_INDEX[source_id.lower()][0] in allowed: found_ids.add(source_id)
        if found_gap:
            missing_ids = sorted(expected_ids - found_ids)
            if missing_ids: errors.append("Grouped Gap Analysis sheets are missing Question IDs: " + ", ".join(missing_ids))
        if errors: raise ValueError("Standardized workbook validation failed:\n- " + "\n- ".join(errors))
        return {"valid": True, "inventory_tools": len(tool_names), "profiling_systems": len(profile_names), "gap_questions": len(found_ids), "warnings": warnings}

    # Accept a standalone Profiling workbook (for example EWARS1.xlsx). These
    # files contain the complete A1-L1 questionnaire but intentionally omit
    # Inventory and Gap Analysis sheets. Phase 1 will create the named system
    # entry and Phase 2 will import all responses and comments.
    standalone_profile_sheets = []
    for sheet in workbook.worksheets:
        ids = {
            str(sheet.cell(row, 2).value or "").strip().upper()
            for row in range(2, min(sheet.max_row, 100) + 1)
        }
        headers = {normalized_key(cell.value) for cell in sheet[1]}
        if len(ids & set(PROFILE_ID_MAP)) >= 60 and {"question", "response"}.issubset(headers):
            standalone_profile_sheets.append(sheet)
    if standalone_profile_sheets:
        profile_names = []
        for sheet in standalone_profile_sheets:
            ids = []
            official_name = ""
            for row in sheet.iter_rows(min_row=2, max_row=min(sheet.max_row, 200), values_only=True):
                source_id = str(row[1] or "").strip().upper() if len(row) > 1 else ""
                if source_id in PROFILE_ID_MAP:
                    ids.append(source_id)
                    if source_id == "A1" and len(row) > 5:
                        official_name = str(row[5] or "").strip()
            missing_ids = sorted(set(PROFILE_ID_MAP) - set(ids))
            if missing_ids:
                errors.append(f"{sheet.title} is missing Profiling field IDs: {', '.join(missing_ids)}")
            if not official_name:
                errors.append(f"{sheet.title} has no system name in field A1")
            else:
                profile_names.append(official_name)
        if errors:
            raise ValueError("Profiling workbook validation failed:\n- " + "\n- ".join(errors))
        warnings.append("Profiling-only workbook: Inventory entries will be created from the completed system tabs.")
        return {
            "valid": True,
            "inventory_tools": len(profile_names),
            "profiling_systems": len(profile_names),
            "gap_questions": 0,
            "warnings": warnings,
        }

    inventory_sheet = next((sheet for sheet in workbook.worksheets if sheet.title.strip().casefold() == "inventory"), None)
    profile_sheets = [sheet for sheet in workbook.worksheets if sheet.title.upper().startswith("SP_")]
    gap_sheet = next((sheet for sheet in workbook.worksheets if sheet.title.strip().casefold() == "gap analysis responses"), None)
    if not inventory_sheet: errors.append("Required sheet 'Inventory' is missing")
    if not profile_sheets: errors.append("No Profiling sheets were found; sheet names must start with 'SP_'")
    if not gap_sheet: errors.append("Required sheet 'Gap Analysis Responses' is missing")
    tool_names = []
    if inventory_sheet:
        headers = [normalized_key(cell.value) for cell in inventory_sheet[1]]
        required = ["country", "reporting_period", "overall_comments", "tool_system_name", "unit_responsible", "type_of_system", "surveillance_functions", "geographical_coverage", "point_of_data_entry", "data_captured", "information_users", "technology", "has_api", "linked_to_other_systems", "comments"]
        missing = [header for header in required if header not in headers]
        if missing: errors.append("Inventory is missing columns: " + ", ".join(missing))
        name_index = headers.index("tool_system_name") if "tool_system_name" in headers else None
        if name_index is not None:
            tool_names = [str(row[name_index].value or "").strip() for row in inventory_sheet.iter_rows(min_row=2) if str(row[name_index].value or "").strip()]
            folded = [name.casefold() for name in tool_names]
            if len(folded) != len(set(folded)): errors.append("Inventory contains duplicate Tool / System names")
        if not tool_names: errors.append("Inventory contains no named tools or systems")
    profile_names = []
    expected_profile_ids = set(PROFILE_ID_MAP)
    for sheet in profile_sheets:
        headers = [normalized_key(cell.value) for cell in sheet[1]]
        if not {"question", "response"}.issubset(headers):
            errors.append(f"{sheet.title} must contain Question and Response columns"); continue
        question_index, response_index = headers.index("question"), headers.index("response")
        id_index = question_index - 1
        ids, official_name = [], ""
        for row in sheet.iter_rows(min_row=2):
            source_id = str(row[id_index].value or "").strip().upper() if id_index >= 0 else ""
            if source_id: ids.append(source_id)
            if source_id == "A1": official_name = str(row[response_index].value or "").strip()
        missing_ids = sorted(expected_profile_ids - set(ids))
        unknown_ids = sorted(set(ids) - expected_profile_ids)
        if missing_ids: errors.append(f"{sheet.title} is missing Profiling field IDs: {', '.join(missing_ids)}")
        if unknown_ids: errors.append(f"{sheet.title} contains unknown Profiling field IDs: {', '.join(unknown_ids)}")
        if len(ids) != len(set(ids)): errors.append(f"{sheet.title} contains duplicate Profiling field IDs")
        if not official_name: errors.append(f"{sheet.title} has no system name in field A1")
        else: profile_names.append(official_name)
    if tool_names and profile_names:
        inventory_set, profile_set = {name.casefold() for name in tool_names}, {name.casefold() for name in profile_names}
        for name in sorted(inventory_set - profile_set): errors.append(f"No Profiling sheet matches Inventory tool: {name}")
        for name in sorted(profile_set - inventory_set): errors.append(f"Profiling system is not listed in Inventory: {name}")
        if len(profile_names) != len(set(name.casefold() for name in profile_names)): errors.append("Profiling contains duplicate system names")
    gap_rows = 0
    if gap_sheet:
        headers = [normalized_key(cell.value) for cell in gap_sheet[1]]
        required = {"domain", "question_id", "question", "participant_group", "response", "explanation", "comments", "last_updated"}
        missing = sorted(required - set(headers))
        if missing: errors.append("Gap Analysis Responses is missing columns: " + ", ".join(missing))
        if "question_id" in headers:
            index = headers.index("question_id")
            ids = [str(row[index].value or "").strip().upper() for row in gap_sheet.iter_rows(min_row=2) if str(row[index].value or "").strip()]
            expected = {str(question.get("source_id", question["id"])).upper() for domain in DOMAINS for question in domain["questions"]}
            missing_ids, unknown_ids = sorted(expected - set(ids)), sorted(set(ids) - expected)
            if missing_ids: errors.append("Gap Analysis is missing Question IDs: " + ", ".join(missing_ids))
            if unknown_ids: errors.append("Gap Analysis contains unknown Question IDs: " + ", ".join(unknown_ids))
            if len(ids) != len(set(ids)): errors.append("Gap Analysis contains duplicate Question IDs")
            gap_rows = len(ids)
    if errors:
        raise ValueError("Standardized workbook validation failed:\n- " + "\n- ".join(errors))
    return {"valid": True, "inventory_tools": len(tool_names), "profiling_systems": len(profile_names), "gap_questions": gap_rows, "warnings": warnings}


def import_phase1_excel(file_data):
    workbook = workbook_from_upload(file_data)
    aliases = {"system": "inventory_name", "tool_system_name": "inventory_name", "tool_name": "inventory_name", "system_name": "inventory_name", "type_of_system": "system_type", "surveillance_functions_select_all": "surveillance_functions", "information_user": "information_users", "linked_systems": "linked_to_other_systems", "linked_to_other_systems": "linked_to_other_systems", "overall_comments": "overall_comments"}
    allowed = {"inventory_name", "unit_responsible", "system_type", "surveillance_functions", "geographical_coverage", "point_of_data_entry", "data_captured", "information_users", "technology", "has_api", "linked_to_other_systems", "comments"}
    existing_phase1 = profile_state().get("phase1", {})
    phase1 = {"country_name": existing_phase1.get("country_name", ""), "reporting_period": existing_phase1.get("reporting_period", ""), "comments": existing_phase1.get("comments", ""), "tools": []}
    sheet, header_row, headers = None, None, None
    for candidate in workbook.worksheets:
        for row_index, row in enumerate(candidate.iter_rows(min_row=1, max_row=min(20, candidate.max_row), values_only=True), 1):
            candidate_headers = [aliases.get(normalized_key(value), normalized_key(value)) for value in row]
            if "inventory_name" in candidate_headers:
                sheet, header_row, headers = candidate, row_index, candidate_headers; break
        if sheet is not None: break
    if sheet is None:
        seen = set()
        for candidate in workbook.worksheets:
            ids = {
                str(candidate.cell(row, 2).value or "").strip().upper()
                for row in range(2, min(candidate.max_row, 100) + 1)
            }
            headers_found = {normalized_key(cell.value) for cell in candidate[1]}
            if len(ids & set(PROFILE_ID_MAP)) < 60 or not {"question", "response"}.issubset(headers_found):
                continue
            name = next((
                str(candidate.cell(row, 6).value or "").strip()
                for row in range(2, min(candidate.max_row, 100) + 1)
                if str(candidate.cell(row, 2).value or "").strip().upper() == "A1"
            ), "")
            if name and name.casefold() not in seen:
                seen.add(name.casefold())
                phase1["tools"].append({"inventory_name": name})
        if not phase1["tools"]:
            raise ValueError("No Inventory table or completed Profiling system tabs were found in the workbook")
        return save_profile_phase("phase1", phase1)
    # Row 1 may contain the Inventory headers, in which case B1 and D1 are
    # labels rather than Member State and reporting-period metadata.
    if sheet.title.strip().casefold() == "his inventory" and header_row > 1:
        phase1["country_name"] = str(sheet["B1"].value or phase1["country_name"]).strip()
        phase1["reporting_period"] = str(sheet["D1"].value or phase1["reporting_period"]).strip()
    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        if not any(value not in (None, "") for value in row): continue
        tool = {}
        for key, value in zip(headers, row):
            if value in (None, ""): continue
            if key in {"country", "country_name"}: phase1["country_name"] = str(value).strip()
            elif key == "reporting_period": phase1["reporting_period"] = str(value).strip()
            elif key == "overall_comments": phase1["comments"] = str(value).strip()
            elif key in allowed:
                text = str(value).strip()
                # A slash can be part of a single valid label (for example,
                # "IDSR/weekly syndromic surveillance"), so only split actual
                # list separators here.
                tool[key] = [x.strip() for x in re.split(r"[;,\n]+", text) if x.strip()] if key in {"system_type", "surveillance_functions"} else text
        if any(tool.values()): phase1["tools"].append(tool)
    if not phase1["tools"]:
        seen = set()
        for candidate in workbook.worksheets:
            ids = {str(candidate.cell(row, 2).value or "").strip().upper() for row in range(2, min(candidate.max_row, 80) + 1)}
            if len(ids & set(PROFILE_ID_MAP)) < 60: continue
            name = next((str(candidate.cell(row, 6).value or "").strip() for row in range(2, min(candidate.max_row, 80) + 1) if str(candidate.cell(row, 2).value or "").strip().upper() == "A1"), "")
            if name and name.casefold() not in seen:
                seen.add(name.casefold()); phase1["tools"].append({"inventory_name": name})
    if not phase1["tools"]:
        raise ValueError("The standard Inventory sheet and profiling tabs contain no named tools")
    return save_profile_phase("phase1", phase1)


def import_phase2_excel(file_data):
    workbook = workbook_from_upload(file_data)
    phase1, assessments = parse_phase2_workbook(workbook, profile_state()["phase1"])
    if not assessments:
        raise ValueError("No completed Profiling tabs matched a system in the imported Inventory")
    save_profile_phase("phase1", phase1)
    return save_profile_phase("phase2", assessments)


def _set_excel_value(cell, value):
    """Write assessment text literally, without Excel formula interpretation."""
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def _append_excel_values(sheet, values):
    # These generated rows contain labels/data only, never template formulas.
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def export_workbook():
    if openpyxl is None:
        raise RuntimeError("openpyxl is required for Excel export")
    state = current_state()
    report = report_data()
    workbook = openpyxl.Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Gap Analysis — Gaps and Recommendations"])
    for section in report["sections"]:
        summary.append([])
        _append_excel_values(summary, [section["domain"], section["group"]])
        summary.append(["Gaps", "Recommendations"])
        length = max(len(section["gaps"]), len(section["recommendations"]), 1)
        for i in range(length):
            _append_excel_values(summary, [section["gaps"][i] if i < len(section["gaps"]) else "", section["recommendations"][i] if i < len(section["recommendations"]) else ""])
    detail = workbook.create_sheet("Responses")
    detail.append(["Domain", "Question ID", "Question", "Group", "Response", "Explanation", "Comment / Gap / Recommendation", "Updated"])
    response_map = {item["question_id"]: item for item in state["responses"]}
    assignment_map = {item["domain_id"]: item for item in state["assignments"]}
    for domain in DOMAINS:
        for question in domain["questions"]:
            response = response_map.get(question["id"], {})
            _append_excel_values(detail, [domain["name"], question.get("source_id", question["id"]).upper(), question["text"], assignment_map.get(domain["id"], {}).get("group_name", ""), response.get("response", ""), response.get("explanation", ""), response.get("comment", ""), response.get("updated_at", "")])
    profile = profile_state()
    phase1 = workbook.create_sheet("Inventory")
    _append_excel_values(phase1, ["Country", profile["phase1"].get("country_name", ""), "Reporting Period", profile["phase1"].get("reporting_period", "")])
    tool_headers = ["Tool / System Name", "Unit Responsible", "Type of System", "Surveillance Functions", "Geographical Coverage", "Point of Data Entry", "Data Captured", "Information Users", "Technology", "Has API", "Linked Systems", "Comments"]
    tool_keys = ["inventory_name", "unit_responsible", "system_type", "surveillance_functions", "geographical_coverage", "point_of_data_entry", "data_captured", "information_users", "technology", "has_api", "linked_to_other_systems", "comments"]
    phase1.append(tool_headers)
    for tool in profile["phase1"].get("tools", []):
        _append_excel_values(phase1, ["; ".join(tool.get(key, [])) if isinstance(tool.get(key), list) else tool.get(key, "") for key in tool_keys])
    phase2 = workbook.create_sheet("Profiling")
    assessments = profile.get("phase2", [])
    phase2_keys = list(dict.fromkeys(key for item in assessments for key in item.keys()))
    _append_excel_values(phase2, [key.replace("_", " ").title() for key in phase2_keys])
    for assessment in assessments:
        _append_excel_values(phase2, ["; ".join(assessment.get(key, [])) if isinstance(assessment.get(key), list) else assessment.get(key, "") for key in phase2_keys])
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
            cell.fill = openpyxl.styles.PatternFill("solid", fgColor="176B45")
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(60, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
            for cell in column:
                cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_profile_workbook(phase):
    profile = profile_state()
    workbook = openpyxl.Workbook()
    if phase == 1:
        sheet = workbook.active; sheet.title = "Inventory"
        headers = ["Country", "Reporting Period", "Tool / System Name", "Unit Responsible", "Type of System", "Surveillance Functions", "Geographical Coverage", "Point of Data Entry", "Data Captured", "Information Users", "Technology", "Has API", "Linked to Other Systems", "Comments"]
        sheet.append(headers)
        keys = ["inventory_name", "unit_responsible", "system_type", "surveillance_functions", "geographical_coverage", "point_of_data_entry", "data_captured", "information_users", "technology", "has_api", "linked_to_other_systems", "comments"]
        for tool in profile["phase1"].get("tools", []):
            values = ["; ".join(tool.get(k, [])) if isinstance(tool.get(k), list) else tool.get(k, "") for k in keys]
            _append_excel_values(sheet, [profile["phase1"].get("country_name", ""), profile["phase1"].get("reporting_period", ""), *values])
    else:
        assessments = profile.get("phase2", [])
        for index, assessment in enumerate(assessments or [{}]):
            sheet = workbook.active if index == 0 else workbook.create_sheet()
            sheet.title = re.sub(r"[\\*?:/\[\]]", "-", str(assessment.get("official_name") or f"Assessment {index+1}"))[:31]
            sheet.append(["index", "Section", "#", "Question", "Response Type", "Expected Responses / Options", "Response", "Comments"])
            for row_index, (key, value) in enumerate(assessment.items()):
                if key == "official_name": continue
                is_multi = isinstance(value, list)
                _append_excel_values(sheet, [row_index, "", "", key.replace("_", " ").title(), "Multi-select" if is_multi else "Text", "", ", ".join(value) if is_multi else value, ""])
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF"); cell.fill = openpyxl.styles.PatternFill("solid", fgColor="176B45")
        for column in sheet.columns:
            letter = column[0].column_letter; sheet.column_dimensions[letter].width = min(55, max(12, max(len(str(c.value or "")) for c in column) + 2))
    output = io.BytesIO(); workbook.save(output); return output.getvalue()


def export_phase3_workbook():
    if openpyxl is None:
        raise RuntimeError("openpyxl is required for Excel export")
    state, summary = current_state(), report_data()
    workbook = openpyxl.Workbook()
    overview = workbook.active; overview.title = "Gap Analysis Summary"
    _append_excel_values(overview, ["Gap Analysis", state["assessment"].get("title", "")])
    _append_excel_values(overview, ["Scope", state["assessment"].get("scope", "")])
    overview.append([]); overview.append(["Domain", "Participant Group", "Gaps / Other Findings", "Recommendations"])
    for item in summary["sections"]:
        _append_excel_values(overview, [item["domain"], item["group"], "\n".join(item["gaps"] + item["other"]), "\n".join(item["recommendations"])])
    details = workbook.create_sheet("Gap Analysis Responses")
    details.append(["Domain", "Question ID", "Question", "Participant Group", "Response", "Explanation", "Comment / Gap / Recommendation", "Last Updated"])
    assignments = {item["domain_id"]: item for item in state["assignments"]}
    responses = {item["question_id"]: item for item in state["responses"]}
    for domain in DOMAINS:
        group = assignments.get(domain["id"], {}).get("group_name", "Unassigned")
        for question in domain["questions"]:
            response = responses.get(question["id"], {})
            value = response.get("response", "")
            try:
                parsed = json.loads(value)
                if isinstance(parsed, list): value = "; ".join(map(str, parsed))
            except (TypeError, ValueError): pass
            _append_excel_values(details, [domain["name"], question.get("source_id", question["id"]).upper(), question["text"], group, value, response.get("explanation", ""), response.get("comment", ""), response.get("updated_at", "")])
    for sheet in workbook.worksheets:
        header_row = 4 if sheet.title == "Gap Analysis Summary" else 1
        sheet.freeze_panes = f"A{header_row + 1}"
        for cell in sheet[header_row]:
            cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
            cell.fill = openpyxl.styles.PatternFill("solid", fgColor="176B45")
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(55, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
            for cell in column:
                cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
    output = io.BytesIO(); workbook.save(output); return output.getvalue()


def _safe_sheet_title(value, fallback):
    cleaned = re.sub(r"[\\*?:/\[\]]", "-", str(value or fallback)).strip() or fallback
    return cleaned[:31]


def _style_roundtrip_workbook(workbook, header_rows):
    for sheet in workbook.worksheets:
        header_row = header_rows.get(sheet.title)
        if not header_row:
            continue
        sheet.freeze_panes = f"A{header_row + 1}"
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(sheet.max_column)}{sheet.max_row}"
        for cell in sheet[header_row]:
            cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
            cell.fill = openpyxl.styles.PatternFill("solid", fgColor="176B45")
            cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="center")
        for row in sheet.iter_rows(min_row=header_row + 1):
            for cell in row:
                cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(55, max(12, max(len(str(cell.value or "")) for cell in column) + 2))


def roundtrip_workbook(phases=(1, 2, 3)):
    """Build Excel files in the exact layouts consumed by the web importers."""
    if openpyxl is None:
        raise RuntimeError("openpyxl is required for Excel export")
    profile, state, summary = profile_state(), current_state(), report_data()
    workbook = openpyxl.Workbook(); workbook.remove(workbook.active)
    header_rows = {}
    if 1 in phases:
        sheet = workbook.create_sheet("Inventory")
        headers = ["Country", "Reporting Period", "Overall Comments", "Tool / System Name", "Unit Responsible", "Type of System", "Surveillance Functions", "Geographical Coverage", "Point of Data Entry", "Data Captured", "Information Users", "Technology", "Has API", "Linked to Other Systems", "Comments"]
        keys = ["inventory_name", "unit_responsible", "system_type", "surveillance_functions", "geographical_coverage", "point_of_data_entry", "data_captured", "information_users", "technology", "has_api", "linked_to_other_systems", "comments"]
        sheet.append(headers); header_rows[sheet.title] = 1
        tools = profile["phase1"].get("tools", []) or [{}]
        for tool in tools:
            values = ["; ".join(map(str, tool.get(key, []))) if isinstance(tool.get(key), list) else tool.get(key, "") for key in keys]
            _append_excel_values(sheet, [profile["phase1"].get("country_name", ""), profile["phase1"].get("reporting_period", ""), profile["phase1"].get("comments", ""), *values])
    if 2 in phases:
        saved_assessments = profile.get("phase2", [])
        inventory_tools = [tool for tool in profile["phase1"].get("tools", []) if str(tool.get("inventory_name") or "").strip()]
        assessments = []
        for tool in inventory_tools:
            assessment = next((item for item in saved_assessments if item.get("_tool_id") == tool.get("_id")), None)
            if assessment is None:
                assessment = next((item for item in saved_assessments if str(item.get("official_name") or "").strip().casefold() == str(tool.get("inventory_name") or "").strip().casefold()), None)
            assessment = dict(assessment or {})
            assessment["official_name"] = assessment.get("official_name") or tool.get("inventory_name", "")
            assessment["_tool_id"] = tool.get("_id", "")
            assessments.append(assessment)
        if not assessments:
            assessments = [dict(item) for item in saved_assessments] or [{"official_name": ""}]
        used_titles = set()
        for index, assessment in enumerate(assessments, 1):
            base = _safe_sheet_title(f"SP_{index}_{assessment.get('official_name') or 'Assessment'}", f"SP_{index}_Assessment")
            title = base; suffix = 2
            while title.casefold() in used_titles:
                title = _safe_sheet_title(f"{base[:27]}_{suffix}", f"SP_{index}_{suffix}"); suffix += 1
            used_titles.add(title.casefold()); sheet = workbook.create_sheet(title)
            sheet.append(["Index", "Section", "#", "Question", "Response Type", "Expected Responses / Options", "Response", "Comments"]); header_rows[title] = 1
            for row_index, (source_id, key) in enumerate(PROFILE_ID_MAP.items(), 1):
                value = assessment.get(key, "")
                is_multi = key in PROFILE_MULTI_FIELDS
                comments = (assessment.get("field_comments") or {}).get(key, "")
                if key == "data_exchange" and not comments:
                    comments = assessment.get("exchange_comments", "")
                _append_excel_values(sheet, [row_index, "", source_id, key.replace("_", " ").title(), "Multi-select" if is_multi else "Text", "", "; ".join(map(str, value)) if isinstance(value, list) else value, comments])
    if 3 in phases:
        sheet = workbook.create_sheet("Gap Analysis Responses")
        sheet.append(["Domain", "Question ID", "Question", "Participant Group", "Response", "Explanation", "Comments", "Last Updated"]); header_rows[sheet.title] = 1
        assignments = {item["domain_id"]: item for item in state["assignments"]}; responses = {item["question_id"]: item for item in state["responses"]}
        for domain in DOMAINS:
            group = assignments.get(domain["id"], {}).get("group_name", "")
            for question in domain["questions"]:
                response = responses.get(question["id"], {}); value = response.get("response", "")
                try:
                    parsed = json.loads(value)
                    if isinstance(parsed, list): value = "; ".join(map(str, parsed))
                except (TypeError, ValueError): pass
                _append_excel_values(sheet, [domain["name"], question.get("source_id", question["id"]).upper(), question["text"], group, value, response.get("explanation", ""), response.get("comment", ""), response.get("updated_at", "")])
        summary_sheet = workbook.create_sheet("Gap Analysis Summary")
        summary_sheet.append(["Domain", "Participant Group", "Gaps / Other Findings", "Recommendations"]); header_rows[summary_sheet.title] = 1
        for item in summary["sections"]:
            _append_excel_values(summary_sheet, [item["domain"], item["group"], "\n".join(item["gaps"] + item["other"]), "\n".join(item["recommendations"])])
    _style_roundtrip_workbook(workbook, header_rows)
    output = io.BytesIO(); workbook.save(output); return output.getvalue()


def _display_excel_value(value):
    if isinstance(value, list):
        return "; ".join(map(str, value))
    return value if value not in (None, "") else ""


def _split_gap_comment(value):
    gaps, recommendations = [], []
    for raw_line in str(value or "").splitlines():
        line = raw_line.strip()
        if not line: continue
        if re.match(r"^recommendations?\s*:", line, re.I): recommendations.append(re.sub(r"^recommendations?\s*:\s*", "", line, flags=re.I))
        elif re.match(r"^gaps?\s*:", line, re.I): gaps.append(re.sub(r"^gaps?\s*:\s*", "", line, flags=re.I))
        else: gaps.append(line)
    return "\n".join(gaps), "\n".join(recommendations)


def standard_format_workbook(phases=(1, 2, 3)):
    """Export using the supplied Africa CDC standard workbook layout."""
    if openpyxl is None:
        raise RuntimeError("openpyxl is required for Excel export")
    if not STANDARD_TEMPLATE_PATH.exists():
        raise RuntimeError("The standard Excel template is missing")
    phases = set(phases)
    workbook = openpyxl.load_workbook(STANDARD_TEMPLATE_PATH)
    profile, state = profile_state(), current_state()
    inventory = workbook["HIS Inventory"]
    template = workbook["System Profilling"]
    standard_gap_names = {name for name in STANDARD_GAP_SHEETS}

    # Remove the filled example tool tabs supplied in the reference workbook.
    for sheet in list(workbook.worksheets):
        if sheet.title not in {"HIS Inventory", "System Profilling"} and sheet.title.strip().casefold() not in standard_gap_names:
            workbook.remove(sheet)

    if 1 in phases:
        _set_excel_value(inventory["B1"], profile["phase1"].get("country_name", ""))
        _set_excel_value(inventory["D1"], profile["phase1"].get("reporting_period", ""))
        for row in inventory.iter_rows(min_row=3, max_row=inventory.max_row, min_col=1, max_col=12):
            for cell in row: cell.value = None
        keys = ["inventory_name", "unit_responsible", "system_type", "surveillance_functions", "geographical_coverage", "point_of_data_entry", "data_captured", "information_users", "technology", "has_api", "linked_to_other_systems", "comments"]
        for row_index, tool in enumerate(profile["phase1"].get("tools", []), 3):
            for column, key in enumerate(keys, 1): _set_excel_value(inventory.cell(row_index, column), _display_excel_value(tool.get(key)))
    else:
        workbook.remove(inventory)

    if 2 in phases:
        saved_profiles = profile.get("phase2", [])
        tools = [tool for tool in profile["phase1"].get("tools", []) if str(tool.get("inventory_name") or "").strip()]
        if not tools: tools = [{"_id": "", "inventory_name": "tool 1"}]
        used = set(workbook.sheetnames)
        for index, tool in enumerate(tools, 1):
            assessment = next((item for item in saved_profiles if item.get("_tool_id") == tool.get("_id")), {})
            title = _safe_sheet_title(f"tool {index} - {tool.get('inventory_name', '')}", f"tool {index}")
            base, suffix = title, 2
            while title in used:
                title = _safe_sheet_title(f"{base[:27]} {suffix}", f"tool {index}-{suffix}"); suffix += 1
            used.add(title); sheet = workbook.copy_worksheet(template); sheet.title = title; sheet.sheet_state = "visible"
            for row_index in range(2, 71):
                source_id = str(sheet.cell(row_index, 2).value or "").strip().upper()
                key = PROFILE_ID_MAP.get(source_id)
                if not key: continue
                value = assessment.get(key, tool.get("inventory_name", "") if source_id == "A1" else "")
                _set_excel_value(sheet.cell(row_index, 6), _display_excel_value(value))
                comments = (assessment.get("field_comments") or {}).get(key, "")
                if source_id == "G7" and not comments:
                    comments = assessment.get("exchange_comments", "")
                _set_excel_value(sheet.cell(row_index, 7), comments)
        template.sheet_state = "hidden"
    else:
        workbook.remove(template)

    if 3 in phases:
        assignments = {item["domain_id"]: item for item in state["assignments"]}
        responses = {item["question_id"]: item for item in state["responses"]}
        source_to_question = {str(question.get("source_id", question["id"])).upper(): (domain["id"], question["id"]) for domain in DOMAINS for question in domain["questions"]}
        for sheet_name, (allowed_domains, _) in STANDARD_GAP_SHEETS.items():
            sheet = next(item for item in workbook.worksheets if item.title.strip().casefold() == sheet_name)
            for row_index in range(1, min(sheet.max_row, 150) + 1):
                source_id = str(sheet.cell(row_index, 1).value or "").strip().upper()
                mapped = source_to_question.get(source_id)
                if not mapped: continue
                for column in (5, 6, 7): sheet.cell(row_index, column).value = None
                domain_id, question_id = mapped
                if domain_id not in allowed_domains: continue
                response = responses.get(question_id, {})
                value = response.get("response", "")
                try:
                    parsed = json.loads(value)
                    if isinstance(parsed, list): value = "; ".join(map(str, parsed))
                except (TypeError, ValueError): pass
                explanation = str(response.get("explanation") or "").strip()
                if explanation: value = f"{value}\nExplanation: {explanation}".strip()
                gap, recommendation = _split_gap_comment(response.get("comment", ""))
                _set_excel_value(sheet.cell(row_index, 5), value); _set_excel_value(sheet.cell(row_index, 6), gap); _set_excel_value(sheet.cell(row_index, 7), recommendation)
    else:
        for sheet in list(workbook.worksheets):
            if sheet.title.strip().casefold() in standard_gap_names: workbook.remove(sheet)

    output = io.BytesIO(); workbook.save(output); return output.getvalue()


REDCAP_PRIMARY_WRITE_PATHS = {
    "/api/profile/phase1", "/api/profile/phase1/tool", "/api/profile/phase2", "/api/profile/phase2/tool",
    "/api/profile/import", "/api/import/phase1", "/api/import/phase2", "/api/import",
    "/api/groups", "/api/inventory/assign", "/api/inventory/unassign", "/api/profiling/assign",
    "/api/profiling/unassign", "/api/assessment", "/api/assign", "/api/unassign", "/api/reassign",
    "/api/responses", "/api/accounts/reset",
}


def workspace_snapshot():
    return {"profile": profile_state(), "state": current_state()}


def analysis_section_text(section):
    profile = profile_state()
    phase1 = profile.get("phase1") or {}
    full_text = complete_report_text()
    lines = full_text.splitlines()
    profiling_index = next((index for index, line in enumerate(lines) if line.strip() == "Profiling"), len(lines))
    gap_index = next((index for index, line in enumerate(lines) if line.strip().startswith("Gap Analysis:")), len(lines))
    metadata = [
        lines[0] if lines else "Africa CDC Surveillance Systems Assessment",
        f"Country: {phase1.get('country_name') or 'Not provided'}",
        f"Reporting Period: {phase1.get('reporting_period') or 'Not provided'}",
        "",
    ]
    if section == "inventory":
        return "\n".join(lines[:profiling_index]).strip()
    if section == "profiling":
        return "\n".join(metadata + lines[profiling_index:gap_index]).strip()
    gap_lines = report_text().splitlines()
    if gap_lines:
        gap_lines = gap_lines[1:]
    return "\n".join(metadata + gap_lines).strip()


def analysis_assistant_context(section, profile, state):
    evidence = []
    # The Unified DHIS2 comparison is available from every report section, so
    # always include the named Inventory records.
    for index, tool in enumerate((profile.get("phase1") or {}).get("tools") or [], 1):
        values = {key.replace("_", " ").title(): value for key, value in tool.items() if not key.startswith("_") and value not in (None, "", [])}
        name = str(tool.get("inventory_name") or f"Inventory system {index}").strip()
        if values: evidence.append({"record": f"Inventory: {name}", "evidence_type": "inventory", "values": values})
    if section == "profiling":
        for index, item in enumerate(profile.get("phase2") or [], 1):
            values = {key.replace("_", " ").title(): value for key, value in item.items() if not key.startswith("_") and value not in (None, "", [])}
            name = str(item.get("official_name") or f"Profiling system {index}").strip()
            if values: evidence.append({"record": f"Profiling: {name}", "evidence_type": "profiling", "values": values})
    elif section == "gap":
        responses = {str(item.get("question_id")): item for item in state.get("responses") or []}
        for domain in state.get("domains") or []:
            for question in domain.get("questions") or []:
                saved = responses.get(str(question.get("id")), {})
                values = {key: saved.get(key) for key in ("response", "explanation", "comment") if saved.get(key) not in (None, "", [])}
                if values:
                    evidence.append({"domain": domain.get("name"), "question": question.get("text"), **values})
    return json.dumps(evidence, ensure_ascii=False, separators=(",", ":"))


def ollama_executable():
    candidates = [
        shutil.which("ollama"),
        str(Path(os.environ.get("LOCALAPPDATA", "")) / "Programs" / "Ollama" / "ollama.exe"),
        str(Path(os.environ.get("PROGRAMFILES", "")) / "Ollama" / "ollama.exe"),
    ]
    return next((item for item in candidates if item and Path(item).is_file()), None)


def keep_display_awake_for_local_ai():
    if os.name != "nt":
        return False
    try:
        import ctypes
        flags = 0x80000000 | 0x00000001 | 0x00000002
        return bool(ctypes.windll.kernel32.SetThreadExecutionState(flags))
    except (AttributeError, OSError):
        return False


def restore_normal_power_state(local_ai_guard_active):
    if not local_ai_guard_active or os.name != "nt":
        return
    try:
        import ctypes
        ctypes.windll.kernel32.SetThreadExecutionState(0x80000000)
    except (AttributeError, OSError):
        pass


def ensure_ollama_service():
    parsed = urlparse(OLLAMA_URL)
    if parsed.hostname not in {"127.0.0.1", "localhost", "::1"}:
        return
    try:
        with urlopen(Request(f"{OLLAMA_URL}/api/tags", method="GET"), timeout=2):
            return
    except Exception:
        pass
    with OLLAMA_START_LOCK:
        try:
            with urlopen(Request(f"{OLLAMA_URL}/api/tags", method="GET"), timeout=2):
                return
        except Exception:
            pass
        executable = ollama_executable()
        if not executable:
            return
        kwargs = {"stdin": subprocess.DEVNULL, "stdout": subprocess.DEVNULL, "stderr": subprocess.DEVNULL}
        if os.name == "nt":
            kwargs["creationflags"] = subprocess.CREATE_NO_WINDOW | subprocess.DETACHED_PROCESS
        subprocess.Popen([executable, "serve"], **kwargs)
        for _ in range(16):
            time.sleep(0.5)
            try:
                with urlopen(Request(f"{OLLAMA_URL}/api/tags", method="GET"), timeout=2):
                    return
            except Exception:
                continue


def ollama_status():
    ensure_ollama_service()
    try:
        request = Request(f"{OLLAMA_URL}/api/tags", method="GET")
        with urlopen(request, timeout=2) as response:
            models = json.loads(response.read().decode("utf-8")).get("models", [])
        names = {str(item.get("name") or item.get("model") or "") for item in models}
        # Chat uses the configured tag verbatim; another tag in the same family
        # does not prove this exact model is installed.
        available = OLLAMA_MODEL in names
        return {"running": True, "available": available, "model": OLLAMA_MODEL}
    except Exception:
        return {"running": False, "available": False, "model": OLLAMA_MODEL}


def assistant_provider():
    provider = app_setting("assistant_provider", "local").strip().lower()
    return provider if provider in {"local", "openai"} else "local"


def assistant_status():
    provider = assistant_provider()
    if provider == "openai":
        return {"provider": "openai", "available": bool(OPENAI_API_KEY), "configured": bool(OPENAI_API_KEY), "model": OPENAI_MODEL, "online": True}
    status = ollama_status()
    return {**status, "provider": "local", "configured": True, "online": False}


def configure_assistant_provider(provider, approved=False):
    provider = str(provider or "local").strip().lower()
    if provider not in {"local", "openai"}:
        raise ValueError("Select Local Ollama or OpenAI")
    if provider == "openai":
        if not OPENAI_API_KEY:
            raise ValueError("AFRICA_CDC_OPENAI_API_KEY is not configured in .env")
        if not approved:
            raise ValueError("Confirm that online processing is authorised before enabling OpenAI")
    save_app_setting("assistant_provider", provider)
    return assistant_status()


def recommendation_theme_instructions(question):
    if not re.search(r"recommend|consolidat|group.*theme|coverage check", str(question), re.I):
        return ""
    return (
        "\nRECOMMENDATIONS: Group related issues into up to eight evidence-supported themes: "
        "Governance and coordination; Sustainable financing; Workforce and technical support; "
        "Infrastructure and reporting continuity; System integration and transition; Surveillance "
        "configuration and data quality; Analysis, dissemination and feedback; Data access, protection "
        "and recovery. Omit unsupported themes. Under each provide a consolidated finding with exact "
        "source references, one Proposed recommendation, two or three specific practical actions and "
        "a Proposed monitoring measure. Preserve differences between systems. Distinguish respondent "
        "suggestions from established findings. Do not invent owners, budgets, targets or deadlines. "
        "Do not reproduce the full questionnaire or a long numbered recommendation list. Keep a short "
        "Unresolved issues section for unmatched, contradictory or insufficient evidence. Do not claim "
        "complete coverage from truncated source material. Only when a coverage check is requested, "
        "map each supplied original recommendation to its theme and state combined, needs validation "
        "or unresolved; a keyword match alone does not establish that the action addresses every detail.\n"
    )


def openai_report_answer(payload):
    if not OPENAI_API_KEY:
        raise RuntimeError("OpenAI is not configured")
    report = str(payload.get("report") or "").strip()[:60000]
    question = str(payload.get("question") or "").strip()[:2000]
    if not report or not question:
        raise ValueError("Generate a report and enter a question first")
    instructions = (
        "You are an Africa CDC public-health surveillance report analyst. Use only the generated report supplied in "
        "this request. Do not use outside facts, browse the web, invent evidence, or imply access to REDCap. Preserve "
        "recorded system names, figures, denominators, country, reporting period, and qualifications. Distinguish "
        "recorded evidence, interpretation, and proposed action. State evidence limitations and answer directly."
    )
    body = json.dumps({
        "model": OPENAI_MODEL,
        "store": False,
        "max_output_tokens": 3500,
        "input": [
            {"role": "system", "content": instructions + recommendation_theme_instructions(question)},
            {"role": "user", "content": f"GENERATED REPORT:\n{report}\n\nREQUEST:\n{question}"},
        ],
    }).encode("utf-8")
    request = Request(OPENAI_RESPONSES_URL, data=body, headers={"Content-Type": "application/json", "Authorization": f"Bearer {OPENAI_API_KEY}"}, method="POST")
    try:
        with urlopen(request, timeout=300) as response:
            result = json.loads(response.read().decode("utf-8"))
    except HTTPError as exc:
        detail = exc.read().decode("utf-8", "replace")
        try: detail = json.loads(detail).get("error", {}).get("message") or detail
        except Exception: pass
        raise RuntimeError(f"OpenAI rejected the request: {detail}") from exc
    except (URLError, TimeoutError, OSError) as exc:
        raise RuntimeError("The online AI provider is unavailable") from exc
    answer = str(result.get("output_text") or "").strip()
    if not answer:
        answer = "".join(
            str(content.get("text") or "")
            for item in result.get("output") or [] if item.get("type") == "message"
            for content in item.get("content") or [] if content.get("type") == "output_text"
        ).strip()
    if not answer:
        raise RuntimeError("OpenAI returned no report analysis")
    return answer


def stream_online_report_chat(handler, payload):
    answer = openai_report_answer(payload)
    handler.send_response(200)
    handler.send_header("Content-Type", "application/x-ndjson; charset=utf-8")
    handler.send_header("Cache-Control", "no-store")
    handler.end_headers()
    packet = {"delta": answer, "done": True, "model": OPENAI_MODEL, "provider": "openai"}
    handler.wfile.write((json.dumps(packet, ensure_ascii=False) + "\n").encode("utf-8")); handler.wfile.flush()


def compact_thematic_evidence(raw_evidence, total_limit=9000):
    """Keep coverage of every record while fitting small local-model contexts."""
    try:
        records = json.loads(raw_evidence)
    except (TypeError, json.JSONDecodeError):
        return str(raw_evidence or "")[:total_limit]
    if not isinstance(records, list):
        return str(raw_evidence or "")[:total_limit]

    def clipped(value, limit):
        text = re.sub(r"\s+", " ", str(value or "")).strip()
        if len(text) <= limit:
            return text
        cut = text[:limit].rsplit(" ", 1)[0].rstrip(" ,;:")
        return cut + "…"

    valid_records = [record for record in records if isinstance(record, dict)]
    # Divide the available prompt space across all records so late domains are
    # never silently dropped when comments are long.
    per_record = max(70, (total_limit - len(valid_records)) // max(1, len(valid_records)))
    lines = []
    for index, record in enumerate(valid_records, 1):
        present_fields = sum(bool(record.get(key)) for key in ("response", "explanation", "comment"))
        label_length = sum(
            len(label) + 3 for key, label in (("response", "Finding: "), ("explanation", "Explanation: "), ("comment", "Gap/action: "))
            if record.get(key)
        )
        domain_budget = min(40, max(12, per_record // 4))
        fixed_length = len(str(index)) + domain_budget + label_length + 3
        detail_budget = max(12 * max(1, present_fields), per_record - fixed_length)
        parts = [f"[{index}] {clipped(record.get('domain'), domain_budget)}"]
        if record.get("response"):
            parts.append("Finding: " + clipped(record["response"], max(8, int(detail_budget * 0.42))))
        if record.get("explanation"):
            parts.append("Explanation: " + clipped(record["explanation"], max(8, int(detail_budget * 0.22))))
        if record.get("comment"):
            parts.append("Gap/action: " + clipped(record["comment"], max(8, int(detail_budget * 0.36))))
        lines.append(" | ".join(parts))
    final_budget = max(1, (total_limit - len(lines) + 1) // max(1, len(lines)))
    return "\n".join(line[:final_budget] for line in lines)


def compact_recommendation_material(report, evidence, total_limit=9000):
    """Select recommendation-relevant evidence without sending the full report to a CPU-only model."""
    terms = re.compile(
        r"recommend|finding|gap|risk|challenge|limitation|priority|action|phase.?out|api|exchange|interop|"
        r"offline|policy|govern|ownership|security|training|capacity|fund|coverage|laborator|contact|specimen|dhis2",
        re.I,
    )
    selected = []
    seen = set()
    for raw_line in str(report or "").splitlines():
        line = re.sub(r"\s+", " ", raw_line).strip()
        if not line or (not terms.search(line) and not re.match(r"^(Overview|Profiling|Gap Analysis|Summary)", line, re.I)):
            continue
        key = line.casefold()
        if key not in seen:
            seen.add(key)
            selected.append(line)
    report_text = "\n".join(selected)
    evidence_text = compact_thematic_evidence(evidence, total_limit=1800) if evidence else ""
    combined = f"RELEVANT REPORT FINDINGS:\n{report_text}\n\nCOMPACT SUPPORTING EVIDENCE:\n{evidence_text}"
    return combined[:total_limit] + ("\nSOURCE LIMITATION: Evidence was shortened; complete coverage cannot be verified." if len(combined) > total_limit else "")


def assistant_numeric_grounding_warning(answer, source_material):
    answer = str(answer or "")
    source = re.sub(r"\s+", " ", str(source_material or "")).casefold()
    expressions = set()
    for pattern in (r"\b\d+(?:\.\d+)?\s*%", r"\b\d+\s+(?:of|out of)\s+\d+\b"):
        expressions.update(re.sub(r"\s+", " ", item).strip() for item in re.findall(pattern, answer, re.I))
    unsupported = sorted(item for item in expressions if item.casefold() not in source)
    if not unsupported:
        return ""
    return (
        "\n\nEvidence check warning: The following numerical expressions were not found verbatim in the supplied "
        f"report evidence and must be independently verified before use: {', '.join(unsupported)}."
    )


def ollama_report_chat(payload):
    ensure_ollama_service()
    report = str(payload.get("report") or "").strip()[:60000]
    evidence = str(payload.get("evidence") or "").strip()[:100000]
    question = str(payload.get("question") or "").strip()[:2000]
    if not report or not question:
        raise ValueError("Generate a report and enter a question first")
    history = []
    for item in list(payload.get("history") or [])[-6:]:
        role = str(item.get("role") or "")
        content = str(item.get("content") or "").strip()[:3000]
        if role in {"user", "assistant"} and content:
            history.append({"role": role, "content": content})
    system = (
        "You are a senior public-health surveillance and digital-health assessment report specialist supporting "
        "Africa CDC. Work exclusively from the GENERATED REPORT below. Never use outside facts, invent evidence, "
        "change scores, or imply that an action was recorded when it was inferred. If evidence is missing, state "
        "that limitation. Treat Inventory, Profiling, and Gap Analysis as distinct assessment components.\n\n"
        "PROFESSIONAL REPORTING STANDARD\n"
        "1. Use precise, neutral, decision-oriented language suitable for ministries of health and partners.\n"
        "2. Synthesize related evidence instead of repeating questionnaire responses system by system.\n"
        "3. Preserve system names, institutions, figures, denominators, periods, and qualifications exactly.\n"
        "4. Separate: recorded evidence; analytical interpretation; proposed recommendation.\n"
        "5. Link every recommendation to a documented finding and make it specific, feasible, and actionable.\n"
        "6. Do not claim causation, national representativeness, or implementation status unless recorded.\n"
        "7. When comparing systems or domains, explain the comparison basis and note incomplete data.\n\n"
        "TABLE AND EVIDENCE INTERPRETATION\n"
        "When the user asks for discussion, insights, interpretation, or a professional report, examine every table "
        "and substantive section in the supplied report. Do not merely restate rows. For each table: name the table "
        "or subject; summarize the strongest pattern; identify meaningful variation, outliers, strengths, gaps, and "
        "missing values; explain the operational or strategic implication; and give a traceable action where supported. "
        "Use exact figures and denominators when present. If a table has insufficient variation or evidence for an "
        "inference, say so. After the table-level discussion, synthesize cross-table relationships and tensions. "
        "Never describe a table that is not present. A request for 'all tables' requires a numbered subsection for "
        "every actual table in the report, followed by Overall Insights and Decision Implications. A denominator of "
        "applicable yes/no responses is a response count, not a number of systems, sites, or domains. Do not convert it "
        "into another unit. Never describe affirmative answers or scored yes/no answers as systems. Use only an explicit "
        "system-level count when stating how many systems report a capability, and distinguish the Inventory total from "
        "the Profiling denominator. Do not call a narrative list a table. Recommendations may add professional judgement, but "
        "must be labelled 'Proposed' and must not introduce a named site, institution, technology, standard, target, "
        "or implementation claim absent from the report.\n\n"
        "AVAILABLE ANALYSES\n"
        "- Professional report: Title, Executive Summary, Scope and Evidence Base, Key Findings, Discussion of every "
        "table and major result, Thematic Analysis, Cross-cutting Issues, Prioritized Recommendations, Action Plan, "
        "Limitations, and Conclusion. The discussion must add interpretation rather than copy the report.\n"
        "- Thematic analysis: identify recurring themes from the evidence; for each give supporting evidence, affected "
        "systems/domains, operational implication, and priority. Do not fabricate frequencies.\n"
        "- Qualitative thematic analysis of open-ended responses: read all supplied responses, explanations, and "
        "comments; code recurring ideas into themes and subthemes; retain important divergent or minority evidence; "
        "and distinguish respondent statements from your interpretation. For each theme provide: definition, evidence "
        "summary, affected systems/domains, prevalence stated only as an exact count when calculable, contrasting or "
        "negative cases, interpretation, operational implication, and proposed response. Then provide cross-cutting "
        "insights, relationships between themes, evidence limitations, and prioritized actions. Paraphrase evidence "
        "unless a very short quotation is necessary. Do not equate repetition with importance and do not fabricate "
        "participant intent.\n"
        "- Action plan: provide a readable table with Action, Evidence/Rationale, Priority, Suggested Lead, Timeframe, "
        "and Verifiable Indicator. Label leads and timeframes as proposed unless explicitly recorded.\n"
        "- Executive summary: concise context, strongest findings, implications, and highest-priority actions.\n"
        "- Recommendations: consolidate duplicates, order by urgency and dependency, and retain traceability to findings.\n"
        "- Unified DHIS2 integration analysis: the Unified DHIS2 Surveillance Toolkit is an added-value surveillance "
        "layer for DHIS2, not evidence that every toolkit function belongs to core DHIS2. The primary comparison is "
        "between each existing surveillance tool and the Unified DHIS2 Surveillance Toolkit. When the report or user "
        "supplies the toolkit's full functionality, treat it as the authoritative added-value capability baseline "
        "for the comparison. Do not discard it merely because the capability has not yet been documented "
        "as configured or used in the selected country. Construct a tool-level matrix with one row per unique existing "
        "tool and separate columns for: Tool; Recorded functions; Functions covered by the Unified DHIS2 "
        "Toolkit; Required functions not covered by toolkit; Country validation gap; and Proposed tool disposition. "
        "The proposed disposition must decide the future of the "
        "existing tool, not say how to configure DHIS2. "
        "A toolkit capability that is not evidenced as "
        "locally configured must be labelled 'Toolkit supports; local configuration not verified', never 'No DHIS2 "
        "capability' or 'retain externally' solely because local implementation evidence is absent. Attribute each "
        "capability to core DHIS2, the added-value toolkit, or both; never silently present a toolkit-added function "
        "as a native core-DHIS2 function. Conversely, do "
        "not claim that a generally known DHIS2 feature is in the supplied toolkit unless the user-provided toolkit "
        "baseline or report records it. For every relevant "
        "tool classify the proposed disposition as: Retire after validated replacement; Integrate with Unified "
        "DHIS2 and retain; Retain as an external specialist application; Retain temporarily pending validation; or "
        "Integrate with Unified DHIS2; or Evidence insufficient. Use Evidence insufficient when the assessment records "
        "no usable function or workflow evidence for the tool. When an existing tool "
        "is named DHIS2 or HMIS/DHIS2, treat Unified DHIS2 as enhancing that existing platform. Recommend retaining "
        "and enhancing it through Unified DHIS2, never retiring or replacing it merely because the toolkit covers its functions. Never assume a DHIS2 "
        "module, feature, API, interoperability standard, or country implementation that neither the supplied toolkit "
        "baseline nor the report documents. Identify functions available in another tool but absent from the supplied "
        "Unified DHIS2 baseline as functional-gap risks. When another tool provides a required function that the "
        "toolkit does not provide, explicitly identify the function and source tool as an integration requirement; "
        "if integration is not evidenced or feasible, propose retaining that specialist function externally pending "
        "validation. When both provide the function, assess duplication, workflow continuity, data migration, or "
        "authoritative-source decisions. When only the toolkit provides it, identify the configuration and adoption "
        "requirement. Identify documented toolkit functions that are not verified "
        "in-country as configuration or validation gaps rather than product-capability gaps. Consider workflow "
        "continuity, identifiers, data ownership, data exchange, offline "
        "operation, laboratory linkage, alerts, case/event workflows, security, governance, support, and migration "
        "where those topics occur in the report. Then provide: (a) capability and gap matrix; (b) integration "
        "dependencies and risks; (c) short-term recommendations for 0-6 months focused on validation, governance, "
        "minimum viable interoperability, and continuity; and (d) long-term recommendations for 6-24 months focused "
        "on architecture, standards-based integration, phased migration, capacity, sustainability, monitoring, and "
        "decommissioning only after validated replacement. Mark all timeframes as proposed.\n"
        "Follow the user's requested length and format. Use clear headings, short paragraphs, bullets, or a plain-text "
        "table as appropriate. Return only the requested deliverable, without explaining these instructions.\n\n"
        f"GENERATED REPORT:\n{report}\n\nUNDERLYING SELECTED-ASSESSMENT EVIDENCE:\n{evidence or 'Not supplied'}"
    )
    body = json.dumps({
        "model": OLLAMA_MODEL,
        "stream": False,
        "think": False,
        "messages": [{"role": "system", "content": system + recommendation_theme_instructions(question)}, *history, {"role": "user", "content": question}],
        "options": {"temperature": 0.1, "num_ctx": 16384, "num_predict": 1200},
        "keep_alive": "5m",
    }).encode("utf-8")
    request = Request(f"{OLLAMA_URL}/api/chat", data=body, headers={"Content-Type": "application/json"}, method="POST")
    try:
        with urlopen(request, timeout=420) as response:
            result = json.loads(response.read().decode("utf-8"))
    except (URLError, TimeoutError, OSError) as exc:
        raise RuntimeError("The local AI model is unavailable") from exc
    answer = str((result.get("message") or {}).get("content") or "").strip()
    if not answer:
        raise RuntimeError("The local AI model returned no response")
    return {"answer": answer, "model": OLLAMA_MODEL}


def stream_grounded_recommendations(handler, payload):
    report = str(payload.get("report") or "").strip()
    evidence = str(payload.get("evidence") or "").strip()
    if not report:
        raise ValueError("Generate a report first")

    def ask_model(messages):
        body = json.dumps({
            "model": OLLAMA_MODEL, "stream": False, "think": False, "format": "json",
            "messages": messages,
            "options": {"temperature": 0.1, "num_ctx": 4096, "num_predict": 256},
            "keep_alive": "5m",
        }).encode("utf-8")
        request = Request(f"{OLLAMA_URL}/api/chat", data=body,
                          headers={"Content-Type": "application/json"}, method="POST")
        with urlopen(request, timeout=300) as response:
            result = json.loads(response.read().decode("utf-8"))
        if result.get("done_reason") == "length":
            raise ValueError("Model reached its output limit; draft was not accepted")
        return str((result.get("message") or {}).get("content") or "")

    handler.send_response(200)
    handler.send_header("Content-Type", "application/x-ndjson; charset=utf-8")
    handler.send_header("Cache-Control", "no-store")
    handler.end_headers()
    awake_guard = keep_display_awake_for_local_ai()
    try:
        for packet in draft_report(report, evidence, ask_model):
            if "audit" in packet:
                packet["coverage_check"] = format_audit(packet.pop("audit"))
            packet["model"] = OLLAMA_MODEL
            handler.wfile.write((json.dumps(packet, ensure_ascii=False) + "\n").encode("utf-8"))
            handler.wfile.flush()
    except (BrokenPipeError, ConnectionResetError):
        return
    finally:
        restore_normal_power_state(awake_guard)


def stream_ollama_report_chat(handler, payload):
    ensure_ollama_service()
    if is_recommendation_request(payload.get("question", "")):
        return stream_grounded_recommendations(handler, payload)
    report = str(payload.get("report") or "").strip()[:60000]
    evidence = str(payload.get("evidence") or "").strip()[:100000]
    question = str(payload.get("question") or "").strip()[:2000]
    section = str(payload.get("section") or "").strip().lower()
    if section not in {"inventory", "profiling", "gap"}:
        if re.search(r"\bGap Analysis\b", report, re.I):
            section = "gap"
        elif re.search(r"\bProfiling\b", report, re.I):
            section = "profiling"
        else:
            section = "inventory"
    if not report or not question:
        raise ValueError("Generate a report and enter a question first")
    complex_request = bool(re.search(
        r"professional report|thematic|open-ended|all tables|dhis2|integration analysis|action plan|cross-cutting|"
        r"what analys(?:is|es)|analys(?:is|e|es|ing)|insight|discussion|discuss",
        question, re.I,
    ))
    thematic_request = bool(re.search(r"thematic|open-ended", question, re.I))
    table_request = bool(re.search(r"\btable(?:s)?\b|discuss(?:ion)?", question, re.I))
    integration_request = bool(re.search(r"dhis2|integration analysis|capability.*gap|functional.*gap", question, re.I))
    recommendation_request = bool(re.search(
        r"\brecommend(?:ation|ations|ed|ing)?\b|"
        r"(?:consolidat|prioriti[sz]|deduplicat|link).{0,50}recommend|"
        r"recommend.{0,50}(?:supporting finding|evidence|prioriti[sz]|consolidat)",
        question, re.I,
    ))
    supporting_evidence_request = bool(re.search(r"professional report|dhis2|integration analysis|cross-cutting", question, re.I)) or recommendation_request
    history = []
    for item in list(payload.get("history") or [])[-4:]:
        role = str(item.get("role") or "")
        content = str(item.get("content") or "").strip()[:2000]
        if role in {"user", "assistant"} and content:
            history.append({"role": role, "content": content})
    required_applications = ""
    if thematic_request and evidence:
        # The gap-analysis evidence is normally much richer than the generated
        # summary. Supplying only its first few records biases the model toward
        # whichever domains happen to occur first and produces incomplete themes.
        thematic_evidence = compact_thematic_evidence(evidence)
        source_material = (
            f"REPORT CONTEXT:\n{report[:2500]}\n\n"
            f"COMPLETE COMPACT EVIDENCE ({thematic_evidence.count(chr(10)) + 1} records):\n{thematic_evidence}"
        )
    elif integration_request and evidence:
        try:
            inventory_records = [item for item in json.loads(evidence) if isinstance(item, dict)]
        except (TypeError, json.JSONDecodeError):
            inventory_records = []
        application_names = []
        seen_application_names = set()
        integration_records = []
        for index, item in enumerate(inventory_records, 1):
            values = item.get("values") if isinstance(item.get("values"), dict) else {}
            if item.get("evidence_type") != "inventory" and "Inventory Name" not in values:
                continue
            name = str(values.get("Inventory Name") or item.get("record") or f"Inventory system {index}").strip()
            name = re.sub(r"^Inventory:\s*", "", name, flags=re.I)
            normalized_name = re.sub(r"\s+", " ", name).casefold()
            if normalized_name in seen_application_names:
                continue
            seen_application_names.add(normalized_name)
            application_names.append(name)
            selected = {key: values[key] for key in (
                "Inventory Name", "Surveillance Functions", "Data Captured", "Technology", "Has Api",
                "Linked To Other Systems", "Comments",
            ) if key in values}
            integration_records.append({"tool": name, "values": selected})
        required_applications = "; ".join(application_names)
        compact_integration_evidence = json.dumps(integration_records, ensure_ascii=False, separators=(",", ":"))
        source_material = (
            f"GENERATED REPORT:\n{report[:8000]}\n\n"
            f"REQUIRED APPLICATION LIST ({len(application_names)} rows):\n{required_applications or 'No named applications'}\n\n"
            f"COMPLETE COMPACT INVENTORY EVIDENCE:\n{compact_integration_evidence[:18000]}"
        )
    elif recommendation_request:
        source_material = compact_recommendation_material(report, evidence)
    elif supporting_evidence_request and evidence:
        source_material = f"GENERATED REPORT:\n{report}\n\nSUPPORTING EVIDENCE:\n{evidence[:6500]}"
    else:
        source_material = f"GENERATED REPORT:\n{report}"
    system = (
        "You are an Africa CDC public-health surveillance report analyst. Use only the supplied report and evidence. "
        "Never invent facts, figures, system functions, institutions, or implementation status. Distinguish recorded "
        "evidence, analytical interpretation, and proposed actions. State when evidence is insufficient. Preserve exact "
        "system names, figures, denominators, country, and reporting period. Interpret findings instead of repeating them. "
        "For every substantive conclusion, recommendation, or comparison, cite the exact report section, table, system, "
        "field, or evidence record using an 'Evidence reference:' label. If no precise reference exists, omit the claim "
        "or state 'Evidence unavailable'. Treat blank, missing, unknown, and not-confirmed values as unknown, never as No "
        "or as proof that a capability is absent. Do not make causal claims unless the report explicitly records the "
        "outcome; otherwise label the statement 'Interpretation'. Copy reported figures exactly. For any newly calculated "
        "figure, show its numerator, denominator, and calculation and label it 'Calculated from report evidence'. Silently "
        "verify every system name, figure, denominator, and evidence reference before responding, and remove any statement "
        "that cannot pass this check. "
        "Link recommendations to recorded evidence and label inferred actions, leads, and timeframes as Proposed. "
        "For thematic analysis, analyze the complete supplied evidence across all domains before writing. Consolidate "
            "overlapping findings into cross-cutting themes rather than following questionnaire order. Begin with a "
        "one-sentence priority method. Then provide a Markdown table with exactly these columns: Theme; Supporting "
        "evidence; Affected systems or domains; Operational implication; Priority. Ground each row in multiple relevant "
        "records where available, name the systems and domains actually recorded, and assign Critical, High, Medium, or "
        "Low based on urgency, breadth of impact, and risk to outbreak detection or response. After the table, provide "
        "Cross-cutting interpretation, a numbered Recommended priority sequence, and Evidence limitations. Explicitly "
        "identify material contradictions or divergent evidence. Do not stop after the first few themes and do not "
        "fabricate counts. For table "
        "discussion, discuss each actual table, patterns, variation, gaps, implications, and cross-table insights; a "
        "denominator of applicable responses is not a count of systems. Never describe affirmative answers or scored "
        "yes/no answers as systems; use the report's explicit system-level capability count. Distinguish the Inventory "
        "total from the Profiling denominator. For DHIS2 analysis, use the supplied authoritative "
        "toolkit baseline, including case investigation, contact follow-up, and case classification. Treat Unified DHIS2 "
        "as enhancing an existing DHIS2-based tool. Use the earlier six-column matrix format with Country validation "
        "gap and Proposed tool disposition as the final two columns. Then separate "
        "proposed 0-6 month and 6-24 month actions. Be concise and "
        "answer the user's requested deliverable directly.\n\n"
        + source_material
    )
    if recommendation_request:
        system = (
            "You are an Africa CDC surveillance assessment analyst. Use only the supplied findings. Never invent or "
            "broaden evidence, and never assign one system's finding to another. Consolidate genuine duplicates and "
            "link every proposed action to explicit supporting evidence. Preserve system names and qualifications. "
            "For every item, include an Evidence reference naming the exact report section, table, system, field, or "
            "evidence record. Treat blank, missing, unknown, and not-confirmed values as unknown, never as No. Copy "
            "reported figures exactly; for calculations show the numerator and denominator. Omit claims that cannot be "
            "traced to the supplied material. If evidence is incomplete or contradictory, state the validation need. Return the complete requested "
            "deliverable before stopping.\n\n" + source_material
        )
    if thematic_request:
        thematic_scope = {
            "profiling": (
                "This is a Profiling report. Develop evidence-led themes from the detailed system profiles: governance "
                "and ownership; surveillance workflows and functions; data collection and quality; interoperability, "
                "APIs and standards; analytics and reporting; security and privacy; offline access and infrastructure; "
                "and documentation, training, support and sustainability. Do not impose Gap Analysis domain findings or "
                "scores. Provide at least 7 distinct themes and include more whenever additional unique, well-supported "
                "themes are identified. Consolidate genuine overlaps and do not create duplicate or weak themes."
            ),
            "gap": (
                "This is a Gap Analysis report. Develop cross-domain themes from recorded gaps, strengths, explanations, "
                "comments and recommendations. Cover governance, architecture, infrastructure, workforce, data quality, "
                "surveillance workflows, analytics, security, financing and coverage only where supported. Consolidate "
                "related domain findings rather than following questionnaire order. Provide at least 7 distinct themes "
                "and include more whenever additional unique, well-supported themes are identified. Do not create "
                "duplicate or weak themes."
            ),
            "inventory": (
                "This is an Inventory report. Develop portfolio-level themes from system purposes and types, surveillance "
                "function coverage, responsible institutions, geographic or operational coverage, duplication, "
                "complementarity and documented portfolio gaps. Do not infer detailed capabilities or Gap Analysis "
                "findings that Inventory did not record. Provide at least 7 distinct themes and include more whenever "
                "additional unique, well-supported themes are identified; if fewer than 7 are supported, "
                "state the evidence limitation explicitly rather than inventing themes."
            ),
        }[section]
        user_message = (
            f"{question}\n\n"
            "MANDATORY OUTPUT CONTRACT: Read every supplied evidence record before answering. Return one complete "
            f"analysis covering the entire {section} evidence base, not a continuation or a sample. {thematic_scope} "
            "Use a compact Markdown table with these columns: Theme | Supporting evidence | Affected systems or domains | "
            "Operational implication | Priority. Keep each cell concise but substantive. In Supporting evidence use no "
            "more than three short semicolon-separated evidence points; do not use bullets or HTML line breaks inside "
            "table cells. Keep Affected systems or domains to names only and Operational implication to one sentence. Cite "
            "recorded system names, figures, institutions and concrete gaps. Use Critical/High/Medium/Low priorities, "
            "judged by urgency, breadth and outbreak-detection or response risk. After the table add exactly three short "
            "sections: Cross-cutting interpretation; Recommended priority sequence (numbered); Evidence limitations and "
            "contradictions. Sequence actions by dependency: governance and financing enable infrastructure and workforce, "
            "which enable interoperability, quality and response performance. Treat technologies named in respondent "
            "recommendations, including Starlink, as proposals requiring feasibility assessment, not settled actions. "
            "Do not equate absence of a public laboratory-results portal with a requirement to expose identifiable data; "
            "state privacy and authorization safeguards. If one record says role-based access is absent and another says "
            "it exists, report the inconsistency instead of choosing one. Distinguish EWARS not capturing complete case "
            "investigation/outcome details from case-investigation forms being absent. Do not claim government-led "
            "training is absent when refresher training is recorded. Use plain ASCII punctuation only: straight quotes, "
            "hyphens, and apostrophes; never output curly punctuation or mojibake. Do not write long theme-by-theme essays. "
            "Make Cross-cutting interpretation one paragraph of no more than 120 words. Give at least 13 concise, "
            "evidence-linked items in Recommended priority sequence, with each item no longer than 25 words. Cover "
            "distinct actions across the supported themes, order them by urgency and dependency, and do not pad the "
            "list with duplicates. Make Evidence limitations and contradictions one paragraph of no more "
            "than 100 words. Do not stop before the final limitations section."
        )
    elif integration_request:
        user_message = (
            f"{question}\n\n"
            "AUTHORITATIVE UNIFIED DHIS2 SURVEILLANCE TOOLKIT BASELINE: Treat the following as documented toolkit "
            "functionality for this comparison: (1) multi-disease surveillance for notifiable and epidemic-prone "
            "diseases, with mpox, cholera and Ebola as initial reference configurations; (2) continuous routine and "
            "rapidly adaptable outbreak surveillance; (3) case notification; (4) case investigation covering "
            "demographic, clinical, epidemiological, exposure, travel and risk-factor information; (5) case "
            "classification, including suspected, probable, confirmed, discarded and recovered; (6) contact "
            "registration linked to source cases; (7) contact tracing assignments and activities; (8) repeated contact "
            "follow-up covering symptoms, status, missed visits and completion; (9) laboratory management covering "
            "specimen collection, requests, identification, transportation, testing and results; (10) case outcome "
            "monitoring covering hospitalization, isolation, recovery and death; (11) unified records connecting "
            "notification, investigation, laboratory, contacts and outcomes in one database; and (12) disease-specific "
            "forms and workflows. Do not mark any of these functions as absent from the toolkit. Distinguish documented "
            "toolkit capability from configuration or verified use in the selected country.\n\n"
            "MANDATORY DHIS2 ANALYSIS CONTRACT: Use only capabilities explicitly recorded in this assessment. Do not "
            "use general knowledge of DHIS2 and do not invent modules, workflows, APIs, native features, or form designs. "
            "The Technology field identifies platform implementation; a system that exchanges data with DHIS2 is not "
            "therefore a DHIS2 system. Do not classify eLIMS, EWARS, Go.Data, ODK, Kobo, ONA, Eyer, CoDA, NIS, VLSM, or "
            "paper forms as DHIS2 unless the record explicitly says their technology is DHIS2. Distinguish: function "
            "recorded in an existing DHIS2-based system; function recorded only in another system; and DHIS2 replacement "
            "capability not documented. Absence of evidence is not evidence that DHIS2 cannot perform a function. "
            "Produce one matrix row per unique recorded tool, not one row per function. Produce exactly one Markdown "
            "table and show its header only once. Use exactly these columns: Tool | Recorded functions | Functions "
            "covered by Unified DHIS2 Toolkit | Required functions not covered by toolkit | Country validation gap | "
            "Proposed tool disposition. Keep product-capability gaps separate from country configuration or validation "
            "gaps. The Proposed tool disposition must advise what to do with the existing tool and must never say "
            "Configure in DHIS2. Choose only: Retire after "
            "validated replacement; Integrate with Unified DHIS2 and retain; Retain as an external specialist application; "
            "Retain temporarily pending validation; Integrate with Unified DHIS2; or Evidence insufficient. Use "
            "Evidence insufficient when no usable function or workflow evidence is recorded for a tool, and state the "
            "specific validation need in the Country validation gap column. If all required recorded functions of an "
            "application are covered by the toolkit, recommend Retire after validated replacement, subject to data "
            "migration, workflow equivalence, offline continuity, user acceptance and cutover validation. If the "
            "tool performs a required function outside the toolkit baseline, recommend integration and retention "
            "or external retention and name that missing function. If an existing tool is named DHIS2 or HMIS/DHIS2, "
            "recommend retaining it and enhancing its functionality through Unified DHIS2; never recommend retiring "
            "or replacing that DHIS2 platform merely because the toolkit covers its functions. "
            f"The matrix body must contain exactly one row for every unique tool in this required list: "
            f"{required_applications or 'No named tools'}. Do not combine tools into one row, duplicate a tool, omit a "
            "listed tool, add a tool not in the list, repeat the header, or write 'Same as above'. If a listed tool has little evidence, "
            "still include it and mark its disposition Evidence insufficient. After drafting, count the matrix rows and silently correct "
            "the table until its row count equals the required unique-tool count. The column Required functions not "
            "covered by toolkit may contain only required functions absent from the authoritative 12-function toolkit "
            "baseline. Do not place functions missing from the existing tool, functions unconfigured in-country, or "
            "unvalidated functions in that column; describe those validation needs in the narrative after the matrix. If the toolkit baseline covers "
            "all required functions, write None in the toolkit-gap column. "
            "API operation, identifiers, workflow equivalence, offline performance, or migration readiness is not "
            "recorded. Recommend integration when another tool has a required function outside the authoritative toolkit "
            "baseline or provides a specialist workflow that should remain authoritative. Do not recommend integration "
            "merely because toolkit configuration or country use has not been verified. Interpret obvious typographical "
            "variants such as 'Yas' as the intended 'Yes', "
            "while noting the normalization only when it materially affects the conclusion. Treat blank API fields as "
            "unknown, not as evidence that an API is absent. Do not recommend retaining paper "
            "as the preferred end state; describe it only as a continuity fallback when supported. Keep the capability "
            "matrix compact. Then provide: Integration dependencies and risks; Proposed 0-6 month actions; Proposed 6-24 "
            "month actions; Validation gates and indicators; Evidence limitations. Put governance, workflow validation, "
            "identifier mapping, data ownership, offline continuity, security, and user acceptance before migration or "
            "decommissioning. Do not recommend decommissioning until functional equivalence and continuity are validated. "
            "Keep the complete answer below 1,500 words and do not stop before Evidence limitations. Use ASCII punctuation."
        )
    elif recommendation_request:
        user_message = (
            f"{question}\n\n"
            "Follow the thematic recommendation structure. Retain each distinct issue within its relevant "
            "theme or explicitly flag it as unresolved. Use evidence references, proposed practical actions "
            "and monitoring measures. Do not invent urgency rankings or silently omit issues to meet the "
            "theme limit. End with evidence limitations."
        )
    elif table_request:
        user_message = (
            f"{question}\n\n"
            "MANDATORY TABLE-DISCUSSION CONTRACT: First identify every actual table present in the supplied report. "
            "Discuss all of them in report order and finish the whole response. For each table use its numbered exact "
            "title followed by one concise analytical paragraph of no more than 100 words. Integrate the strongest "
            "pattern, meaningful variation or outlier, key evidence gap, operational implication, and one supported "
            "action without repeating subheadings. Use exact figures, system names, numerators and denominators where "
            "reported. Do not repeat the same finding under multiple tables; place it where it is best supported. "
            "Do not invent tables, values, system capabilities, or causal explanations. A field-completeness percentage "
            "measures documentation completeness, not system performance or implementation quality. Do not treat an "
            "optional or undocumented field as a confirmed operational deficiency. Distinguish a count of recorded tools, "
            "functions, fields, or responses from a count of facilities, deployments, users, or geographic areas. When a "
            "table contains multi-select functions, state that totals can exceed the number of systems. When evidence is "
            "insufficient, say what must be validated. End with exactly three sections: Cross-table insights; Decision "
            "implications; Evidence limitations. Limit each closing section to one short paragraph and keep the complete "
            "response below 1,000 words. Do not stop before Evidence limitations. Use plain ASCII punctuation."
        )
    else:
        user_message = question
    body = json.dumps({
        "model": OLLAMA_MODEL,
        "stream": True,
        "think": False,
        "messages": [{"role": "system", "content": system + recommendation_theme_instructions(question)}, *history, {"role": "user", "content": user_message}],
        "options": {
            "temperature": 0.1,
            # The 4B local model becomes substantially slower and uses several
            # extra GB at 12K context. Thematic evidence is compacted above so
            # the complete record set fits this smaller, faster context.
            "num_ctx": 6144 if thematic_request else (8192 if integration_request else (8192 if recommendation_request else (12288 if complex_request else 6144))),
            "num_predict": 3000 if thematic_request else (3000 if integration_request else (3000 if recommendation_request else (2400 if table_request else (2400 if complex_request else 700)))),
        },
        "keep_alive": "5m",
    }).encode("utf-8")
    request = Request(f"{OLLAMA_URL}/api/chat", data=body, headers={"Content-Type": "application/json"}, method="POST")
    local_ai_guard_active = keep_display_awake_for_local_ai()
    try:
        response = urlopen(request, timeout=300)
    except HTTPError as exc:
        restore_normal_power_state(local_ai_guard_active)
        detail = exc.read().decode("utf-8", "replace").strip()
        try:
            detail = json.loads(detail).get("error") or detail
        except json.JSONDecodeError:
            pass
        raise RuntimeError(f"Ollama rejected the request: {detail or exc.reason}") from exc
    except (URLError, TimeoutError, OSError) as exc:
        restore_normal_power_state(local_ai_guard_active)
        reason = str(getattr(exc, "reason", exc)).lower()
        if "timed out" in reason or "timeout" in reason:
            raise RuntimeError(
                "The local AI model did not begin responding within five minutes. "
                "Wait for any earlier analysis to finish, then try again."
            ) from exc
        raise RuntimeError(
            "Could not connect to the local AI service. Confirm Ollama is running and the configured model is installed."
        ) from exc
    handler.send_response(200)
    handler.send_header("Content-Type", "application/x-ndjson; charset=utf-8")
    handler.send_header("Cache-Control", "no-store")
    handler.end_headers()
    answer_parts = []
    try:
        with response:
            for raw_line in response:
                if not raw_line.strip():
                    continue
                item = json.loads(raw_line.decode("utf-8"))
                content = str((item.get("message") or {}).get("content") or "")
                if content:
                    answer_parts.append(content)
                packet = {"delta": content, "done": bool(item.get("done")), "model": OLLAMA_MODEL}
                handler.wfile.write((json.dumps(packet, ensure_ascii=False) + "\n").encode("utf-8"))
                handler.wfile.flush()
            warning = assistant_numeric_grounding_warning("".join(answer_parts), source_material)
            if warning:
                packet = {"delta": warning, "done": True, "model": OLLAMA_MODEL, "grounding_warning": True}
                handler.wfile.write((json.dumps(packet, ensure_ascii=False) + "\n").encode("utf-8"))
                handler.wfile.flush()
    except (BrokenPipeError, ConnectionResetError):
        return
    finally:
        restore_normal_power_state(local_ai_guard_active)


def _analysis_result(section, source):
    profile = profile_state()
    state = current_state()
    phase1 = profile.get("phase1") or {}
    return {
        "section": section,
        "source": source,
        "country_name": phase1.get("country_name", ""),
        "reporting_period": phase1.get("reporting_period", ""),
        "text": analysis_section_text(section),
        "profile": profile,
        "state": state,
        "summary": report_data(),
        "assistant_context": analysis_assistant_context(section, profile, state),
    }


def filtered_analysis_report(section, country_name, reporting_period, use_current=False):
    section = str(section or "").strip().lower()
    if section not in {"inventory", "profiling", "gap"}:
        raise ValueError("Select Inventory, Profiling or Gap Analysis")
    country_name = str(country_name or "").strip()
    reporting_period = str(reporting_period or "").strip()
    if not country_name or not reporting_period:
        raise ValueError("Country and reporting period are required")

    current_phase1 = profile_state().get("phase1") or {}
    current_matches = (
        bool(use_current)
        and str(current_phase1.get("country_name") or "").strip().casefold() == country_name.casefold()
        and str(current_phase1.get("reporting_period") or "").strip().casefold() == reporting_period.casefold()
    )
    if current_matches:
        return _analysis_result(section, "current")

    # Report generation must not replace the user's active data-entry workspace.
    # Holding the runtime lock also prevents concurrent requests from observing
    # the short-lived report workspace.
    with _RUNTIME_LOCK:
        snapshot = workspace_snapshot()
        original_project_id = active_project_id()
        with connect() as db:
            setting_rows = db.execute(
                "SELECT key,value FROM app_settings WHERE key IN ('redcap_last_record_id','redcap_last_pull_at')"
            ).fetchall()
        original_settings = {row["key"]: row["value"] for row in setting_rows}
        try:
            loaded = pull_current_project_from_redcap(
                country_name, reporting_period, reset_group_accounts=False
            )
            result = _analysis_result(section, "redcap")
            result["record_id"] = loaded["record_id"]
            return result
        finally:
            with connect() as db:
                db.execute(
                    "UPDATE app_settings SET value=? WHERE key='active_project_id'",
                    (str(original_project_id),),
                )
                for key in ("redcap_last_record_id", "redcap_last_pull_at"):
                    if key in original_settings:
                        db.execute(
                            "INSERT INTO app_settings(key,value) VALUES(?,?) "
                            "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                            (key, original_settings[key]),
                        )
                    else:
                        db.execute("DELETE FROM app_settings WHERE key=?", (key,))
            restore_project(snapshot, reset_group_accounts=False)


def filtered_analysis_docx(section, country_name, reporting_period, use_current=False):
    if Document is None:
        raise RuntimeError("python-docx is required for Word export")
    from report_builder import build_report
    result = filtered_analysis_report(
        section, country_name, reporting_period, use_current=use_current
    )
    phase = {"inventory": 1, "profiling": 2, "gap": 3}[result["section"]]
    return build_report(
        result["profile"], result["state"], DOMAINS, result["summary"],
        STATIC_DIR / "africa_cdc_logo_full.png", phases=(phase,),
    )


def assistant_answer_docx(payload):
    if Document is None:
        raise RuntimeError("python-docx is required for Word export")
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    answer = str(payload.get("answer") or "").strip()
    if not answer:
        raise ValueError("Generate an assistant response before downloading it")
    document = Document(); section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.65)
    normal = document.styles["Normal"]; normal.font.name = "Calibri"; normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.1
    for name, size, color in (("Title", 22, "0B2545"), ("Heading 1", 16, "2E74B5"), ("Heading 2", 13, "2E74B5"), ("Heading 3", 12, "1F4D78")):
        style = document.styles[name]; style.font.name = "Calibri"; style.font.size = Pt(size); style.font.color.rgb = RGBColor.from_string(color)
    document.add_heading(str(payload.get("title") or "Report Assistant Analysis").strip()[:160], 0)
    metadata = " | ".join(str(payload.get(key) or "").strip() for key in ("country_name", "reporting_period") if payload.get(key))
    if metadata:
        paragraph = document.add_paragraph(metadata); paragraph.runs[0].italic = True
    lines = answer.replace("\r", "").split("\n"); index = 0
    while index < len(lines):
        raw = lines[index].strip()
        if not raw: index += 1; continue
        if "|" in raw and index + 1 < len(lines) and re.match(r"^\|?[\s|:-]+\|?$", lines[index + 1].strip()):
            table_lines = [raw]; index += 2
            while index < len(lines) and "|" in lines[index].strip():
                table_lines.append(lines[index].strip()); index += 1
            rows = [[re.sub(r"\*\*([^*]+)\*\*", r"\1", cell.strip()) for cell in line.strip("|").split("|")] for line in table_lines]
            columns = max(len(row) for row in rows); table = document.add_table(rows=1, cols=columns)
            table.style = "Table Grid"; table.autofit = False; table.alignment = WD_TABLE_ALIGNMENT.CENTER
            if columns == 5:
                widths = [1.4, 2.8, 2.0, 2.75, 0.75]
            elif columns == 6:
                widths = [1.1, 1.9, 2.0, 1.8, 1.4, 1.5]
            else:
                widths = [9.7 / columns] * columns
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for col, value in enumerate(rows[0]):
                table.rows[0].cells[col].width = Inches(widths[col])
                table.rows[0].cells[col].text = value
                table.rows[0].cells[col].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for run in table.rows[0].cells[col].paragraphs[0].runs:
                    run.bold = True; run.font.size = Pt(8.5); run.font.color.rgb = RGBColor(255, 255, 255)
                shading = OxmlElement("w:shd"); shading.set(qn("w:fill"), "176B45"); table.rows[0].cells[col]._tc.get_or_add_tcPr().append(shading)
            for values in rows[1:]:
                cells = table.add_row().cells
                for col, value in enumerate(values):
                    cells[col].width = Inches(widths[col]); cells[col].text = value
                    cells[col].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    for paragraph in cells[col].paragraphs:
                        paragraph.paragraph_format.space_after = Pt(0)
                        for run in paragraph.runs: run.font.size = Pt(8.5)
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", raw); numbered_heading = re.match(r"^(\d+)\.\s+(.+)$", raw)
        if heading:
            document.add_heading(re.sub(r"\*\*([^*]+)\*\*", r"\1", heading.group(2)), level=len(heading.group(1)))
        elif raw.rstrip(":").lower() in {"cross-cutting interpretation", "recommended priority sequence", "evidence limitations and contradictions"}:
            document.add_heading(raw.rstrip(":"), level=1)
        elif numbered_heading and len(raw) < 240:
            document.add_paragraph(numbered_heading.group(2), style="List Number")
        elif re.match(r"^[-*]\s+", raw):
            document.add_paragraph(re.sub(r"^[-*]\s+", "", raw), style="List Bullet")
        else:
            document.add_paragraph(re.sub(r"\*\*([^*]+)\*\*", r"\1", raw))
        index += 1
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Africa CDC Surveillance Systems Assessment")
    output = io.BytesIO(); document.save(output); return output.getvalue()


class RequestRejected(ValueError):
    def __init__(self, message, status=400):
        super().__init__(message)
        self.status = status


class LoginAttemptLimiter:
    """Bounded, process-local limits; expired windows never lock accounts forever."""

    WINDOW_SECONDS = 60
    PAIR_LIMIT = 8
    ACCOUNT_LIMIT = 10
    PEER_LIMIT = 60
    MAX_BUCKETS = 10000

    def __init__(self):
        self.lock = threading.Lock()
        self.buckets = {}

    def allow(self, peer, username):
        now = time.monotonic()
        account = hashlib.sha256(username.encode("utf-8")).hexdigest()
        limits = ((('pair', peer, account), self.PAIR_LIMIT),
                  (('account', account), self.ACCOUNT_LIMIT),
                  (('peer', peer), self.PEER_LIMIT))
        with self.lock:
            self.buckets = {key: item for key, item in self.buckets.items() if item[0] > now}
            retry = max((int(self.buckets[key][0] - now) + 1
                         for key, limit in limits
                         if key in self.buckets and self.buckets[key][1] >= limit), default=0)
            if retry:
                return retry
            if len(self.buckets) + sum(key not in self.buckets for key, _ in limits) > self.MAX_BUCKETS:
                return self.WINDOW_SECONDS
            for key, _ in limits:
                expires, count = self.buckets.get(key, (now + self.WINDOW_SECONDS, 0))
                self.buckets[key] = (expires, count + 1)
        return 0


class ExclusiveThreadingHTTPServer(ThreadingHTTPServer):
    allow_reuse_address = False
    MAX_REQUESTS = 32

    def __init__(self, *args, **kwargs):
        self.request_slots = threading.BoundedSemaphore(self.MAX_REQUESTS)
        self.password_slots = threading.BoundedSemaphore(2)
        self.login_attempts = LoginAttemptLimiter()
        super().__init__(*args, **kwargs)

    def server_bind(self):
        if os.name == "nt":
            self.socket.setsockopt(socket.SOL_SOCKET, socket.SO_EXCLUSIVEADDRUSE, 1)
        super().server_bind()

    def process_request(self, request, client_address):
        if not self.request_slots.acquire(blocking=False):
            try:
                request.settimeout(1)
                body = b'{"error":"The server is busy. Please retry shortly."}'
                response = ("HTTP/1.0 503 Service Unavailable\r\n"
                            "Content-Type: application/json; charset=utf-8\r\n"
                            "Retry-After: 5\r\nConnection: close\r\n"
                            f"Content-Length: {len(body)}\r\n\r\n").encode("ascii")
                request.sendall(response + body)
            except OSError:
                pass
            finally:
                self.shutdown_request(request)
            return
        try:
            super().process_request(request, client_address)
        except Exception:
            self.request_slots.release()
            raise

    def process_request_thread(self, request, client_address):
        try:
            super().process_request_thread(request, client_address)
        finally:
            self.request_slots.release()


class Handler(SimpleHTTPRequestHandler):
    MAX_JSON_BYTES = 16 * 1024 * 1024
    REQUEST_READ_SECONDS = 30.0

    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(STATIC_DIR), **kwargs)

    def setup(self):
        super().setup()
        # A socket I/O timeout does not limit the duration of report generation.
        self.connection.settimeout(self.REQUEST_READ_SECONDS)

    def handle_one_request(self):
        # A per-read timeout alone permits indefinitely trickled request headers.
        # Stop only header reads here; body reads have their own absolute deadline.
        def expire_headers():
            try:
                self.connection.shutdown(socket.SHUT_RD)
            except OSError:
                pass
        self.header_deadline = threading.Timer(self.REQUEST_READ_SECONDS, expire_headers)
        self.header_deadline.daemon = True
        self.header_deadline.start()
        try:
            return super().handle_one_request()
        finally:
            self.header_deadline.cancel()

    def parse_request(self):
        try:
            return super().parse_request()
        finally:
            self.header_deadline.cancel()

    def end_headers(self):
        # The app is frequently updated in place; never let the browser retain
        # an obsolete HTML, JavaScript, or stylesheet interface.
        self.send_header("Cache-Control", "no-store, no-cache, must-revalidate, max-age=0")
        self.send_header("Pragma", "no-cache")
        self.send_header("Expires", "0")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("X-Frame-Options", "DENY")
        self.send_header("Referrer-Policy", "same-origin")
        # The only inline script is the existing Print / Save PDF handler.
        print_hash = base64.b64encode(hashlib.sha256(b"window.print()").digest()).decode("ascii")
        self.send_header("Content-Security-Policy", (
            "default-src 'self'; "
            f"script-src 'self' 'unsafe-hashes' 'sha256-{print_hash}'; "
            "style-src 'self' 'unsafe-inline'; img-src 'self' data: blob:; "
            "font-src 'self'; connect-src 'self'; object-src 'none'; "
            "base-uri 'none'; frame-ancestors 'none'; form-action 'self'"
        ))
        super().end_headers()

    def log_message(self, fmt, *args):
        print(f"[{self.log_date_time_string()}] {fmt % args}")

    def write_body(self, data):
        if self.command != "HEAD":
            self.wfile.write(data)

    def copyfile(self, source, outputfile):
        if self.command != "HEAD":
            super().copyfile(source, outputfile)

    def json_response(self, payload, status=200, extra_headers=None):
        if status < 400 and getattr(self, "sync_workspace_on_success", False):
            self.sync_workspace_on_success = False
            try:
                synced = sync_current_project_to_redcap()
                if isinstance(payload, dict):
                    payload["redcap_save"] = {
                        "record_id": synced["record_id"], "saved_at": synced["synced_at"],
                        "temporary": bool(synced.get("temporary")),
                    }
            except Exception as exc:
                snapshot = getattr(self, "workspace_before_write", None)
                if snapshot:
                    restore_project(snapshot, reset_group_accounts=False)
                payload = {"error": f"REDCap did not accept the save: {exc}. Your entered values remain in the browser; retry the save."}
                status = 502
        data = json.dumps(payload).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(data)))
        self.send_header("Cache-Control", "no-store")
        for key, value in (extra_headers or {}).items():
            self.send_header(key, str(value))
        self.end_headers()
        self.write_body(data)

    def body(self):
        if self.headers.get_all("Transfer-Encoding"):
            raise RequestRejected("Transfer-Encoding is not supported")
        lengths = self.headers.get_all("Content-Length", [])
        if not lengths:
            raise RequestRejected("Content-Length is required", 411)
        if len(lengths) != 1 or not re.fullmatch(r"[0-9]{1,10}", lengths[0].strip()):
            raise RequestRejected("Invalid Content-Length")
        length = int(lengths[0])
        if length > self.MAX_JSON_BYTES:
            raise RequestRejected("Request exceeds the 16 MiB JSON upload limit", 413)
        content_types = self.headers.get_all("Content-Type", [])
        if len(content_types) != 1 or content_types[0].split(";", 1)[0].strip().lower() != "application/json":
            raise RequestRejected("Content-Type must be application/json", 415)
        deadline = time.monotonic() + self.REQUEST_READ_SECONDS
        data = bytearray()
        try:
            while len(data) < length:
                remaining = deadline - time.monotonic()
                if remaining <= 0:
                    raise TimeoutError()
                self.connection.settimeout(remaining)
                chunk = self.rfile.read1(min(65536, length - len(data)))
                if not chunk:
                    raise RequestRejected("Incomplete request body")
                data.extend(chunk)
        except TimeoutError:
            raise RequestRejected("Request body read timed out", 408) from None
        finally:
            self.connection.settimeout(self.REQUEST_READ_SECONDS)
        try:
            payload = json.loads(data.decode("utf-8") or "{}")
        except (UnicodeError, ValueError):
            raise RequestRejected("Request body must contain valid JSON") from None
        if not isinstance(payload, dict):
            raise RequestRejected("Request body must be a JSON object")
        return payload

    def normalized_path(self):
        parsed = urlparse(self.path)
        if parsed.scheme or parsed.netloc or not self.path.startswith("/"):
            raise RequestRejected("Invalid request path")
        # Use the same decoding and platform path handling as static serving.
        # Keep self.path unchanged: decoding it twice would create a new bypass.
        try:
            served_path = Path(super().translate_path(self.path)).resolve()
            relative = served_path.relative_to(STATIC_DIR.resolve()).as_posix()
        except (ValueError, OSError):
            raise RequestRejected("Invalid request path") from None
        return "/" if relative == "." else "/" + relative.casefold()

    def validate_request_origin(self):
        if self.headers.get("Sec-Fetch-Site", "").strip().lower() == "cross-site":
            raise RequestRejected("Cross-site requests are not allowed", 403)
        origins = self.headers.get_all("Origin", [])
        if not origins:
            return  # JSON-only command-line/API clients do not send Origin.
        hosts = self.headers.get_all("Host", [])
        if len(origins) != 1 or len(hosts) != 1:
            raise RequestRejected("Invalid request origin", 403)
        scheme = "https" if SECURE_COOKIES else "http"
        try:
            origin = urlparse(origins[0])
            expected = urlparse(scheme + "://" + hosts[0])
            def identity(value):
                if value.scheme not in {"http", "https"} or not value.hostname or value.username or value.password or value.path or value.params or value.query or value.fragment:
                    raise ValueError()
                return (value.scheme, value.hostname.lower(), value.port or (443 if value.scheme == "https" else 80))
            matches = identity(origin) == identity(expected)
        except ValueError:
            matches = False
        if not matches:
            raise RequestRejected("Request origin does not match this application", 403)

    def login(self, payload):
        username = str(payload.get("username") or "").strip().lower()
        retry = self.server.login_attempts.allow(self.client_address[0], username)
        if retry:
            return self.json_response({"error": "Too many sign-in attempts. Please wait and retry."}, 429, {"Retry-After": retry})
        if not self.server.password_slots.acquire(blocking=False):
            return self.json_response({"error": "Sign-in is busy. Please retry shortly."}, 429, {"Retry-After": 2})
        try:
            password = str(payload.get("password") or "")
            with connect() as db:
                row = db.execute("SELECT id,username,display_name,password_hash,role,phase,group_id,project_scope FROM users WHERE username=? COLLATE NOCASE AND active=1", (username,)).fetchone()
            if not row or not password_matches(password, row["password_hash"]):
                return self.json_response({"error": "Incorrect username or password"}, 401)
            with _RUNTIME_LOCK:
                with connect() as db:
                    current = db.execute("SELECT id,username,display_name,password_hash,role,phase,group_id,project_scope FROM users WHERE id=? AND active=1", (row["id"],)).fetchone()
                    if not current or current["password_hash"] != row["password_hash"] or not clerk_workspace_matches(dict(current)):
                        return self.json_response({"error": "Incorrect username or password"}, 401)
                    row = current
                    if password_needs_rehash(row["password_hash"]):
                        db.execute("UPDATE users SET password_hash=? WHERE id=?", (password_digest(password), row["id"]))
                        db.execute("DELETE FROM sessions WHERE user_id=?", (row["id"],))
                    db.execute("UPDATE users SET last_login_at=? WHERE id=?", (utc_now(), row["id"]))
                token = new_session(row["id"])
            data = json.dumps({"ok": True, "user": {"username": row["username"], "display_name": row["display_name"], "role": row["role"], "phase": row["phase"], "group_id": row["group_id"]}}).encode("utf-8")
            self.send_response(200); self.send_header("Content-Type", "application/json; charset=utf-8"); self.auth_cookie(token); self.send_header("Content-Length", str(len(data))); self.end_headers(); self.write_body(data)
        finally:
            self.server.password_slots.release()

    def redirect(self, location):
        self.send_response(303)
        self.send_header("Location", location)
        self.send_header("Content-Length", "0")
        self.end_headers()

    def auth_cookie(self, token):
        secure = "; Secure" if SECURE_COOKIES else ""
        self.send_header("Set-Cookie", f"africa_cdc_session={token}; Path=/; HttpOnly; SameSite=Lax; Max-Age={SESSION_HOURS * 3600}{secure}")

    def clear_auth_cookie(self):
        secure = "; Secure" if SECURE_COOKIES else ""
        self.send_header("Set-Cookie", f"africa_cdc_session=; Path=/; HttpOnly; SameSite=Lax; Max-Age=0{secure}")

    def authenticated_user(self):
        return session_user(self.headers.get("Cookie"))

    def do_GET(self):
        try:
            path = self.normalized_path()
            if path.startswith("/api/") and path != "/api/assistant/status":
                # Authorization and data reads must refer to the same country,
                # including while another request generates a report preview.
                with _RUNTIME_LOCK:
                    return self.dispatch_get(path)
            return self.dispatch_get(path)
        except RequestRejected as exc:
            return self.json_response({"error": str(exc)}, exc.status)

    def do_HEAD(self):
        return self.do_GET()

    def dispatch_get(self, path):
        user = self.authenticated_user()
        if path == "/api/auth/status":
            with connect() as db:
                setup_required = not bool(db.execute("SELECT 1 FROM users LIMIT 1").fetchone())
            return self.json_response({"setup_required": setup_required, "authenticated": bool(user), "user": user})
        if path == "/health":
            return self.json_response({"status": "ok", "app_version": APP_VERSION, "profile_format_version": PROFILE_FORMAT_VERSION})
        if path in {"/", "/login", "/login.html"}:
            self.path = "/login.html"
            return super().do_GET()
        if path == "/app":
            if not user:
                return self.redirect("/login")
            self.path = "/index.html"
            return super().do_GET()
        if path == "/user-manual.html":
            if not user:
                return self.redirect("/login")
            return super().do_GET()
        if path in {"/admin-guide.html", "/update-guide.html"}:
            if not user:
                return self.redirect("/login")
            if user.get("role") != "admin":
                return self.json_response({"error": "Administrator access is required"}, 403)
            self.path = "/admin-guide.html"
            return super().do_GET()
        if not user:
            if path.startswith("/api/"):
                return self.json_response({"error": "Authentication required"}, 401)
            if path == "/index.html":
                return self.redirect("/login")
        if user and user.get("role") not in {"admin", "coordinator"} and (path.startswith("/api/export") or path in {"/api/profile.json", "/api/report", "/api/full-report", "/api/report.txt", "/api/report.docx"}):
            return self.json_response({"error": "Administrator access is required"}, 403)
        if path == "/api/state":
            return self.json_response(scoped_state(user))
        if path == "/api/profile":
            return self.json_response(scoped_profile(user))
        if path == "/api/my-entries":
            try: return self.json_response(my_entries_data(user))
            except PermissionError as exc: return self.json_response({"error": str(exc)}, 403)
        if path == "/api/my-entries.txt":
            try:
                data = my_entries_text(user).encode("utf-8")
                self.send_response(200); self.send_header("Content-Type", "text/plain; charset=utf-8")
                self.send_header("Content-Disposition", 'attachment; filename="My_Group_Entries.txt"')
                self.send_header("Content-Length", str(len(data))); self.end_headers(); self.write_body(data); return
            except PermissionError as exc: return self.json_response({"error": str(exc)}, 403)
        if path == "/api/projects":
            if user.get("role") not in {"admin", "coordinator"}: return self.json_response({"error": "Coordinator access is required"}, 403)
            return self.json_response(country_projects_state())
        if path == "/api/users":
            if user.get("role") != "admin": return self.json_response({"error": "Administrator access is required"}, 403)
            return self.json_response({"users": application_users()})
        if path == "/api/redcap/status":
            if user.get("role") != "admin": return self.json_response({"error": "Administrator access is required"}, 403)
            return self.json_response(redcap_status())
        if path == "/api/assistant/status":
            return self.json_response(assistant_status())
        if path == "/api/profile.json":
            data = json.dumps(full_profile(), ensure_ascii=False, indent=2).encode("utf-8")
            self.send_response(200); self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Content-Disposition", 'attachment; filename="Africa_CDC_Assessment_Profile.json"')
            self.send_header("Content-Length", str(len(data))); self.end_headers(); self.write_body(data); return
        if path == "/api/report.txt":
            data = complete_report_text().encode("utf-8")
            self.send_response(200); self.send_header("Content-Type", "text/plain; charset=utf-8")
            self.send_header("Content-Disposition", 'attachment; filename="Africa_CDC_Assessment_Report.txt"')
            self.send_header("Content-Length", str(len(data))); self.end_headers(); self.write_body(data); return
        if path == "/api/report.docx":
            try:
                data = report_docx()
                self.send_response(200); self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
                self.send_header("Content-Disposition", 'attachment; filename="Africa_CDC_Assessment_Report.docx"')
                self.send_header("Content-Length", str(len(data))); self.end_headers(); self.write_body(data)
            except Exception as exc: self.json_response({"error": str(exc)}, 500)
            return
        if path == "/api/report":
            return self.json_response(report_data())
        if path == "/api/full-report":
            return self.json_response({"text": complete_report_text(), "phase3": report_data()})
        if path == "/api/export.xlsx":
            try:
                data = standard_format_workbook((1, 2, 3))
                self.send_response(200)
                self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                self.send_header("Content-Disposition", 'attachment; filename="Africa_CDC_Gap_Analysis.xlsx"')
                self.send_header("Content-Length", str(len(data)))
                self.end_headers()
                self.write_body(data)
            except Exception as exc:
                self.json_response({"error": str(exc)}, 500)
            return
        if path in {"/api/export-phase1.xlsx", "/api/export-phase2.xlsx", "/api/export-phase3.xlsx"}:
            try:
                phase = 1 if "phase1" in path else 2 if "phase2" in path else 3
                data = standard_format_workbook((phase,))
                self.send_response(200); self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                self.send_header("Content-Disposition", f'attachment; filename="Africa_CDC_Phase_{phase}.xlsx"')
                self.send_header("Content-Length", str(len(data))); self.end_headers(); self.write_body(data)
            except Exception as exc: self.json_response({"error": str(exc)}, 500)
            return
        if path == "/":
            self.path = "/index.html"
        return super().do_GET()

    def do_POST(self):
        try:
            path = self.normalized_path()
            self.validate_request_origin()
            payload = self.body()
            workspace_paths = REDCAP_PRIMARY_WRITE_PATHS | {
                "/api/projects/new", "/api/projects/switch", "/api/redcap/pull",
                "/api/redcap/sync", "/api/accounts/reset", "/api/analysis/report",
                "/api/analysis/docx",
            }
            if path in workspace_paths:
                # Serialize workspace writes through their REDCap acknowledgement.
                # Body reads, sign-in and AI streams do not hold this lock.
                with _RUNTIME_LOCK:
                    return self.dispatch_post(path, payload)
            return self.dispatch_post(path, payload)
        except RequestRejected as exc:
            self.close_connection = True
            return self.json_response({"error": str(exc)}, exc.status)

    def dispatch_post(self, path, payload):
        try:
            if path == "/api/auth/setup":
                username, password, display_name = validate_credentials(payload.get("username"), payload.get("password"), payload.get("display_name"))
                with connect() as db:
                    db.execute("BEGIN IMMEDIATE")
                    if db.execute("SELECT 1 FROM users LIMIT 1").fetchone():
                        return self.json_response({"error": "Administrator setup has already been completed"}, 409)
                    cursor = db.execute("INSERT INTO users(username,display_name,password_hash,role,created_at) VALUES(?,?,?,?,?)", (username, display_name, password_digest(password), "admin", utc_now()))
                    user_id = cursor.lastrowid
                token = new_session(user_id)
                data = json.dumps({"ok": True, "user": {"username": username, "display_name": display_name, "role": "admin"}}).encode("utf-8")
                self.send_response(201); self.send_header("Content-Type", "application/json; charset=utf-8"); self.auth_cookie(token); self.send_header("Content-Length", str(len(data))); self.end_headers(); self.wfile.write(data); return
            if path == "/api/auth/login":
                return self.login(payload)
            if path == "/api/auth/logout":
                cookie = SimpleCookie(); cookie.load(self.headers.get("Cookie") or "")
                morsel = cookie.get("africa_cdc_session")
                if morsel:
                    token_hash = hashlib.sha256(morsel.value.encode("utf-8")).hexdigest()
                    with connect() as db: db.execute("DELETE FROM sessions WHERE token_hash=?", (token_hash,))
                data = b'{"ok":true}'
                self.send_response(200); self.send_header("Content-Type", "application/json; charset=utf-8"); self.clear_auth_cookie(); self.send_header("Content-Length", str(len(data))); self.end_headers(); self.wfile.write(data); return
            user = self.authenticated_user()
            if not user:
                return self.json_response({"error": "Authentication required"}, 401)
            if path == "/api/auth/change-password":
                try:
                    return self.json_response(change_own_password(
                        user["id"], payload.get("current_password"), payload.get("new_password"), payload.get("confirmation")
                    ))
                except ValueError as exc:
                    return self.json_response({"error": str(exc)}, 400)
            admin_only_paths = {"/api/users/create", "/api/users/reset", "/api/users/change-password", "/api/users/active", "/api/redcap/config", "/api/redcap/test", "/api/redcap/sync", "/api/assistant/provider", "/api/assistant/provider/test"}
            coordinator_paths = {"/api/groups", "/api/inventory/assign", "/api/inventory/unassign", "/api/profiling/assign", "/api/profiling/unassign", "/api/assign", "/api/unassign", "/api/reassign", "/api/accounts/reset", "/api/projects/new", "/api/projects/switch", "/api/redcap/pull", "/api/analysis/report", "/api/analysis/docx", "/api/assistant/chat", "/api/assistant/docx"}
            if path in admin_only_paths and user.get("role") != "admin":
                return self.json_response({"error": "Administrator access is required"}, 403)
            if path in coordinator_paths and user.get("role") not in {"admin", "coordinator"}:
                return self.json_response({"error": "Coordinator access is required"}, 403)
            if user.get("role") not in {"admin", "coordinator"}:
                clerk_allowed = {"inventory": {"/api/profile/phase1/tool"}, "profiling": {"/api/profile/phase2/tool"}, "gap": {"/api/responses"}}
                if path not in clerk_allowed.get(user.get("phase"), set()):
                    return self.json_response({"error": "This account cannot modify that assessment phase"}, 403)
                payload["group_id"] = user.get("group_id")
            if path == "/api/analysis/docx":
                try:
                    data = filtered_analysis_docx(
                        payload.get("section"), payload.get("country_name"), payload.get("reporting_period"),
                        payload.get("use_current", False),
                    )
                    section_name = re.sub(r"[^A-Za-z0-9]+", "_", str(payload.get("section") or "Report")).strip("_")
                    country_name = re.sub(r"[^A-Za-z0-9]+", "_", str(payload.get("country_name") or "Country")).strip("_")
                    filename = f"Africa_CDC_{country_name}_{section_name}_Report.docx"
                    self.send_response(200)
                    self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
                    self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
                    self.send_header("Content-Length", str(len(data)))
                    self.end_headers(); self.wfile.write(data); return
                except (ValueError, RuntimeError) as exc:
                    return self.json_response({"error": str(exc)}, 400)
            if path == "/api/assistant/chat":
                try:
                    if assistant_provider() == "openai":
                        return stream_online_report_chat(self, payload)
                    return stream_ollama_report_chat(self, payload)
                except (ValueError, RuntimeError) as exc:
                    return self.json_response({"error": str(exc)}, 503)
            if path == "/api/assistant/docx":
                try:
                    data = assistant_answer_docx(payload)
                    self.send_response(200)
                    self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
                    self.send_header("Content-Disposition", 'attachment; filename="Africa_CDC_Report_Assistant_Analysis.docx"')
                    self.send_header("Content-Length", str(len(data)))
                    self.end_headers(); self.wfile.write(data); return
                except (ValueError, RuntimeError) as exc:
                    return self.json_response({"error": str(exc)}, 400)
            if path in REDCAP_PRIMARY_WRITE_PATHS:
                self.workspace_before_write = workspace_snapshot()
                self.sync_workspace_on_success = True
            if path == "/api/redcap/sync":
                return self.json_response(sync_current_project_to_redcap())
            if path == "/api/users/create":
                return self.json_response(create_coordinator_account(payload.get("username"), payload.get("display_name")), 201)
            if path == "/api/users/reset":
                return self.json_response(reset_coordinator_password(payload.get("user_id")))
            if path == "/api/users/change-password":
                return self.json_response(admin_change_user_password(
                    payload.get("user_id"), payload.get("new_password"), payload.get("confirmation")
                ))
            if path == "/api/users/active":
                return self.json_response(set_coordinator_active(payload.get("user_id"), payload.get("active")))
            if path == "/api/assistant/provider":
                try: return self.json_response(configure_assistant_provider(payload.get("provider"), payload.get("approved", False)))
                except ValueError as exc: return self.json_response({"error": str(exc)}, 400)
            if path == "/api/assistant/provider/test":
                if assistant_provider() == "openai":
                    try:
                        answer = openai_report_answer({"report": "Connection test report. No assessment data is included.", "question": "Reply with: Online AI connection successful."})
                        return self.json_response({"ok": True, "provider": "openai", "model": OPENAI_MODEL, "message": answer[:200]})
                    except RuntimeError as exc: return self.json_response({"error": str(exc)}, 503)
                status = ollama_status()
                return self.json_response({"ok": bool(status.get("available")), **status})
            if path == "/api/profile/phase1":
                return self.json_response(save_profile_phase("phase1", payload))
            if path == "/api/projects/new":
                return self.json_response(create_country_project(
                    payload.get("country_name"), payload.get("reporting_period"), payload.get("open_existing", False)
                ), 201)
            if path == "/api/projects/switch":
                return self.json_response(switch_country_project(payload.get("project_id")))
            if path == "/api/redcap/config":
                global REDCAP_TOKEN
                api_url = str(payload.get("api_url") or "").strip()
                token = str(payload.get("token") or "").strip()
                normalized_url = api_url.rstrip("/").lower()
                if not api_url.lower().startswith("https://") or not (normalized_url.endswith("/api") or normalized_url.endswith("/api/index.php")):
                    return self.json_response({"error": "Use the HTTPS REDCap API URL ending with /api/ or /api/index.php"}, 400)
                save_app_setting("redcap_api_url", api_url if normalized_url.endswith("index.php") else api_url.rstrip("/") + "/")
                if token:
                    with REDCAP_TOKEN_LOCK:
                        REDCAP_TOKEN = token
                elif not REDCAP_TOKEN:
                    return self.json_response({"error": "API token is required the first time"}, 400)
                return self.json_response(redcap_status())
            if path == "/api/redcap/test":
                return self.json_response(test_redcap_connection())
            if path == "/api/redcap/pull":
                return self.json_response(pull_current_project_from_redcap(payload.get("country_name"), payload.get("reporting_period")))
            if path == "/api/analysis/report":
                return self.json_response(filtered_analysis_report(
                    payload.get("section"), payload.get("country_name"), payload.get("reporting_period"),
                    payload.get("use_current", False),
                ))
            if path == "/api/accounts/reset":
                with connect() as db: aid = db.execute("SELECT id FROM assessments ORDER BY id LIMIT 1").fetchone()[0]
                credentials = reset_group_account(aid, str(payload.get("phase") or ""), int(payload.get("group_id") or 0))
                result = current_state(); result["generated_credentials"] = credentials
                return self.json_response(result)
            if path == "/api/profile/phase1/tool":
                try:
                    result = save_inventory_tool(payload, user=user)
                    return self.json_response(scoped_profile(user) if user.get("role") != "admin" else result)
                except PermissionError as exc:
                    return self.json_response({"error": str(exc)}, 403)
            if path == "/api/profile/phase2":
                return self.json_response(save_profile_phase("phase2", payload))
            if path == "/api/profile/phase2/tool":
                try:
                    result = save_profiling_tool(payload, user=user)
                    return self.json_response(scoped_profile(user) if user.get("role") != "admin" else result)
                except PermissionError as exc:
                    return self.json_response({"error": str(exc)}, 403)
            if path == "/api/profile/import":
                return self.json_response(import_full_profile(payload))
            if path == "/api/import/validate":
                return self.json_response(validate_standardized_excel(payload.get("file", {})))
            if path == "/api/import/phase1":
                return self.json_response(import_phase1_excel(payload.get("file", {})))
            if path == "/api/import/phase2":
                return self.json_response(import_phase2_excel(payload.get("file", {})))
            state = current_state()
            aid = state["assessment"]["id"]
            if path == "/api/groups":
                name = str(payload.get("name", "")).strip()
                if not name:
                    return self.json_response({"error": "Group name is required"}, 400)
                with connect() as db:
                    db.execute("INSERT INTO groups(assessment_id, name) VALUES (?, ?)", (aid, name))
                return self.json_response(current_state(), 201)
            if path == "/api/inventory/assign":
                tool_id = str(payload.get("tool_id") or "").strip()
                try:
                    group_id = int(payload.get("group_id"))
                except (TypeError, ValueError):
                    return self.json_response({"error": "A valid participant group is required"}, 400)
                with connect() as db:
                    group = db.execute("SELECT id FROM groups WHERE id=? AND assessment_id=?", (group_id, aid)).fetchone()
                    if not group:
                        return self.json_response({"error": "Participant group not found"}, 404)
                    row = db.execute("SELECT payload FROM profile_data WHERE phase='phase1'").fetchone()
                    phase1 = ensure_inventory_tool_ids(json.loads(row[0]) if row else {})
                    tools = phase1.get("tools", [])
                    index = next((i for i, item in enumerate(tools) if item.get("_id") == tool_id), None)
                    if index is None:
                        tool = dict(payload.get("tool") or {})
                        if not tool or str(tool.get("_id") or "") != tool_id:
                            return self.json_response({"error": "Inventory tool details are required before assignment"}, 400)
                        tools.append(tool)
                    elif payload.get("tool"):
                        tool = dict(payload["tool"])
                        tool["_id"] = tool_id
                        tools[index] = tool
                    for key in ("country_name", "reporting_period", "comments"):
                        if key in payload:
                            phase1[key] = str(payload.get(key) or "")
                    db.execute(
                        "INSERT INTO profile_data(phase,payload,updated_at) VALUES('phase1',?,?) "
                        "ON CONFLICT(phase) DO UPDATE SET payload=excluded.payload,updated_at=excluded.updated_at",
                        (json.dumps(phase1, ensure_ascii=False), utc_now()),
                    )
                    db.execute("INSERT OR REPLACE INTO inventory_tool_assignments(assessment_id,tool_id,group_id,assigned_at) VALUES(?,?,?,?)", (aid, tool_id, group_id, utc_now()))
                result = current_state(); credentials = ensure_group_account(aid, "inventory", group_id)
                if credentials: result["generated_credentials"] = credentials
                return self.json_response(result)
            if path == "/api/inventory/unassign":
                tool_id = str(payload.get("tool_id") or "").strip()
                if not tool_id:
                    return self.json_response({"error": "Inventory tool is required"}, 400)
                with connect() as db:
                    db.execute("DELETE FROM inventory_tool_assignments WHERE assessment_id=? AND tool_id=?", (aid, tool_id))
                return self.json_response(current_state())
            if path in {"/api/profiling/assign", "/api/profiling/unassign"}:
                tool_id = str(payload.get("tool_id") or "").strip()
                tool_ids = {str(tool.get("_id")) for tool in profile_state()["phase1"].get("tools", [])}
                if tool_id not in tool_ids:
                    return self.json_response({"error": "Profiling system not found"}, 404)
                with connect() as db:
                    if path.endswith("/unassign"):
                        db.execute("DELETE FROM profiling_tool_assignments WHERE assessment_id=? AND tool_id=?", (aid, tool_id))
                    else:
                        try:
                            group_id = int(payload.get("group_id"))
                        except (TypeError, ValueError):
                            return self.json_response({"error": "A valid participant group is required"}, 400)
                        group = db.execute("SELECT id FROM groups WHERE id=? AND assessment_id=?", (group_id, aid)).fetchone()
                        if not group:
                            return self.json_response({"error": "Participant group not found"}, 404)
                        db.execute("INSERT OR REPLACE INTO profiling_tool_assignments(assessment_id,tool_id,group_id,assigned_at) VALUES(?,?,?,?)", (aid, tool_id, group_id, utc_now()))
                result = current_state()
                if not path.endswith("/unassign"):
                    credentials = ensure_group_account(aid, "profiling", group_id)
                    if credentials: result["generated_credentials"] = credentials
                return self.json_response(result)
            if path == "/api/assessment":
                with connect() as db:
                    db.execute("UPDATE assessments SET title=?, scope=? WHERE id=?", (str(payload.get("title", "")).strip() or "Africa CDC Gap Analysis", str(payload.get("scope", "")).strip(), aid))
                return self.json_response(current_state())
            if path == "/api/assign":
                domain_id, group_id = payload.get("domain_id"), int(payload.get("group_id"))
                if domain_id not in {domain["id"] for domain in DOMAINS}:
                    return self.json_response({"error": "Unknown domain"}, 400)
                with connect() as db:
                    existing = db.execute("SELECT group_id FROM assignments WHERE assessment_id=? AND domain_id=?", (aid, domain_id)).fetchone()
                    if existing and existing[0] != group_id:
                        return self.json_response({"error": "This domain is already assigned to another group"}, 409)
                    db.execute("INSERT OR REPLACE INTO assignments VALUES (?, ?, ?, ?)", (aid, domain_id, group_id, utc_now()))
                return self.json_response(current_state())
            if path == "/api/unassign":
                with connect() as db:
                    db.execute("DELETE FROM assignments WHERE assessment_id=? AND domain_id=?", (aid, payload.get("domain_id")))
                return self.json_response(current_state())
            if path == "/api/reassign":
                domain_id = payload.get("domain_id")
                try:
                    group_id = int(payload.get("group_id"))
                except (TypeError, ValueError):
                    return self.json_response({"error": "A valid participant group is required"}, 400)
                if domain_id not in {domain["id"] for domain in DOMAINS}:
                    return self.json_response({"error": "Unknown domain"}, 400)
                with connect() as db:
                    group = db.execute(
                        "SELECT id FROM groups WHERE id=? AND assessment_id=?", (group_id, aid)
                    ).fetchone()
                    if not group:
                        return self.json_response({"error": "Participant group not found"}, 404)
                    db.execute(
                        "INSERT OR REPLACE INTO assignments(assessment_id,domain_id,group_id,assigned_at) VALUES(?,?,?,?)",
                        (aid, domain_id, group_id, utc_now()),
                    )
                    db.execute(
                        "UPDATE responses SET group_id=?, updated_at=? WHERE assessment_id=? AND domain_id=?",
                        (group_id, utc_now(), aid, domain_id),
                    )
                result = current_state(); credentials = ensure_group_account(aid, "gap", group_id)
                if credentials: result["generated_credentials"] = credentials
                return self.json_response(result)
            if path == "/api/responses":
                domain_id, group_id = payload.get("domain_id"), int(payload.get("group_id"))
                with connect() as db:
                    assignment = db.execute("SELECT group_id FROM assignments WHERE assessment_id=? AND domain_id=?", (aid, domain_id)).fetchone()
                    if not assignment or assignment[0] != group_id:
                        return self.json_response({"error": "The selected group does not own this domain"}, 403)
                    for item in payload.get("responses", []):
                        qid = item.get("question_id")
                        if qid not in QUESTION_INDEX or QUESTION_INDEX[qid][0] != domain_id:
                            continue
                        response = item.get("response", "")
                        if isinstance(response, list):
                            response = json.dumps(response)
                        db.execute("""
                            INSERT INTO responses VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                            ON CONFLICT(assessment_id, question_id) DO UPDATE SET
                            group_id=excluded.group_id, response=excluded.response, explanation=excluded.explanation,
                            comment=excluded.comment, updated_at=excluded.updated_at
                        """, (aid, domain_id, qid, group_id, str(response), str(item.get("explanation", "")), str(item.get("comment", "")), utc_now()))
                return self.json_response(scoped_state(user))
            if path == "/api/import":
                imported = parse_excel_files(payload.get("files", []))
                default_group = int(payload.get("group_id") or state["groups"][0]["id"])
                with connect() as db:
                    for qid, item in imported.items():
                        assignment = db.execute("SELECT group_id FROM assignments WHERE assessment_id=? AND domain_id=?", (aid, item["domain_id"])).fetchone()
                        group_id = assignment[0] if assignment else default_group
                        group_slot = item.get("group_slot")
                        if group_slot is not None and group_slot < len(state["groups"]):
                            group_id = state["groups"][group_slot]["id"]
                            db.execute("INSERT OR REPLACE INTO assignments VALUES (?, ?, ?, ?)", (aid, item["domain_id"], group_id, utc_now()))
                        imported_group = str(item.get("group_name", "")).strip()
                        if imported_group:
                            db.execute("INSERT OR IGNORE INTO groups(assessment_id,name) VALUES(?,?)", (aid, imported_group))
                            group_id = db.execute("SELECT id FROM groups WHERE assessment_id=? AND name=? COLLATE NOCASE", (aid, imported_group)).fetchone()[0]
                            db.execute("INSERT OR REPLACE INTO assignments VALUES (?, ?, ?, ?)", (aid, item["domain_id"], group_id, utc_now()))
                        elif not assignment and group_slot is None:
                            db.execute("INSERT INTO assignments VALUES (?, ?, ?, ?)", (aid, item["domain_id"], group_id, utc_now()))
                        existing = db.execute("SELECT response, explanation, comment FROM responses WHERE assessment_id=? AND question_id=?", (aid, qid)).fetchone()
                        response = item["response"] or (existing[0] if existing else "")
                        explanation = item["explanation"] or (existing[1] if existing else "")
                        comment = item["comment"] or (existing[2] if existing else "")
                        db.execute("INSERT OR REPLACE INTO responses VALUES (?, ?, ?, ?, ?, ?, ?, ?)", (aid, item["domain_id"], qid, group_id, response, explanation, comment, utc_now()))
                return self.json_response({"imported": len(imported), "state": current_state()})
            return self.json_response({"error": "Not found"}, 404)
        except sqlite3.IntegrityError as exc:
            return self.json_response({"error": str(exc)}, 409)
        except PermissionError as exc:
            return self.json_response({"error": str(exc)}, 403)
        except ValueError as exc:
            return self.json_response({"error": str(exc)}, 400)
        except Exception as exc:
            return self.json_response({"error": str(exc)}, 500)


def main():
    parser = argparse.ArgumentParser(description="Africa CDC shared web assessment")
    parser.add_argument("--host", default="0.0.0.0")
    parser.add_argument("--port", type=int, default=8080)
    parser.add_argument("--no-browser", action="store_true")
    args = parser.parse_args()
    server = None
    try:
        initialize_database()
        local_url = f"http://127.0.0.1:{args.port}"
        try:
            server = ExclusiveThreadingHTTPServer((args.host, args.port), Handler)
        except OSError as exc:
            # A packaged Windows application otherwise presents this expected
            # duplicate-launch condition as an "Unhandled exception" dialog.
            if getattr(exc, "winerror", None) == 10048 or getattr(exc, "errno", None) in (48, 98, 10048):
                print(f"Africa CDC Web Assessment is already running at {local_url}")
                if not args.no_browser:
                    webbrowser.open(local_url)
                return
            raise
        auto_sync_thread = start_redcap_auto_sync()
        print(f"Africa CDC Web Assessment running at {local_url}")
        print(f"Other devices can connect using this computer's IP address and port {args.port}.")
        if auto_sync_thread:
            print(f"Automatic REDCap sync is enabled every {redcap_auto_sync_minutes()} minutes.")
        if not args.no_browser:
            threading.Timer(1.0, lambda: webbrowser.open(local_url)).start()
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    except Exception:
        try:
            with (APP_DIR / "AfricaCDCWebAssessment.log").open("a", encoding="utf-8") as log:
                log.write(f"\n[{datetime.now().isoformat(timespec='seconds')}] Server stopped unexpectedly\n")
                traceback.print_exc(file=log)
        except Exception:
            pass
        raise
    finally:
        if server is not None:
            server.server_close()


if __name__ == "__main__":
    main()

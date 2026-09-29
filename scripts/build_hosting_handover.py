from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output' / 'Africa_CDC_Server_Setup_Guide.docx'
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1)
s.header_distance = s.footer_distance = Inches(.492)
for name, size, color, before, after in [('Normal',11,'222222',0,6),('Title',26,'0B2545',0,10),('Subtitle',12,'555555',0,8),('Heading 1',16,'2E74B5',18,10),('Heading 2',13,'2E74B5',14,7),('Heading 3',12,'1F4D78',10,5)]:
    st=doc.styles[name]; st.font.name='Calibri'; st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(color)
    st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.line_spacing=1.25
    if name.startswith('Heading'): st.paragraph_format.keep_with_next=True
# Named overrides: compact code blocks, small running furniture; no tables or simulated lists.
from docx.enum.style import WD_STYLE_TYPE
st=doc.styles.add_style('Command',WD_STYLE_TYPE.PARAGRAPH)
st.font.name='Consolas'; st.font.size=Pt(8.5)
st.paragraph_format.space_before=Pt(0); st.paragraph_format.space_after=Pt(4); st.paragraph_format.line_spacing=1.0
st.paragraph_format.keep_together=True
hp=s.header.paragraphs[0]; hp.text='AFRICA CDC  |  SERVER HOSTING HANDOVER'
hp.style=doc.styles['Normal']; hp.runs[0].font.size=Pt(8); hp.runs[0].font.color.rgb=RGBColor.from_string('666666')
fp=s.footer.paragraphs[0]; fp.alignment=2
r=fp.add_run('IT setup guide  |  '); r.font.size=Pt(8)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); fp._p.append(field)
def p(t): doc.add_paragraph(t)
def h(t): doc.add_heading(t,2)
def code(t): doc.add_paragraph(t.strip(), 'Command')
def page(t):
    doc.add_page_break(); doc.add_heading(t,1)

doc.add_paragraph('Server setup and\nhosting guide', 'Title')
doc.add_paragraph('Africa CDC Surveillance Systems Assessment Web Application', 'Subtitle')
p('For: IT / server administrator\nPrepared: 24 September 2026\nApplication reviewed: version 1.0.40 | Profile format: version 3')
h('Purpose and scope')
p('Use this guide to host the supplied application ZIP as one central website for Inventory, Profiling and Gap Analysis. Users access the site in a browser. This guide is based on the source files in the working application folder; IT must confirm that the ZIP received contains the same version and files.')
h('Choose one installation route')
p('Ubuntu Server with Nginx: follow pages 2–5, then pages 8–9. Windows Server with IIS: follow page 2, pages 6–7, then pages 8–9. Both routes run the same Python application.')
h('How the application works')
code('Browser → HTTPS reverse proxy → Python on 127.0.0.1:8080\n                                  ├─ SQLite: accounts and settings\n                                  └─ REDCap API: assessment data')
p('REDCap is the primary assessment store. The local africa_cdc_web.db holds accounts, sessions, permissions, groups, assignments and country indexes. Backing up SQLite alone does not back up assessment responses.')
p('Run one application process only. The application maintains assessment state in memory; do not add multiple workers, replicas or load-balanced copies. It is a standalone HTTP server, not a Flask/Django WSGI application.')
p('IT should review the use of Python http.server before public production exposure: Python documents it as unsuitable for production by itself. The proxy configuration here provides TLS and restricted backend access; it is not a security audit. See reference [3] in section 8.')

page('1. Check the package and configuration')
p('Confirm the ZIP contains server.py, report_builder.py, surveillance_app.py, requirements.txt, standard_assessment_template.xlsx, the complete static folder and .env.example. A ZIP containing only AfricaCDCWebAssessment.exe is not a Linux source package; request the source files for that route.')
p('Confirm the domain, server OS, DNS record, TLS certificate, backup location and support owner. Provide outbound HTTPS to tools.africacdc.org. Restrict SSH/RDP to administrator networks. Users should reach port 443; keep port 8080 private. Port 80 is needed only for redirect/certificate workflows.')
h('Use one authoritative settings file')
p('The .env beside server.py overrides existing process or machine environment values. For this guide, use that file on both platforms. Restrict read access to the service account and authorised administrators. Never put it in the web proxy document root or share real values in this guide.')
code('AFRICA_CDC_DB_PATH=/var/lib/africa-cdc-assessment/africa_cdc_web.db\nAFRICA_CDC_REDCAP_API_URL=https://tools.africacdc.org/africacdcrc/api/\nAFRICA_CDC_REDCAP_TOKEN=<approved project token>\nAFRICA_CDC_REDCAP_AUTO_SYNC_MINUTES=0\nAFRICA_CDC_PASSWORD_PEPPER=<secret value>\nAFRICA_CDC_SECURE_COOKIES=true')
p('On Windows, replace only DB_PATH with C:/ProgramData/AfricaCDC/Assessment/africa_cdc_web.db. Use actual values instead of angle-bracket placeholders. Secure cookies require users to sign in over HTTPS, even though the proxy connects to Python over local HTTP.')
h('Decide: fresh installation or migration')
p('Migration: stop the old application, take a consistent SQLite backup and preserve its exact password pepper. Transfer both through approved secure channels. Do not generate a new pepper for existing accounts; changing it prevents existing Argon2id passwords from verifying. Preserve country-to-REDCap record mappings as well.')
p('Fresh installation: use an empty data directory and generate a random pepper before creating any accounts. Use the target virtual environment’s Python to run the following locally, then store the result in the protected .env and secret manager:')
code('python -c "import secrets; print(secrets.token_urlsafe(32))"')
p('Confirm the intended REDCap project and import/export API rights with its administrator. Existing deployment notes identify PID 932; verify this before using the token. A new REDCap project needs the application’s matching fields and repeating instruments, not just a token.')

page('2. Ubuntu: install the source package')
p('The examples use Ubuntu 24.04 LTS, /opt/africa-cdc-assessment for code and /var/lib/africa-cdc-assessment for data. Run installation commands with sudo rights. Paths are examples; keep them consistent if changed.')
h('Install prerequisites and create directories')
code('sudo apt update\nsudo apt install -y python3 python3-venv python3-pip nginx curl\nsudo useradd --system --home /opt/africa-cdc-assessment \\\n  --shell /usr/sbin/nologin africacdc\nsudo install -d -o root -g africacdc -m 750 /opt/africa-cdc-assessment\nsudo install -d -o africacdc -g africacdc -m 700 \\\n  /var/lib/africa-cdc-assessment')
h('Copy and install')
p('Extract the ZIP in a restricted staging folder. Copy the source files listed in section 1 into /opt/africa-cdc-assessment, including the complete static directory. Do not copy a development database, private .env, backups, .git or Windows executables into this fresh installation. Transfer a migration database separately into the data directory.')
code('sudo chown -R root:africacdc /opt/africa-cdc-assessment\nsudo chmod -R u=rwX,g=rX,o= /opt/africa-cdc-assessment\nsudo python3 -m venv /opt/africa-cdc-assessment/.venv\nsudo /opt/africa-cdc-assessment/.venv/bin/python -m pip install \\\n  -r /opt/africa-cdc-assessment/requirements.txt')
p('Dependencies are openpyxl, python-docx, argon2-cffi and python-dotenv. No Node.js build step or MySQL server is required by this application. Record the installed Python and dependency versions after successful acceptance testing.')
h('Set configuration and permissions')
code('sudo cp /opt/africa-cdc-assessment/.env.example \\\n  /opt/africa-cdc-assessment/.env\nsudo chown root:africacdc /opt/africa-cdc-assessment/.env\nsudo chmod 640 /opt/africa-cdc-assessment/.env\nsudo nano /opt/africa-cdc-assessment/.env')
p('Populate the settings in section 1. If migrating, place the stopped-instance database at the configured DB_PATH and give africacdc ownership with mode 600. The containing data directory must be writable so SQLite can create its journal files. Do not start the old and new installations against the same production workspace.')

page('3. Ubuntu: run continuously with systemd')
p('Create /etc/systemd/system/africa-cdc-assessment.service with the following contents. The application loads its protected adjacent .env itself.')
code('[Unit]\nDescription=Africa CDC Assessment Web Application\nAfter=network-online.target\nWants=network-online.target\n\n[Service]\nType=simple\nUser=africacdc\nGroup=africacdc\nWorkingDirectory=/opt/africa-cdc-assessment\nExecStart=/opt/africa-cdc-assessment/.venv/bin/python -u /opt/africa-cdc-assessment/server.py --host 127.0.0.1 --port 8080 --no-browser\nRestart=on-failure\nRestartSec=5\nUMask=0077\nNoNewPrivileges=true\nPrivateTmp=true\nProtectSystem=strict\nReadWritePaths=/var/lib/africa-cdc-assessment\n\n[Install]\nWantedBy=multi-user.target')
p('ExecStart must be one logical line in the service file, even if it wraps visually in this document. Do not use start_web_app.bat on Linux.')
h('Enable, start and check')
code('sudo systemctl daemon-reload\nsudo systemctl enable --now africa-cdc-assessment\nsudo systemctl status africa-cdc-assessment --no-pager\ncurl --fail http://127.0.0.1:8080/health\nsudo ss -ltnp | grep :8080')
p('Expected result: the service is active, /health returns JSON containing "status": "ok", and the listener is bound to 127.0.0.1:8080. A healthy endpoint does not prove REDCap saving works; complete section 7.')
h('Logs and service controls')
code('sudo journalctl -u africa-cdc-assessment -n 100 --no-pager\nsudo systemctl restart africa-cdc-assessment\nsudo systemctl stop africa-cdc-assessment')
p('With the read-only application directory, startup errors are captured in the systemd journal; do not rely on the optional application-adjacent log file. Keep the website restricted to IT until HTTPS and the first administrator account are configured.')

page('4. Ubuntu: publish through HTTPS / Nginx')
p('Replace assessment.example.org with the approved hostname. Obtain a valid certificate and key through the organisation’s certificate process and substitute their actual paths below. Ensure renewal is automated and tested.')
p('Create /etc/nginx/sites-available/africa-cdc-assessment:')
code('server {\n    listen 80;\n    server_name assessment.example.org;\n    return 301 https://assessment.example.org$request_uri;\n}\nserver {\n    listen 443 ssl;\n    server_name assessment.example.org;\n    ssl_certificate /etc/ssl/assessment/fullchain.pem;\n    ssl_certificate_key /etc/ssl/assessment/privkey.pem;\n    client_max_body_size 50m;\n    location / {\n        proxy_pass http://127.0.0.1:8080;\n        proxy_http_version 1.1;\n        proxy_set_header Host $host;\n        proxy_set_header X-Real-IP $remote_addr;\n        proxy_set_header X-Forwarded-For $proxy_add_x_forwarded_for;\n        proxy_set_header X-Forwarded-Proto $scheme;\n        proxy_read_timeout 180s;\n        proxy_send_timeout 180s;\n        proxy_cache off;\n    }\n}')
p('This is a starting proxy configuration, not a complete organisation-wide TLS policy. Nginx proxy directives are documented in reference [1]. The 50 MB limit and 180-second timeouts are deployment choices; validate them with representative imports and reports.')
code('sudo ln -s /etc/nginx/sites-available/africa-cdc-assessment \\\n  /etc/nginx/sites-enabled/africa-cdc-assessment\nsudo nginx -t\nsudo systemctl reload nginx\ncurl --fail https://assessment.example.org/health')
p('Resolve any conflicting site bindings without deleting unrelated hosted sites. Publish the app at the hostname root, not under an untested subpath such as /assessment. Allow 443 at host/network firewalls and 80 only where required. Do not expose 8080 or map the application folder as a static website.')

page('5. Windows Server: install and start')
p('Use an organisation-supported Windows Server release and a supported 64-bit Python 3 installation; Python 3.12 is a practical baseline to validate. Install IIS plus Microsoft URL Rewrite and Application Request Routing (ARR). See reference [2] for the proxy prerequisites.')
h('Prepare directories and access')
p('Place the source files listed in section 1 in C:\\Apps\\AfricaCDCAssessment. Create C:\\ProgramData\\AfricaCDC\\Assessment for writable data and C:\\inetpub\\AfricaCDCAssessmentProxy for the IIS proxy configuration only.')
p('Use a dedicated non-administrator service account. Grant Read & Execute on code/Python and Modify on the data directory. Restrict the adjacent .env to that account and authorised administrators. IIS must not serve the code, .env or database folders directly.')
h('Create the virtual environment')
code("Set-Location -LiteralPath 'C:\\Apps\\AfricaCDCAssessment'\npython -m venv .venv\n& '.\\.venv\\Scripts\\python.exe' -m pip install -r requirements.txt\nCopy-Item -LiteralPath '.env.example' -Destination '.env'\nnotepad .env")
p('Enter section 1 settings using the Windows DB_PATH. For migration, stop the old instance and copy its consistent database to that path; keep the original pepper. For a fresh site, the database is created automatically.')
h('Configure Task Scheduler')
p('Create a task named Africa CDC Assessment Web Application. Use the service account, Run whether user is logged on or not, an At startup trigger and a 30-second delay. Grant Log on as a batch job. Administrator privileges should not be needed at runtime.')
code('Program: C:\\Apps\\AfricaCDCAssessment\\.venv\\Scripts\\python.exe\nArguments: -u server.py --host 127.0.0.1 --port 8080 --no-browser\nStart in: C:\\Apps\\AfricaCDCAssessment')
p('Set restart on failure every minute for at least three attempts, remove the task duration limit and select Do not start a new instance. Configure log capture using the organisation’s task/service logging method; monitor task failures. Start the task and check:')
code("Invoke-RestMethod 'http://127.0.0.1:8080/health'\nGet-NetTCPConnection -LocalPort 8080 -State Listen")
p('Expected: status ok and a loopback listener. Complete IIS configuration before browser sign-in; secure cookies are intentionally enabled for HTTPS.')

page('6. Windows Server: IIS reverse proxy')
p('In IIS Manager, select the server, open Application Request Routing Cache > Server Proxy Settings and enable proxy. Set a 180-second proxy timeout as a starting value; validate long reports. Create an IIS site using C:\\inetpub\\AfricaCDCAssessmentProxy as its physical path.')
p('Bind the approved hostname on port 443 with its valid TLS certificate (SNI where needed), and on port 80 for redirection. Give the site identity read access only to the proxy directory. Create web.config there using the following content; replace assessment.example.org:')
code('''<?xml version="1.0" encoding="UTF-8"?>
<configuration>
  <system.webServer>
    <security><requestFiltering>
      <requestLimits maxAllowedContentLength="52428800" />
    </requestFiltering></security>
    <rewrite><rules>
      <rule name="HTTPS redirect" stopProcessing="true">
        <match url="(.*)" />
        <conditions>
          <add input="{HTTPS}" pattern="off" ignoreCase="true" />
        </conditions>
        <action type="Redirect"
          url="https://assessment.example.org/{R:1}"
          redirectType="Permanent" />
      </rule>
      <rule name="Assessment proxy" stopProcessing="true">
        <match url="(.*)" />
        <action type="Rewrite"
          url="http://127.0.0.1:8080/{R:1}"
          appendQueryString="true" />
      </rule>
    </rules></rewrite>
  </system.webServer>
</configuration>''')
p('Keep authenticated pages/API responses out of any proxy cache. Allow inbound HTTPS in Windows and network firewalls; do not allow public 8080. Open the HTTPS hostname and complete section 7. IIS URL Rewrite/ARR setup is documented in reference [2].')
p('Reboot once during acceptance testing. Verify that IIS and the Python task both recover without an interactive login and that the HTTPS health check still succeeds. Configure certificate-expiry monitoring and renewal ownership.')

page('7. REDCap and acceptance checks')
h('Confirm the REDCap project before writing')
p('The application expects its existing data dictionary, including country_assessment and repeating instruments working_groups, inventory_tools, system_profiling, work_assignments and gap_analysis_responses. Ask the REDCap administrator to verify the schema and API permissions. Do not import a dictionary over a live project without their review.')
p('Keep one authoritative deployment. Record IDs can be generated from local project IDs (acdc-N); a fresh local database pointed at an existing production REDCap project needs careful reconciliation to avoid record collisions. Preserve the database and mappings when migrating.')
h('Complete before opening access to participants')
p('Access: HTTPS certificate is valid; HTTP redirects to HTTPS; backend 8080 is not externally accessible. Confirm the first administrator is created by the authorised owner while access is restricted, or verify the migrated administrator can sign in. There is no shared default administrator password.')
p('Permissions: create or verify a clerk/group account, assign its phase and items, and confirm it cannot access unauthorised functions. Test logout and sign-in again.')
p('Persistence: use an approved UAT country/period in the agreed REDCap project. Enter sample Inventory, Profiling and Gap Analysis data; save successfully; verify the corresponding REDCap record. Restart the app, reload the same country/period and confirm the data returns.')
p('Temporary workspace: select the exact name Test Country. Confirm the temporary-only save notice, no REDCap record creation and disappearance of its data after restart. This check does not replace the persistent UAT test.')
p('Outputs: test a representative workbook import, comments and blank responses; download Excel and Word reports; check results from a second browser/device. Confirm the expected language options display.')
p('Recovery: perform a controlled REDCap outage test in UAT. A Member State save must fail visibly, not claim success. Keep entered values in the browser, restore connectivity and retry. Verify the REDCap record before closing the page.')
h('Optional report assistant')
p('Ollama runs on the server, not users’ devices. If required, install it under the organisation’s process, choose an installed model tag and set AFRICA_CDC_OLLAMA_URL (normally http://127.0.0.1:11434) and AFRICA_CDC_OLLAMA_MODEL. Restart the app and use Test provider, then generate a report. Size resources for the selected model; no capacity benchmark is supplied.')
p('Ollama is optional: the app has a built-in fallback. Keep its port private. Leave AFRICA_CDC_OPENAI_API_KEY unset unless the optional online analysis feature has been approved and separately configured.')

page('8. Operations, recovery and handover')
h('Backups and updates')
p('Schedule daily protected SQLite backups and confirm REDCap has its own assessment-data backup and restore process. Use a consistent SQLite backup, or stop the app before copying its database. Keep an encrypted off-server copy. Preserve the pepper and API credentials separately in the approved secret manager.')
p('Before an update, confirm all users have saved to REDCap, stop the application and retain the old code, dependencies and a consistent database backup. Replace application files only; preserve .env, the data directory and backups. Install requirements, restart and repeat the health, login, save/load and export checks. If rollback is needed, restore a compatible code/database pair and reconcile any REDCap writes made since the backup.')
p('Test restores in isolation with a test REDCap project/token, never an uncontrolled second writer to production. Record backup retention, acceptable data loss and recovery time with the application owner.')
h('Common problems')
p('502 / site unavailable: check the Python service/task, local /health, proxy binding, certificate and firewall. Address already in use: stop the duplicate process; keep one instance. SQLite read-only: check DB_PATH and the service account’s write access to its parent directory.')
p('Login loop: use HTTPS and confirm secure-cookie configuration. Migrated passwords fail: check the exact original pepper and .env precedence. REDCap saves fail: check outbound HTTPS, endpoint, token rights and matching schema; retain browser entries for retry. A healthy /health alone does not verify REDCap.')
h('IT handover record')
p('Production URL: ____________________    Server / OS: ____________________\nDeployed version / date: _____________________________________________\nIT owner / incident contact: __________________________________________\nApplication owner: __________________    REDCap owner: _________________\nBackup location / schedule: __________________________________________\nRestore test / reboot test date: _______________________________________\nTLS renewal owner: __________________    Acceptance sign-off: ____________')
p('Share the production URL, user access instructions and support contacts with the application owner. Transfer secrets separately. This guide documents setup; it does not certify deployment or testing on the recipient’s server.')
h('Technical references')
for t in ['[1] Nginx proxy module: https://nginx.org/en/docs/http/ngx_http_proxy_module.html', '[2] Microsoft IIS reverse proxy: https://learn.microsoft.com/en-us/iis/extensions/url-rewrite-module/reverse-proxy-with-url-rewrite-v2-and-application-request-routing', '[3] Python HTTP server: https://docs.python.org/3/library/http.server.html']:
    pp=doc.add_paragraph(t); pp.paragraph_format.line_spacing=1; pp.paragraph_format.space_after=Pt(3)
    for rr in pp.runs: rr.font.size=Pt(8)
doc.core_properties.title='Africa CDC Server Setup and Hosting Guide'
doc.core_properties.subject='IT deployment handover for application version 1.0.40'
doc.core_properties.author='Africa CDC Assessment Project'
OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print(OUT)
